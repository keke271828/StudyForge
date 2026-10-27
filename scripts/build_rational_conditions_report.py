"""Build an offline record of bounded rational tools and rejected prompt changes.

Saved cloud responses stay untouched. Local tool summaries and human teaching
notes are explicitly separate from model outputs. This script makes no request.
"""
import base64
import hashlib
from html import escape
import json
from pathlib import Path
import re
import shutil
from xml.etree import ElementTree as ET

from docx import Document
from lxml import etree

from scripts.build_teaching_comparison import TEMPLATE, inline_script

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'build/rational-conditions-acceptance-20261017'
TITLE = '有理式条件核验与补充讲解实际验收'

NOTES = {
    'chemistry_lecture': (
        '单位和给定数值的计算保留；5336字符，扩展很多。把常见情境中的恒容等条件归到模型的普遍适用边界，还断言随时间变化，超过原资料所给的信息。不能整体判为正确。',
        '4607字符、这次稍快；仍直接说模型假定恒容、充分混合、温度均匀，并增加原资料没有给出的过程断言。候选没有稳定解决本次关注的问题，因此未加入默认教学规则。'),
    'code_followup': (
        '两个静态结果都是2，区分函数依赖输入与不同输入必有不同输出。没有声称实际执行，623字符仍可压缩。',
        '两个结果与关键区分正确，461字符更直接，这次稍快。单个编程例较好，不足以抵消化学例中的补充条件问题。'),
    'historical_review': (
        '交付视图包含旧正文、核查后的局部替换与软件说明；不是模型原始JSON。发现两类问题，只有第0项局部修订成功；第1项一处原句含额外数学定界符，未通过精确匹配，因此整组未替换。仍漏掉“v不是常数”等目标表述，不能认证整篇。',
        '同一旧正文与快速核查，定位“才”的必要性措辞和单位总结，两项均由本机局部替换；这是一次核查，没有额外全文重写或独立复核。仍未覆盖本次关注的全部过度概括，单位问题也须结合前文已有约定解释，不能只按问题数量判优。'),
    'valid_control': (
        '正确短文在核查后保留，返回passed=true、issues为空。这里只说明这一份控制材料未被误改，不证明核查能识别所有正确证明。',
        '同样保留正确短文，未增加问题；本次6.00秒反而比原核查4.21秒慢。候选不能据此宣称普遍提速。'),
}


def saved_text(folder, result):
    path = folder / result['answer_file']
    raw = path.read_bytes()
    # Undo exactly the write_text Windows newline translation. Universal
    # newline reads also fold an existing CRLF in the older review draft and
    # would not reproduce the string whose fingerprint the experiment saved.
    original = raw.decode('utf-8').replace('\r\n', '\n')
    assert hashlib.sha256(original.encode()).hexdigest() == result['answer_sha256']
    return raw.decode('utf-8'), hashlib.sha256(raw).hexdigest()


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    native_dir = ROOT / 'build/rational-agent-real-20261017-1'
    component_dir = ROOT / 'build/supplement-conditions-real-20261017-1'
    native_cases = json.loads((native_dir / 'cases.json').read_text(encoding='utf-8'))
    native = json.loads((native_dir / 'results.json').read_text(encoding='utf-8'))
    inputs = json.loads((component_dir / 'inputs.json').read_text(encoding='utf-8'))
    components = json.loads((component_dir / 'results.json').read_text(encoding='utf-8'))
    assert len(native) == 2 and sum(r['model_requests'] for r in native) == 5
    assert len(components) == 8 and all(r['status'] == 'completed' and r['model_requests'] == 1 for r in components)
    groups, table_rows = [], []

    for case, result in zip(native_cases, native):
        assert case['id'] == result['case'] and result['status'] == 'completed'
        text, sha = saved_text(native_dir, result)
        calls = result['math_calls']
        call = next(c for c in calls if c['arguments']['operation'] == 'rational_identity')
        detail = call['result']['results']
        if case['id'] == 'cancelled_hole':
            tool_text = ('# 本机工具结果摘要\n\n这是对实际工具记录的人工整理，不是另一个模型的回答。\n\n'
                         '$$\\frac{x^2-1}{x-1}=x+1\\qquad(x\\ne1).$$\n\n'
                         '交叉相乘后的差为0；原式仍要求$x-1\\ne0$。右式没有分母限制。工具只核对共同有定义时的恒等，未自动判定两个函数的定义域相同，也不把$x=1$补进原函数。\n\n'
                         '下方请求记录保留完整输入和结果；本机核对本身没有模型调用，右侧导师续答仍有模型请求。')
            note = ('实际导师解释了缺口、原定义域与极限值的区别，没有出题。仍暴露了工具字段名，也把未定义的代入写成f(1)=0/0；更好的措辞是“直接代入出现0/0，原函数此处未定义”。不能把这份样例说成完全没有措辞问题。')
        else:
            tool_text = ('# 本机工具结果摘要\n\n这是对实际工具记录的人工整理，不是另一个模型的回答。\n\n'
                         '$$t\\cdot\\frac1t=1\\qquad(t\\ne0).$$\n\n'
                         '精确恒等核对保留原分母非零条件。原题给定$t>0$，工具没有应用或替换这个更小的范围；导师仍须按原题范围讨论。另核对两条有理式恒等与五个算式实例，数值实例不替代一般证明。\n\n'
                         '该完整任务实际用了3次模型请求和4次本机核对。工具没有证明导数法则、整篇推理、现实装置或教学质量。')
            note = ('真实完整DOCX任务成功并保留t>0；但5611字符明显过长。把表格第三行错叫第二行，将题设范围写成“我们额外加的限制”，又无依据断言现实量不会精确满足倒数关系。原始文件保留用于验收，含待核对问题，不作为定稿教材。')
        assert detail['identical_on_common_domain'] and detail['domain_comparison'] == 'not_compared'
        common = {'case': case['id'], 'response_model': 'deepseek-flash', 'input_chars': 0}
        groups.append({'round': '已接入实际Agent', 'case': case['id'], 'title': case['title'],
                       'question': case['question'], 'source': case['source'] + '\n\n左侧为本机工具摘要，右侧为实际导师交付，不是两个模型方案的质量对照。',
                       'rows': [
                           {**common, 'arm': 'local_tool', 'label': '本机工具摘要（人工整理，零模型调用）',
                            'text': tool_text, 'sha256': hashlib.sha256(tool_text.encode()).hexdigest(),
                            'seconds': '未单独计时', 'output_chars': len(tool_text), 'usage': [],
                            'request': calls, 'review': '只确认有界有理式在公共定义域上的恒等与原分母条件；不负责认证整篇讲解或学生掌握程度。'},
                           {**common, 'arm': 'actual_delivery', 'label': '真实导师／DOCX交付（未润色）',
                            'text': text, 'sha256': sha, 'seconds': round(result['seconds'], 2),
                            'output_chars': result['characters'], 'usage': [t['usage'] for t in result['http_attempts']],
                            'request': {'source': case['source'], 'question': case['question'],
                                        'actual_request_count': result['model_requests'], 'traces': result['http_attempts'],
                                        'public_final_model_texts': json.loads((native_dir / (case['id'] + '-public-final-calls.json')).read_text(encoding='utf-8'))},
                            'review': note},
                       ]})
        table_rows.append('<tr><td>实际Agent</td><td>' + escape(case['title']) + '</td><td>本机工具：不单独计时</td><td>'
                          + f"{result['seconds']:.2f}秒／{result['model_requests']}次模型请求／{len(calls)}次本机核对" + '</td></tr>')

    for case in inputs['cases']:
        by_arm = {r['arm']: r for r in components if r['case'] == case['id']}
        rows = []
        for i, arm in enumerate(('baseline', 'candidate')):
            result = by_arm[arm]
            text, sha = saved_text(component_dir, result)
            records = json.loads((component_dir / (case['id'] + '-' + arm + '-records.json')).read_text(encoding='utf-8'))
            assert len(records) == 1
            if case['role'] != 'review':
                assert records[0]['response'] == text.replace('\r\n', '\n')
            rows.append({'case': case['id'], 'arm': arm,
                         'label': ('原规则' if arm == 'baseline' else '候选措辞（未采用）') + (' · 核查后交付' if case['role'] == 'review' else ' · 模型原始回答'),
                         'text': text, 'sha256': sha, 'review': NOTES[case['id']][i],
                         'seconds': round(result['seconds'], 2), 'output_chars': result['characters'],
                         'input_chars': len(records[0]['system']) + len(records[0]['user']),
                         'response_model': result['http_attempts'][0]['returned_model'],
                         'usage': [t['usage'] for t in result['http_attempts']],
                         'request': {'actual_public_input_and_raw_response': records, 'workflow_reports': result.get('reports', [])}})
        source = case['source'] + '\n\n预先固定的核对点：\n' + '\n'.join(case['checks'])
        if case.get('draft'):
            source += '\n\n送入核查的完整旧正文（AI内容）：\n' + case['draft']
        groups.append({'round': '未采用的提示词实验', 'case': case['id'], 'title': case['title'],
                       'question': case['question'], 'source': source, 'rows': rows})
        a, b = by_arm['baseline'], by_arm['candidate']
        table_rows.append('<tr><td>措辞对照</td><td>' + escape(case['title']) + '</td><td>'
                          + f"{a['seconds']:.2f}秒／{a['characters']}字符" + '</td><td>'
                          + f"{b['seconds']:.2f}秒／{b['characters']}字符" + '</td></tr>')

    latest, suites = {}, {}
    for name in ('rational-tools-final-regression-20261017.xml', 'rational-tools-context-regression-20261017.xml'):
        cases = ET.parse(ROOT / 'build' / name).getroot().findall('.//testcase')
        suites[name] = len(cases)
        for case in cases:
            assert all(case.find(k) is None for k in ('failure', 'error', 'skipped'))
            latest[(case.get('classname'), case.get('name'))] = case
    assert len(latest) == 250
    ui_dir = ROOT / 'build/rational-tools-ui-20261017-2'
    ui = json.loads((ui_dir / 'ui-results.json').read_text(encoding='utf-8'))
    assert ui['passed'] and len(ui['checks']) == 4 and ui['controlled_model_requests'] == 4
    for name in ('电脑端数学核验.png', '手机端数学核验.png'):
        shutil.copyfile(ui_dir / name, OUT / name)
    docx_name = '变化因子与恒定乘积_实际讲解.docx'
    shutil.copyfile(native_dir / docx_name, OUT / docx_name)
    doc = Document(OUT / docx_name)
    text = ''.join(etree.fromstring(doc.part.blob).xpath('//*[local-name()="t"]/text()'))
    assert len(doc.paragraphs) > 20 and len(doc.tables) == 1 and 't>0' in re.sub(r'\s+', '', text)

    intro = '''<header><div class="eyebrow">STUDYFORGE · CONDITIONS &amp; TEACHING · 2026-10-17</div>
<h1>约分以后，原来的限制还在</h1><p>新增有界有理式核对；保留原始讲解，也公开工具核对不了的教学问题。</p>
<p><a href="../../验收最新版.vbs">最新版源码验收入口</a> · <a href="../../docs/有理式条件核验与补充讲解验收.md">详细记录</a>。本轮未打包。</p>
<div class="notice">本轮共<strong>13次真实模型请求</strong>：8次提示词组件实验、2项实际Agent任务合计5次请求。下面有12份验收视图，其中2份是本机工具的人工摘要、2份是真实Agent交付、8份是组件对照交付；不是12次独立模型回答。模型为deepseek-flash／high，未做网页端比较。<strong>提示词候选没有稳定胜出，未采用。</strong></div></header>'''
    evidence = r'''<details class="prompt"><summary>展开：这次加了什么，实际讲解还差在哪里</summary>
<h2>已经接入共同Agent的能力</h2><p>此前本机数学工具能核对精确算术、多项式恒等和形式偏导，但不能保留有理式约分前的分母限制。本次新增有理式恒等操作：以有理数精确运算，在两边都有定义时交叉相乘核对，并分别保留原来的非零条件。约分、乘零、负幂或嵌套除法都不抹掉限制。</p>
<p>例如(x²−1)/(x−1)和x+1在x≠1时相等，原函数在1处仍无定义；t·(1/t)=1仍有t≠0的条件，题设给定t&gt;0时还须按原范围讨论。工具不会把“尚未比较定义域”误报为“已证明定义域不同”。它不求分母根，不检查完整定理、极限、不等式或真实系统，也不按工具结果给学生评分。</p>
<p>这是本机按需计算，不是每份讲解新增一轮审查。只读权限可用，关闭Agent工具时不调用。限定6变量、每式2000字符／160语法节点、每个中间多项式总次数24／128项，并限制总计算量和结果大小；不执行代码、不联网、不启动额外模型。模型选择调用工具时仍可能需要续答，右侧实际任务分别用了2次、3次模型请求，不能把本机核对称为整个任务零token。</p>
<h2>为什么没有继续叠加提示词</h2><p>用化学完整讲解、编程追问、真实旧讲解快速核查和正确短文控制，分别比较原措辞与“补充假设不能反过来当普遍必要条件”的候选，共8次请求。编程候选更短且正确，但化学仍添加未给定的情境条件，两版审查也漏掉目标表述；正确控制两版都保留。没有足够证据采用候选。</p>
<p>核查视图是“旧正文＋本机局部替换＋软件说明”的交付文本，不是原始模型JSON。每份展开记录同时保留模型原始JSON、完整输入和修订报告，二者可区分。原核查的一项修订因原句数学定界符不一致，没有精确匹配成功，仍按未修正保留，没有猜测替换。</p>
<h2>真实Word仍有的问题</h2><p><strong>实际DOCX成功导出并保留原题范围，但不是定稿教材。</strong>讲解5611字符，扩展到导数、对数导数、梯度及多次总结；表格第三行被误称第二行，把原题范围称作“我们额外加的限制”，又无依据外推真实量不满足倒数关系。工具核对通过不等于这些文字正确。此处原样交付，未润色后冒充模型原答。</p>
<p><a href="变化因子与恒定乘积_实际讲解.docx">下载原始Word验收样例（含上述待核对问题）</a>。已重新读取文档结构与公式文字，65个段落、1张表，t&gt;0保留；本轮没有在Microsoft Word里进行视觉验收。</p>
<details><summary>人工整理：同一问题可以怎样更直接地讲清（非模型原始回答）</summary>
<p>两个因子各自变化，只说明各自取值不固定；乘积还取决于两者怎样配合。原题给定t&gt;0、u=t、v=1/t，直接代入便有P=t/t=1。t从1变成2时，u从1变成2，v从1变成1/2，乘积都为1。这就反驳了“两个因子都变，所以乘积一定变”。</p>
<p>关键是v始终是u的倒数，而不是仅仅“一增一减”；例如u=t、v=1/t²时乘积为1/t，也是一增一减但乘积在变。分母要求t≠0，原题已经指定更小范围t&gt;0，继续保留即可。这个数学反例没有决定未给出条件的真实系统如何变化。只有需要变化率或几何推广时，再展开导数与等值线。</p></details>
<h2>真实耗时与范围</h2><div class="report-table"><table><thead><tr><th>阶段</th><th>内容</th><th>原规则／本机摘要</th><th>候选／实际任务</th></tr></thead><tbody>__ROWS__</tbody></table></div>
<p>每组每臂只测一次。8次组件请求与原生Agent任务的组织方式不同，不能混作A/B速度比较。缓存、推理长度、服务排队都会影响时间；不把12.09秒和42.27秒推广成书籍、视频或所有文档的耗时。保存了公共回答、请求或输入指纹与服务返回用量，没有保存私有思考正文。</p>
<h2>验收与打开方式</h2><p>250项最终相关程序检查通过，两批171项与125项按测试标识去重。覆盖原分母条件、有界运算、非法输入、精确数值独立对照、只读／禁用权限、私有核查结果隔离、原生服务工具续答、完整DOCX任务、快速核查、对话画像范围和客户端释放。受控模型用于这些程序流程，不当作真实教学质量证据。</p>
<p>4项实际Electron电脑／配对HTTPS手机宽度界面检查通过：工具完成与续答、重载保留回答、手机流程、公开记录不包含私有计算结果且不生成分数／题单。该界面测试共4次受控模型响应、零外部模型请求。手机为390px浏览器验收，未替代真机安卓测试。</p>
<p>指南和手机页面版本v55同步。退出旧源码实例后，双击本页顶部“最新版源码验收入口”，可问导师“请用核验工具检查约分前后的分母条件”。已安装的旧EXE不会随源码更新自动变化。本轮没有测出全应用RAM下降比例，也没有宣布视频处理和教学质量问题全部解决。</p>
<details><summary>实际界面截图（受控正文，仅验证交互）</summary>
<figure><img src="电脑端数学核验.png" alt="电脑数学工具完成后等待导师续答的界面"><figcaption>电脑：本机数学核验已完成、导师仍在续答时的截图；之后实际等待完整回答并验证重载。</figcaption></figure>
<figure><img class="phone-shot" src="手机端数学核验.png" alt="配对HTTPS手机保留数学定义域条件"><figcaption>手机宽度：完整受控回答与原定义域条件；公式无渲染错误。</figcaption></figure></details></details>'''
    evidence = evidence.replace('__ROWS__', ''.join(table_rows))
    page = re.sub(r'<header>.*?</header>', lambda _: intro, TEMPLATE, flags=re.S).replace('<main>', '<main>' + evidence, 1)
    page = page.replace('<title>StudyForge · 讲解质量真实对照</title>', f'<title>{TITLE} · StudyForge</title>')
    page = page.replace('grid-template-columns:repeat(3,minmax(0,1fr))', 'grid-template-columns:repeat(2,minmax(0,1fr))')
    page = page.replace('</style>', '.report-table{overflow-x:auto}table{border-collapse:collapse;width:100%;font-size:14px}td,th{padding:10px;border:1px solid #d2e0d6;text-align:left}a{color:#236b5b;overflow-wrap:anywhere}.prompt{overflow-wrap:anywhere}.lesson-math.display-math{display:block;overflow-x:auto;overflow-y:hidden;padding:6px 0}.katex-display{overflow:visible}figure{margin:20px 0}figure img{display:block;max-width:100%;height:auto;border:1px solid #d8e2df;border-radius:12px}figure .phone-shot{max-width:min(100%,390px)}figcaption{font-size:14px;color:#53656a}</style>', 1)
    css = (ROOT / 'web/vendor/katex.css').read_text(encoding='utf-8')
    css = re.sub(r'url\(fonts/([^)]*\.woff2)\)', lambda m: 'url(data:font/woff2;base64,' + base64.b64encode((ROOT / 'web/vendor/fonts' / m[1]).read_bytes()).decode() + ')', css)
    css = re.sub(r',url\(fonts/[^)]*\) format\("(?:woff|truetype)"\)', '', css)
    page = page.replace('__KATEX_CSS__', css).replace('__PAYLOAD__', json.dumps({'groups': groups, 'count': 12, 'failures': []}, ensure_ascii=False).replace('<', '\\u003c'))
    page = page.replace('__LIBRARIES__', '\n'.join(inline_script((ROOT / 'web/vendor' / name).read_text(encoding='utf-8')) for name in ('marked.js', 'purify.js', 'katex.js')))
    page = page.replace('__RENDERER__', inline_script((ROOT / 'web/lesson_render.js').read_text(encoding='utf-8')))
    page = page.replace('份原始回答', '份验收视图').replace('秒 · 正文 ', '秒 · 交付 ')
    page = page.replace("row.seconds+'秒 ·", "row.seconds+(typeof row.seconds==='number'?'秒':'')+' ·")
    page = page.replace('下载这份原始回答', '下载这份交付文本').replace('查看本次提供的完整合成原资料', '查看资料、要求与本次核对范围')
    page = page.replace('核对实际请求、用量与原文指纹', '核对公开输入／结果、用量与交付指纹')
    page = page.replace('所有数据与字体已嵌入此文件；阅读、切换与下载不发送模型请求。修订前有缺陷的回答保留用于对照，切勿把它们当作定稿教材。', '验收视图与字体已内嵌，Word和截图在同一文件夹。工具摘要为人工整理，核查交付含局部修订，原始模型响应可展开核对。查看本页不请求模型；有缺陷的讲解不当定稿教材。')
    target = OUT / (TITLE + '.html')
    target.write_text(page, encoding='utf-8')
    manifest = {'answers': 12, 'groups': 6, 'report_table_rows': [6], 'program_checks': len(latest), 'suites': suites,
                'ui_checks': len(ui['checks']), 'real_model_requests': 13, 'component_requests': 8,
                'native_agent_requests': 5, 'native_math_calls': sum(len(r['math_calls']) for r in native),
                'candidate_policy_adopted': False, 'package_created': False,
                'native_docx_structurally_reopened': True, 'docx_positive_scope_retained': True,
                'docx_paragraphs': len(doc.paragraphs), 'docx_tables': len(doc.tables),
                'output_sha256': hashlib.sha256(target.read_bytes()).hexdigest()}
    (OUT / 'manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({'path': str(target), **manifest}, ensure_ascii=True))


if __name__ == '__main__':
    main()
