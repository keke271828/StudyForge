"""Real provider smoke using an isolated project; DeepSeek forbids Codex calls."""
import asyncio
import json
import os
import tempfile
from pathlib import Path
from contextlib import ExitStack
from unittest.mock import patch

root = Path(__file__).resolve().parents[1]
provider = os.environ.get('STUDYFORGE_SMOKE_PROVIDER', 'codex')
codex_attempts = []
target = Path(tempfile.mkdtemp(prefix='shared-agent-', dir=root/'build'))
os.environ['STUDYFORGE_DATA_DIR'] = str(target/'data')

from studyforge import server
from studyforge.agent_tools import AgentRuntime, search_learning_web
from studyforge.codex_bridge import CodexBridge
from studyforge.ai import save_settings
from studyforge.db import create_job, create_project, execute, j, many, now, one, uid


async def chat(value):
    response = await server.chat_stream(value)
    events = [json.loads(line) async for line in response.body_iterator]
    failures = [item['value'] for item in events if item['type']=='error']
    if failures: raise RuntimeError(str(failures))
    return events


async def main():
    model_override = os.environ.get('STUDYFORGE_SMOKE_MODEL')
    save_settings({'chat_provider': provider, **({'deepseek_model': model_override} if provider=='deepseek' and model_override else {})})
    project = create_project('统一导师真实验收'); source=uid(); original=target/'因果资料.txt'
    original.write_text('相关描述两个变量共同变化，不能单凭相关推出因果。\n共同原因可以同时影响两者。\n观察到冰淇淋销量和溺水人数都增加，天气可能是共同原因。\n末尾关键条件：实验结果能否推广到其他人群涉及外部效度。', encoding='utf-8')
    execute('INSERT INTO sources(id,project_id,name,kind,original_path,segments_json,status,created_at) VALUES (?,?,?,?,?,?,?,?)',
            (source,project['id'],original.name,'.txt',str(original),'[]','ready',now()))
    model=(server.get_settings().get('codex_local_default') or server.settings().get('codex_model')) if provider=='codex' else server.settings().get('deepseek_model')
    effort=None
    read=server.ReadIn(project_id=project['id'],source_ids=[source],provider=provider,model=model,generation_quality='custom')
    read_job=create_job(project['id'],'read',read.model_dump())
    await server._read_job(read_job,read)
    assert one('SELECT status FROM jobs WHERE id=?',(read_job,))['status']=='completed'
    first=await chat(server.ChatIn(project_id=project['id'],source_ids=[source],model=model,reasoning_effort=effort,provider=provider,
        text='请实际调用 sf_read_source 读取当前原资料，随后调用 sf_write_text 在项目工作目录创建 lesson.txt，写入“相关不能证明因果”。之后用两句话解释共同原因，不生成文档，也不联网。'))
    conversation=next(item['value'] for item in first if item['type']=='conversation')
    first_thread=one('SELECT remote_thread_id FROM conversations WHERE id=?',(conversation,))['remote_thread_id']
    value=server.GenerateIn(project_id=project['id'],source_ids=[source],conversation_id=conversation,model=model,reasoning_effort=effort,provider=provider,
        instructions='只用原资料写一份简短讲解，约300字，包含相关、因果、共同原因的概念对照表，解释外部效度；不用习题。',generation_quality='custom')
    job=create_job(project['id'],'generate',value.model_dump())
    await server._generate_job(job,value)
    completed=one('SELECT * FROM jobs WHERE id=?',(job,))
    if completed['status']!='completed': raise RuntimeError(completed['error'])
    followup=await chat(server.ChatIn(project_id=project['id'],source_ids=[source],conversation_id=conversation,model=model,provider=provider,
        text='接着你刚才生成的文档，用两句话解释为什么还要讨论外部效度。不要重新生成文档。'))
    end_thread=one('SELECT remote_thread_id FROM conversations WHERE id=?',(conversation,))['remote_thread_id']
    tools=many('SELECT tool,status FROM agent_tool_events WHERE conversation_id=?',(conversation,))
    assert first_thread==end_thread
    assert any(row['tool']=='sf_write_text' and row['status']=='completed' for row in tools)
    assert any(row['tool']=='sf_read_source' and row['status']=='completed' for row in tools)
    assert many('SELECT * FROM mentor_handoffs')==[]
    assert len(many('SELECT * FROM generation_dialogue_outputs'))==1
    assert not codex_attempts, '禁止调用 Codex，包括被其他代码捕获的失败调用'
    lecture=one("SELECT path FROM artifacts WHERE kind='explanation'")['path']
    from docx import Document
    doc=Document(lecture); assert doc.tables
    search=await search_learning_web('correlation causation common cause',3)
    result={'provider':provider,'model':model,'codex_disabled':provider=='deepseek','same_dialogue':True,'same_remote_thread':True if provider=='codex' else None,'tools':tools,'word_tables':len(doc.tables),
            'lecture':lecture,'followup':next(item['value'] for item in followup if item['type']=='replace'),
            'search_results':len(search['results']),'search_failures':search['failures']}
    (target/'result.json').write_text(j(result),encoding='utf-8')
    print(j(result),flush=True)
    await server.bridge.close()

with ExitStack() as stack:
    if provider=='deepseek':
        async def forbid_codex(*args,**kwargs):
            codex_attempts.append(True)
            raise AssertionError('DeepSeek 独立验收禁止调用 Codex')
        actual_which=server.shutil.which
        stack.enter_context(patch.object(CodexBridge,'start',forbid_codex))
        stack.enter_context(patch.object(server.model_catalog,'get_models',forbid_codex))
        stack.enter_context(patch.object(server.model_catalog,'validate_effort',forbid_codex))
        stack.enter_context(patch.object(server.shutil,'which',lambda command:None if command=='codex' else actual_which(command)))
    asyncio.run(main())
