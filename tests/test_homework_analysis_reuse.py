"""Preserve validated per-question teaching while a different cause is clarified."""
import asyncio
import copy
import json

import pytest

from studyforge import homework
from studyforge.db import j,unj


def data():
    worksheet={'id':'set-a','title':'分题补讲','difficulty':2,
        'questions_json':j([{'index':i,'topic':topic,'question':question} for i,(topic,question) in enumerate([
            ('方程','解2x+3=7并写出步骤。'),('加法','计算2+3并说明步骤。'),('电通量','净电荷为零是否说明电场处处为零？说明理由。')])]),
        'solutions_json':j([{'index':i,'solution':solution,'rubric':'说明正确依据。'} for i,solution in enumerate([
            'x=2，先减3再除以2。','2+3=5。','不能。均匀电场穿入与穿出通量抵消，但场非零。'])])}
    row={'id':'attempt-a','exercise_set_id':'set-a','project_id':'project-a','conversation_id':'chat-a',
        'source_version':homework.version(worksheet),'answers_json':j(['4','2+3=4','净通量为零，所以电场处处为零。']),
        'analysis_json':'{}','options_json':j({'provider':'deepseek','model':'fixture-model','reasoning_effort':'high',
            'teaching':{'learning_mode':'deep','explanation_style':'balanced','output_language':'zh'},
            'profile':{'items':[{'topic':'电通量','summary':'曾询问通量概念，仅为历史线索。'}]}})}
    grade={'scores':[30,70,40],'score':46.7,'feedback':['缺少步骤。','计算错误。','净通量与局部场混淆。']}
    return row,worksheet,grade


class Teacher:
    model='fixture-model';effort='high'
    def __init__(self):self.calls=[]
    async def complete(self,system,user,**kwargs):
        packet=json.loads(user);self.calls.append(copy.deepcopy(packet));items=[]
        for index in packet['wrong_indices']:
            causes=[cause for cause in packet['student_causes'] if cause['index']==index]
            uncertain=index==0 and not causes
            items.append({'index':index,'kind':'unknown' if uncertain else 'implementation' if index<2 else 'terminology',
                'confidence':'uncertain' if uncertain else 'self_report' if causes else 'supported',
                'cause':'原因尚需澄清。' if uncertain else '只根据对应的本次作答或自报定位。',
                'evidence':([] if uncertain else [{'source':'questionnaire' if causes else 'answer',
                    'quote':causes[0]['text'] if causes else packet['student_answers'][index]}]),
                'clarification':'你当时怎样得到这个结果？',
                'explanation':f'第{index+1}题的第{len(self.calls)}次补讲。'+
                    ('先核对等式两边同减3，再同时除以2；x=2，代回2×2+3=7。每次操作都要保持等式两边相等。' if index==0 else
                     '从2开始依次加三个1，得到3、4、5，所以2+3=5。用逆运算5减3等于2来核对；这能检查本次加法，不能说明全部计算都已掌握。' if index==1 else
                     '通量计算的是穿过闭合面的有向总量；均匀场一侧进入、一侧出去，净通量相消，场仍然不为零。整体积分为零不能推出各点都为零。'),
                'check':'核对原题条件并将结果代回，区分本次检查与一般证明。'})
        return j({'items':items})


def first():
    row,worksheet,grade=data();teacher=Teacher()
    result=asyncio.run(homework.analyze(teacher,row,worksheet,grade,[]));row['analysis_json']=j(result)
    return row,worksheet,grade,teacher,result


CAUSE=[{'index':0,'kinds':['implementation'],'text':'我算出7减3为4后忘了再除以2。'}]


def test_clarifying_one_question_keeps_other_validated_teaching_verbatim():
    row,worksheet,grade,teacher,previous=first()
    result=asyncio.run(homework.analyze(teacher,row,worksheet,grade,CAUSE))
    assert [call['wrong_indices'] for call in teacher.calls]==[[0,1,2],[0]]
    assert result['items'][1:]==previous['items'][1:]
    assert result['items'][0]['confidence']=='self_report'
    assert result['analysis_reuse']['reused_indices']==[1,2]
    assert result['analysis_reuse']['requested_indices']==[0]
    # Keep original context, including cross-question dependencies, in the request.
    assert teacher.calls[0]['questions']==teacher.calls[1]['questions']
    assert teacher.calls[0]['reference_solutions']==teacher.calls[1]['reference_solutions']
    assert teacher.calls[0]['student_answers']==teacher.calls[1]['student_answers']


def test_identical_inputs_need_no_second_model_request_even_for_unresolved_guidance():
    row,worksheet,grade,teacher,previous=first()
    result=asyncio.run(homework.analyze(teacher,row,worksheet,grade,[]))
    assert len(teacher.calls)==1 and result['items']==previous['items']
    assert result['analysis_reuse']['reused_indices']==[0,1,2]
    assert result['items'][0]['confidence']=='uncertain'  # Reuse never turns it into a diagnosis.


def test_result_is_independent_of_mutations_to_a_previous_return_value():
    row,worksheet,grade,teacher,previous=first()
    result=asyncio.run(homework.analyze(teacher,row,worksheet,grade,CAUSE))
    result['items'][1]['explanation']='外部改动'
    assert previous['items'][1]['explanation']!='外部改动'


@pytest.mark.parametrize('change',[
    'attempt','project','conversation','set','version','questions','solutions','answers','grade',
    'style','mode','language','profile','model','provider','effort',
])
def test_changed_identity_or_teaching_basis_does_not_reuse_stale_results(change):
    row,worksheet,grade,teacher,_=first()
    if change in ('attempt','project','conversation','set'):
        key={'attempt':'id','project':'project_id','conversation':'conversation_id','set':'exercise_set_id'}[change];row[key]+='-new'
    elif change=='version':row['source_version']='changed-version'
    elif change in ('questions','solutions'):
        key=change+'_json';items=unj(worksheet[key]);items[0]['question' if change=='questions' else 'solution']+='新条件。';worksheet[key]=j(items)
    elif change=='answers':
        answers=unj(row['answers_json']);answers[0]='8';row['answers_json']=j(answers)
    elif change=='grade':grade['feedback'][0]='更正后的评分依据。'
    else:
        options=unj(row['options_json'])
        if change in ('style','mode','language'):
            options['teaching'][{'style':'explanation_style','mode':'learning_mode','language':'output_language'}[change]]={'style':'clear','mode':'research','language':'en'}[change]
        elif change=='profile':options['profile']={'items':[]}
        else:options[{'model':'model','provider':'provider','effort':'reasoning_effort'}[change]]='changed'
        row['options_json']=j(options)
    result=asyncio.run(homework.analyze(teacher,row,worksheet,grade,CAUSE))
    assert teacher.calls[-1]['wrong_indices']==[0,1,2] and not result['analysis_reuse']['reused_indices']


def test_changed_actual_client_model_invalidates_old_teaching():
    row,worksheet,grade,teacher,_=first();teacher.model='another-service-model'
    asyncio.run(homework.analyze(teacher,row,worksheet,grade,CAUSE))
    assert teacher.calls[-1]['wrong_indices']==[0,1,2]


@pytest.mark.parametrize('corruption',['legacy','version','receipt','duplicate','edited-item'])
def test_legacy_or_damaged_evidence_falls_back_instead_of_trusting_a_partial_cache(corruption):
    row,worksheet,grade,teacher,previous=first()
    if corruption=='legacy':previous.pop('analysis_reuse',None)
    elif corruption=='version':previous['analysis_reuse']['version']='unrecognized'
    elif corruption=='receipt':previous['analysis_reuse']['entries']['1']['input_sha256']='wrong'
    elif corruption=='duplicate':previous['items'].append(copy.deepcopy(previous['items'][1]))
    else:previous['items'][1]['explanation']='被替换的补讲，不能冒称此前已核对。'
    row['analysis_json']=j(previous)
    result=asyncio.run(homework.analyze(teacher,row,worksheet,grade,CAUSE))
    assert 1 in teacher.calls[-1]['wrong_indices']
    assert 1 not in result['analysis_reuse']['reused_indices']
    assert '被替换' not in result['items'][1]['explanation']


def test_only_changed_questionnaire_answer_invalidates_its_own_explanation():
    row,worksheet,grade,teacher,_=first()
    second=asyncio.run(homework.analyze(teacher,row,worksheet,grade,CAUSE));row['analysis_json']=j(second)
    changed=copy.deepcopy(CAUSE);changed[0]['text']='我误把等式两边同减3写成只减右边。'
    third=asyncio.run(homework.analyze(teacher,row,worksheet,grade,changed))
    assert teacher.calls[-1]['wrong_indices']==[0] and third['items'][1:]==second['items'][1:]


def test_saved_state_contains_only_small_hashes_not_another_copy_of_answers_or_profile():
    _,_,_,_,result=first();receipt=j(result['analysis_reuse'])
    assert '净通量为零' not in receipt and '曾询问' not in receipt and '2+3=4' not in receipt
    assert len(receipt)<2000


def test_corrected_grade_with_no_errors_skips_analysis_and_old_cached_items():
    row,worksheet,grade,teacher,_=first();grade.update(scores=[100,100,100],score=100)
    assert asyncio.run(homework.analyze(teacher,row,worksheet,grade,[]))=={'items':[]}
    assert len(teacher.calls)==1


def test_native_questionnaire_resume_preserves_other_question_and_exports_both_documents(monkeypatch):
    from tests.test_homework import fixture,finish
    from studyforge import server
    from studyforge.db import many,one
    from docx import Document
    project,chat,set_id,teacher,value=fixture(monkeypatch)
    async def scenario():
        result=await server.submit_answers(set_id,value)
        assert (await finish(result))['status']=='paused'
        before=homework.get_attempt(result['id'])
        previous=unj(before['analysis_json'])
        assert previous['items'][0]['confidence']=='uncertain'
        assert before['solution_artifact_id'] and not before['feedback_artifact_id']
        response=server.HomeworkCausesIn(responses=[{'index':0,'kinds':['terminology'],'text':'我混淆了电场与整个面的净通量。'}])
        await server.homework_causes(result['id'],response,project['id'])
        final=await finish(result)
        assert final['status']=='completed',final['error']
        return before,homework.get_attempt(result['id']),previous
    before,after,previous=asyncio.run(scenario())
    final=unj(after['analysis_json'])
    assert final['analysis_reuse']['requested_indices']==[0] and final['analysis_reuse']['reused_indices']==[1]
    assert final['items'][1]==previous['items'][1]
    assert after['solution_artifact_id']==before['solution_artifact_id'] and after['feedback_artifact_id']
    assert [json.loads(user)['wrong_indices'] for system,user in teacher.calls if '逐题诊断错因并补讲' in system]==[[0,1],[0]]
    assert len([system for system,_ in teacher.calls if '批改者' in system])==2
    assert len(many('SELECT * FROM submissions'))==1
    artifact=one('SELECT * FROM artifacts WHERE id=?',(after['feedback_artifact_id'],))
    text='\n'.join(paragraph.text for paragraph in Document(artifact['path']).paragraphs)
    assert final['items'][1]['explanation'] in text and '针对性讲解' in text
    assert all(item['project_id']==project['id'] for item in many('SELECT project_id FROM artifacts'))


def test_authenticated_phone_can_resume_only_the_clarified_question_and_download_complete_feedback(monkeypatch):
    from io import BytesIO
    from docx import Document
    from tests.test_homework import fixture
    from tests.test_mobile_access import mobile_client,pair,await_phone_attempt
    from studyforge.db import one
    project,chat,set_id,teacher,value=fixture(monkeypatch)
    manager,client=mobile_client()
    with client:
        pair(manager,client)
        response=client.post('/api/exercises/'+set_id+'/submit',json=value.model_dump())
        assert response.status_code==200,response.text
        submitted=response.json();assert submitted['solution_artifact'] and not submitted['feedback_artifact']
        attempt=await_phone_attempt(client,submitted,{'awaiting_causes'})
        before=unj(homework.get_attempt(attempt['id'])['analysis_json'])
        route='/api/homework/'+attempt['id']+'/causes?project_id='+project['id']
        response=client.post(route,json={'responses':[{'index':0,'kinds':['terminology'],'text':'我混淆了电场和整个闭合面的净通量。'}]})
        assert response.status_code==200,response.text
        done=await_phone_attempt(client,attempt,{'completed'})
        after=unj(homework.get_attempt(attempt['id'])['analysis_json'])
        assert after['analysis_reuse']['requested_indices']==[0] and after['analysis_reuse']['reused_indices']==[1]
        assert after['items'][1]==before['items'][1]
        download=client.get(done['feedback_artifact']['download_url'])
        assert download.status_code==200
        words='\n'.join(p.text for p in Document(BytesIO(download.content)).paragraphs)
        assert '第 1 题' in words and '第 2 题' in words and after['items'][1]['explanation'] in words
        assert one('SELECT COUNT(*) AS n FROM submissions')['n']==1
