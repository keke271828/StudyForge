"""Real DeepSeek tools and later-turn image reading, with Codex unavailable."""
import asyncio
import json
import os
import tempfile
from pathlib import Path
from unittest.mock import patch

root = Path(__file__).resolve().parents[1]
target = Path(tempfile.mkdtemp(prefix='agent-media-', dir=root / 'build'))
os.environ['STUDYFORGE_DATA_DIR'] = str(target / 'data')

from docx import Document
from PIL import Image, ImageDraw
from studyforge import server
from studyforge.ai import save_settings
from studyforge.agent_tools import AgentRuntime
from studyforge.codex_bridge import CodexBridge
from studyforge.db import create_project, execute, j, many, now, uid


async def chat(value):
    response = await server.chat_stream(value)
    events = [json.loads(item) async for item in response.body_iterator]
    errors = [item['value'] for item in events if item['type'] == 'error']
    if errors: raise RuntimeError(str(errors))
    return events


async def main():
    save_settings({'chat_provider': 'deepseek', 'deepseek_model': 'deepseek-flash'})
    project = create_project('学习工具真实验收'); source = uid()
    image = target / 'diagram.png'; original = target / '图文原资料.docx'
    canvas = Image.new('RGB', (640, 180), 'white'); draw = ImageDraw.Draw(canvas)
    draw.rectangle((20, 30, 200, 150), fill='blue'); draw.rectangle((440, 30, 620, 150), fill='red')
    draw.text((100, 80), 'A', fill='white'); draw.text((520, 80), 'B', fill='white')
    draw.line((220, 90, 420, 90), fill='black', width=5)
    draw.polygon(((420, 90), (400, 80), (400, 100)), fill='black'); canvas.save(image)
    doc = Document(); doc.add_paragraph('箭头图用于表示从 A 到 B 的方向；仅有这个图不能证明因果关系。')
    doc.add_picture(str(image)); doc.add_paragraph('尾部条件：颜色是识别标记，不是物理量。'); doc.save(original)
    execute('INSERT INTO sources(id,project_id,name,kind,original_path,segments_json,status,created_at) VALUES (?,?,?,?,?,?,?,?)',
            (source, project['id'], original.name, '.docx', str(original), '[]', 'ready', now()))
    print('真实 DeepSeek：读取原文、视觉核对、创建并精确修改文件。', flush=True)
    first = await chat(server.ChatIn(project_id=project['id'], source_ids=[source], provider='deepseek', model='deepseek-flash',
        text='请实际调用 sf_read_source 阅读原文，再用 sf_inspect_media 查看其中一张插图。调用 sf_write_text 在工作目录建立 notes.txt，内容是“初步记录”，随后调用 sf_edit_text 把“初步记录”改成“已核对原图”。完成后描述两块颜色及箭头方向，三句话以内。不联网，不生成成果文档。'))
    conversation = next(item['value'] for item in first if item['type'] == 'conversation')
    print('继续同一对话：再次读取先前插图，检查权限与图片路径恢复。', flush=True)
    later = await chat(server.ChatIn(project_id=project['id'], source_ids=[source], provider='deepseek', model='deepseek-flash',
        conversation_id=conversation, text='继续刚才的图，请实际调用 sf_inspect_media 再核对它的箭头方向，一句话回答。不创建或修改文件，不联网。'))
    tools = many('SELECT tool,status FROM agent_tool_events WHERE conversation_id=? ORDER BY created_at,rowid', (conversation,))
    for name in ['sf_read_source', 'sf_inspect_media', 'sf_write_text', 'sf_edit_text']:
        assert any(row['tool'] == name and row['status'] == 'completed' for row in tools), (name, tools)
    assert sum(row['tool'] == 'sf_inspect_media' and row['status'] == 'completed' for row in tools) >= 2
    assert not any(row['status'] == 'failed' for row in tools), tools
    note = AgentRuntime(project['id']).workspace / 'notes.txt'
    assert note.read_text(encoding='utf-8') == '已核对原图'
    assert not codex_attempts
    first_reply = next(item['value'] for item in first if item['type'] == 'replace')
    assert '蓝' in first_reply and '红' in first_reply, first_reply
    result = {'provider': 'deepseek', 'codex_disabled': True, 'tools': tools, 'file_edited': True,
              'later_image_read': True, 'first_reply': first_reply,
              'followup': next(item['value'] for item in later if item['type'] == 'replace'), 'fixture': str(original)}
    (target / 'result.json').write_text(j(result), encoding='utf-8'); print(j(result), flush=True)


codex_attempts = []
async def forbidden(*args, **kwargs):
    codex_attempts.append(True)
    raise AssertionError('本次验证禁止调用 Codex')


with patch.object(CodexBridge, 'start', forbidden), patch.object(server.model_catalog, 'get_models', forbidden), \
     patch.object(server.model_catalog, 'validate_effort', forbidden):
    asyncio.run(main())
