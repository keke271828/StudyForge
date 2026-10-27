import json
import sqlite3
import zipfile
from pathlib import Path

from fastapi.testclient import TestClient

from studyforge.db import create_job, create_project, execute, j, one, uid, update_mastery
from studyforge.exporters import export_document
from studyforge.parsers import parse_file
from studyforge.server import app, profile_difficulty


def test_database_jobs_and_mastery():
    project = create_project("物理")
    job = create_job(project["id"], "import", {})
    assert one("SELECT status FROM jobs WHERE id=?", (job,))["status"] == "queued"
    assert update_mastery("电场", quiz=79)["score"] == 79
    assert update_mastery("电场", feynman=85)["evidence_count"] == 2
    rated = update_mastery("电场", self_rating=60, difficulty_rating=80)
    assert rated["difficulty_rating"] == 80
    assert profile_difficulty("电场", [rated]) < profile_difficulty("电场", [rated | {"difficulty_rating": 20}])


def test_first_launch_concurrent_project_requests():
    from concurrent.futures import ThreadPoolExecutor
    from threading import Barrier
    from studyforge.db import many
    barrier = Barrier(12)
    def create(index):
        barrier.wait(timeout=10)
        project = create_project(f'并发项目 {index}')
        assert many('SELECT * FROM projects WHERE id=?', (project['id'],))
        return project['id']
    with ThreadPoolExecutor(max_workers=12) as pool:
        created = list(pool.map(create, range(12)))
    assert len(set(created)) == 12
    assert len(many('SELECT * FROM projects')) == 12


def test_structured_imports(tmp_path):
    mm = tmp_path / "map.mm"
    mm.write_text('<map version="1.0.1"><node TEXT="力学"><node TEXT="牛顿定律"/></node></map>', encoding="utf-8")
    assert "力学 / 牛顿定律" in [x["locator"] for x in parse_file(mm)]
    opml = tmp_path / "map.opml"
    opml.write_text('<opml><body><outline text="物理"><outline text="动量"/></outline></body></opml>', encoding="utf-8")
    assert len(parse_file(opml)) == 2
    xmind = tmp_path / "map.xmind"
    with zipfile.ZipFile(xmind, "w") as archive:
        archive.writestr("content.json", json.dumps([{"title": "画布", "rootTopic": {"title": "电磁学", "children": {"attached": [{"title": "法拉第定律"}]}}}]))
    assert any("法拉第定律" in x["locator"] for x in parse_file(xmind))
    old_xmind = tmp_path / "old.xmind"
    with zipfile.ZipFile(old_xmind, "w") as archive:
        archive.writestr("content.xml", '<xmap-content><sheet><title>章节</title><topic><title>力学</title>'
                         '<children><topics type="attached"><topic><title>动量</title></topic></topics></children>'
                         '</topic></sheet></xmap-content>')
    assert "章节 / 力学 / 动量" in [item["locator"] for item in parse_file(old_xmind)]
    notebook = tmp_path / "demo.ipynb"
    notebook.write_text(json.dumps({"cells": [{"cell_type": "code", "source": ["print(1)\n"]}]}), encoding="utf-8")
    assert parse_file(notebook)[0]["kind"] == "code"
    model = tmp_path / "model.m"
    model.write_text("m = 1;\nF = m * a;\n", encoding="utf-8")
    assert parse_file(model)[0]["meta"]["language"] == "m"
    unknown = tmp_path / "other.xyz"
    unknown.write_text("data")
    try:
        parse_file(unknown)
    except ValueError as exc:
        assert "原文件已保留" in str(exc)
    else:
        assert False, "unsupported format was accepted"


def test_anki_import(tmp_path):
    database = tmp_path / "collection.anki2"
    db = sqlite3.connect(database)
    db.execute("CREATE TABLE notes (id integer, flds text, tags text)")
    db.execute("INSERT INTO notes VALUES (1,?,?)", ("正面\x1f背面", " physics "))
    db.commit(); db.close()
    package = tmp_path / "cards.apkg"
    with zipfile.ZipFile(package, "w") as archive:
        archive.write(database, "collection.anki2")
    cards = parse_file(package)
    assert cards[0]["meta"]["fields"] == ["正面", "背面"]


def test_export_formats_reopen(tmp_path):
    import fitz
    from docx import Document
    from pptx import Presentation
    content = "# 电磁学\n\n## 法拉第定律\n磁通量变化产生感应电动势。\n```python\nprint(2 + 2)\n```"
    for fmt in ("docx", "pdf", "pptx", "mm", "apkg", "py", "ipynb", "zip"):
        prepared = '{"cards":[{"question":"材料中什么变化产生感应电动势？","answer":"磁通量变化。","source":"材料给定模型"}]}' if fmt=='apkg' else content
        path = export_document(prepared, fmt, "电磁学", f"export.{fmt}")
        assert path.stat().st_size > (5 if fmt == "py" else 100)
        if fmt == "docx": assert "法拉第" in "\n".join(x.text for x in Document(path).paragraphs)
        if fmt == "pdf": assert len(fitz.open(path)) >= 1
        if fmt == "pptx": assert len(Presentation(path).slides) >= 2
        if fmt == "mm": assert parse_file(path)
        if fmt == "apkg": assert parse_file(path)
        if fmt == "py": assert "print" in path.read_text()


def test_answer_hidden_until_submission():
    project = create_project("测试")
    set_id = uid()
    execute("INSERT INTO exercise_sets VALUES (?,?,?,?,?,?,?,?,?)",
            (set_id, project["id"], "[]", "题组", j([{"index": 0, "question": "2+2=?", "topic": "算术"}]),
             j([{"index": 0, "solution": "4", "rubric": "答对得分"}]), 1, "assigned", "2026-01-01"))
    client = TestClient(app)
    student = client.get(f"/api/exercises/{set_id}").json()
    assert "solutions" not in student
    assert all("solution" not in q and "rubric" not in q for q in student["questions"])


def test_import_api_keeps_unsupported_file():
    project = create_project("测试")
    client = TestClient(app)
    response = client.post("/api/import/file", data={"project_id": project["id"]}, files={"file": ("unknown.xyz", b"abc")})
    assert response.status_code == 200
    source = one("SELECT * FROM sources WHERE id=?", (response.json()["source_id"],))
    assert Path(source["original_path"]).exists()


def test_submission_feedback_is_reopenable_without_answers():
    project = create_project("批改记录")
    set_id = uid()
    execute("INSERT INTO exercise_sets VALUES (?,?,?,?,?,?,?,?,?)",
            (set_id, project["id"], "[]", "练习", j([{"index": 0, "question": "题目", "topic": "力学"}]),
             j([{"index": 0, "solution": "答案", "rubric": "标准"}]), 1, "feynman_required", "2026-01-01"))
    execute("INSERT INTO submissions VALUES (?,?,?,?,?,?,?)",
            (uid(), set_id, j(["我的答案"]), j({"scores": [79], "feedback": ["单位错误"]}),
             79, j({"score": 81, "feedback": "讲解清楚"}), "2026-01-02"))
    response = TestClient(app).get(f"/api/exercises/{set_id}/submissions")
    assert response.status_code == 200
    result = response.json()[0]
    assert result["grade"]["feedback"] == ["单位错误"]
    assert result["feynman"]["score"] == 81
    assert "我的答案" not in response.text


def test_math_and_code_exports(tmp_path):
    samples = {"m": "matlab", "wl": "wolfram", "js": "javascript",
               "ts": "typescript", "cpp": "c++"}
    for fmt, language in samples.items():
        content = f"# 模型\n参数与适用条件。\n```{language}\nF = m * a;\n```"
        path = export_document(content, fmt, "模型", f"model.{fmt}")
        assert "F = m * a" in path.read_text(encoding="utf-8")
