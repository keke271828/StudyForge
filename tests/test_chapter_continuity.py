"""Source envelopes must not change logical chapter scheduling or teaching."""
import pytest

from studyforge import ai, lesson_quality as quality, source_origins as origin


def wrapped(body, name='讲解.txt'):
    return origin.header([{'name':name,'role':'ai_generated_artifact','kind':'explanation'}])+body


@pytest.mark.parametrize('name',['讲解.txt','同名｜资料.txt'])
def test_actual_artifact_batches_recognize_continuations_after_identity_registry(name):
    record={'id':'a','name':name,'original_path':'/not-read',
        'segments':[{'locator':'第1节','text':'\n'.join(f'推导单元{i:04}。' for i in range(600))}],
        'read_manifest':{'content_origin':{'role':'ai_generated_artifact','artifact_kind':'explanation'}}}
    digest,_=ai._assemble_source_digest([record],[record],False)
    parts=ai.explanation_batches(digest,2000)
    assert len(parts)>2 and all(part.startswith(origin.BEGIN) for part in parts)
    plan=quality.chapter_dependencies(parts)
    for index,part in enumerate(parts[1:],1):
        if '[原文续接]' in part:
            assert plan[index]['after']==index-1


def test_identity_filename_cannot_create_a_fake_dependency_for_independent_body():
    parts=[wrapped('\n[来源 A｜第1节]\n定义A。'),
           wrapped('\n[来源 接上文的证明材料.txt｜第1节]\n# 一个独立问题\n定义B。','接上文的证明材料.txt')]
    plan=quality.chapter_dependencies(parts)
    assert [row['after'] for row in plan]==[None,None]


def test_continuation_after_explicit_material_id_is_recognized():
    parts=['\n[来源 A｜行1]\n前部。',wrapped('\n[来源 A｜行1] [资料ID:a]\n[原文续接]\n后部。')]
    assert quality.chapter_dependencies(parts)[1]['after']==0


@pytest.mark.parametrize('raw',['[]','null','false','{}','{"version":"old","materials":[]}','not-json'])
def test_unknown_or_malformed_envelope_is_never_trimmed(raw):
    value=origin.BEGIN+raw+origin.END+'\n[来源 A｜p1]\n[原文续接]\n正文。'
    assert origin.transport(value) is None
    assert ai.explanation_batches(value)==[value]
    assert quality.chapter_dependencies(['前部。',value])[1]['after'] is None


@pytest.mark.parametrize('body,after',[
    ('[原文续接]\n后半部分。',0),('接上文的证明，继续估计。',0),
    ('## 继续上文的证明\n后半部分。',0),
    ('下面继续上文的证明。',0),('本节继续上文的证明。',0),
    ('我们不继续上文的证明，另立定义。',None),
    ('本节不接上文的证明，另立定义。',None),
    ('新概念。\n顺便比较继续上文的证明与其他方法。',None),
    ('```text\n继续上文的证明\n```',None),
    ('> 继续上文的证明\n这是引用的文字。',None),
])
def test_only_the_actual_body_opening_decides_dependency(body,after):
    values=['前部。',wrapped('\n[来源 独立资料｜p1]\n'+body)]
    assert quality.chapter_dependencies(values)[1]['after']==after


def test_versioned_transport_preserves_exact_metadata_body_and_original_string():
    body='\n[来源 讲解｜p1]\n😀正文\r\n\n末尾条件。'
    original=wrapped(body);before=original.encode()
    metadata,actual=origin.transport(original)
    assert actual==body and metadata['materials'][0]['role']=='ai_generated_artifact'
    assert original.encode()==before


def test_independent_parallel_part_does_not_claim_completed_drafts_were_supplied():
    context=quality.chapter_context({'text':'全文来源导航'},['前部原文','本部原文'],1,[],'')
    assert '已写内容衔接' not in context
    assert '已写第1部分' not in context


@pytest.mark.parametrize('kind',['dependent','independent'])
@pytest.mark.parametrize('parallel',[True,False])
def test_native_generation_orders_only_required_parts_and_exports_full_docx(tmp_path,monkeypatch,kind,parallel):
    import asyncio
    import re
    import time
    from docx import Document
    from scripts.chapter_continuity_cases import seed,split,document
    from studyforge import conversations,server
    from studyforge.db import create_job,create_project,j,many,one,unj
    from studyforge.scheduling import model_slot
    ai.save_settings({'agent_tools_enabled':False,'background_learning_enabled':False,'model_concurrency':2})
    project=create_project('章节续接验收');chat=conversations.create(project['id'],'deepseek','fixture')
    ids=seed(project['id'],kind);events=[];packets=[]
    second_entered=asyncio.Event()
    def batches(digest):
        values=split(digest);packets.extend(values);return values
    monkeypatch.setattr(server,'explanation_batches',batches)
    class Client:
        model='fixture';effort=None
        def __init__(self,runtime,conversation):self.agent_runtime,self.conversation=runtime,conversation
        def fork(self):return Client(self.agent_runtime.fork(),self.conversation)
        async def close(self):pass
        async def complete(self,system,user,**kwargs):
            if '只补写整篇文章的系统性总结' in system:return '## 系统性总结\n完整保留两部分的前提和推理。'
            review=kwargs.get('json_mode',False)
            if review:part=int(re.findall(r'DRAFT_PART_(\d+)',user)[-1])
            else:part=int(re.search(r'这是长资料第 (\d+)/',user)[1])
            previous=('部分的已完成讲解，用于续接' in user or '已写内容衔接' in user)
            async with model_slot('test',self.agent_runtime):
                events.append((part,'review' if review else 'draft','start',time.monotonic(),previous))
                if kind=='independent' and parallel and not review:
                    if part==2:second_entered.set()
                    elif part==1:await asyncio.wait_for(second_entered.wait(),3)
                await asyncio.sleep(.06 if part==1 else .01)
                events.append((part,'review' if review else 'draft','end',time.monotonic(),previous))
            return j({'passed':True,'issues':[]}) if review else document(part,previous)
    monkeypatch.setattr(server,'client_for',lambda *a,**kw:Client(kw['agent_runtime'],kw.get('conversation')))
    value=server.GenerateIn(project_id=project['id'],conversation_id=chat['id'],source_ids=ids,provider='deepseek',
        model='fixture',parallel=parallel,review_mode='fast',format='docx',exercises=False)
    job=create_job(project['id'],'generate',value.model_dump())
    asyncio.run(server._generate_job(job,value))
    row=one('SELECT * FROM jobs WHERE id=?',(job,));assert row['status']=='completed',row['error']
    assert len(packets)==2 and all(part.startswith(origin.BEGIN) for part in packets)
    plan=unj(row['result_json'])['chapter_dependencies']
    assert [entry['after'] for entry in plan]==([None,0] if kind=='dependent' else [None,None])
    def event(part,role,phase):return next(item for item in events if item[:3]==(part,role,phase))
    if kind=='dependent' or not parallel:
        assert event(2,'draft','start')[3]>=event(1,'review','end')[3]
        assert event(2,'draft','start')[4]
    else:
        assert events.index(event(2,'draft','start'))<events.index(event(1,'draft','end'))
        assert not event(2,'draft','start')[4]
    artifact=next(item for item in many("SELECT * FROM artifacts WHERE kind='explanation'")
                  if unj(item['metadata_json']).get('generation_job_id')==job)
    words='\n'.join(p.text for p in Document(artifact['path']).paragraphs)
    assert 'DRAFT_PART_1' in words and 'DRAFT_PART_2' in words and '系统性总结' in words
    assert unj(artifact['metadata_json'])['chapter_dependencies']==plan
    assert len(events)==8 and not many('SELECT * FROM exercise_sets')


def test_retry_reuses_first_draft_but_not_a_cached_independent_second_draft(tmp_path,monkeypatch):
    import asyncio
    import re
    from docx import Document
    from scripts.chapter_continuity_cases import seed,split,document
    from studyforge import conversations,server
    from studyforge.db import create_job,create_project,many,one,unj
    ai.save_settings({'agent_tools_enabled':False,'background_learning_enabled':False,'model_concurrency':2})
    project=create_project('旧缓存续写');chat=conversations.create(project['id'],'deepseek','fixture')
    ids=seed(project['id'],'dependent');calls=[];fail_summary=[True]
    class Client:
        model='fixture';effort=None
        def __init__(self,runtime,conversation):self.agent_runtime,self.conversation=runtime,conversation
        def fork(self):return Client(self.agent_runtime.fork(),self.conversation)
        async def close(self):pass
        async def complete(self,system,user,**kwargs):
            if '只补写整篇文章的系统性总结' in system:
                calls.append(('summary',False))
                if fail_summary[0]:raise RuntimeError('synthetic interruption before delivery')
                return '## 系统性总结\n保留原文条件和续写关系。'
            part=int(re.search(r'这是长资料第 (\d+)/',user)[1])
            previous='部分的已完成讲解，用于续接' in user
            calls.append((part,previous));return document(part,previous)
    monkeypatch.setattr(server,'explanation_batches',split)
    monkeypatch.setattr(server,'client_for',lambda *a,**kw:Client(kw['agent_runtime'],kw.get('conversation')))
    actual_plan=quality.chapter_dependencies
    monkeypatch.setattr(quality,'chapter_dependencies',lambda parts:[{'index':i,'after':None,'reason':''} for i in range(len(parts))])
    value=server.GenerateIn(project_id=project['id'],conversation_id=chat['id'],source_ids=ids,provider='deepseek',
        model='fixture',parallel=True,review_mode='express',format='docx',exercises=False)
    job=create_job(project['id'],'generate',value.model_dump())
    asyncio.run(server._generate_job(job,value))
    assert one('SELECT status FROM jobs WHERE id=?',(job,))['status']=='failed'
    assert calls.count((1,False))==calls.count((2,False))==1
    monkeypatch.setattr(quality,'chapter_dependencies',actual_plan);fail_summary[0]=False
    asyncio.run(server._generate_job(job,value))
    status=one('SELECT * FROM jobs WHERE id=?',(job,));assert status['status']=='completed',status['error']
    assert calls.count((1,False))==1 and calls.count((2,False))==1 and calls.count((2,True))==1
    artifact=next(row for row in many("SELECT * FROM artifacts WHERE kind='explanation'")
        if unj(row['metadata_json']).get('generation_job_id')==job)
    words='\n'.join(paragraph.text for paragraph in Document(artifact['path']).paragraphs)
    assert '本次收到前一部分讲解：是' in words.split('DRAFT_PART_2',1)[1]
    assert len(unj(status['result_json'])['chapter_dependencies'])==2
