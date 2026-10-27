import asyncio
import json
import sqlite3
import xml.etree.ElementTree as ET
import zipfile

import fitz
import pytest
from docx import Document
from docx.oxml.ns import qn
from pptx import Presentation

from studyforge import server
from studyforge.ai import generate_explanation
from studyforge.config import artifacts_dir
from studyforge.db import create_project, create_job, execute, j, many, now, one, uid
from studyforge.document_render import blocks, rich_html, EXPLANATION_COLOR, SOURCE_COLOR
from studyforge.exporters import export_document
from studyforge.teaching import finish_explanation, remove_assigned_exercises, summary_section


LESSON = ('# 规训权力\n\n'
          '> 规训权力通过日常安排影响行为。[原文.txt｜行 1-1]\n\n'
          '**讲解：**这是一种对日常行为持续塑造的权力分析。\n\n'
          '## 术语释义\n规训权力：这里指福柯的相关理论，不把所有权力关系当作规训。\n\n'
          '## 相关延伸\n补充背景应有依据，并说明与原文的联系和适用边界。\n\n'
          '## 系统性总结\n理解概念、行为塑造过程和理论边界。')


def test_prompt_teaches_source_discipline_and_explains_terms():
    calls = []
    class Client:
        async def complete(self, system, user):
            calls.append((system, user))
            return LESSON
    asyncio.run(generate_explanation(Client(), '福柯讨论规训权力。', 2,
                 [dict(topic='力学', score=95)], exercises=False))
    system, user = calls[0]
    assert '特别擅长物理' not in system and '先解释问题的物理图景' not in system
    assert all(word in system for word in ('数学', '自然科学', '社会科学', '编程', '提出者', '理论传统', '系统性总结'))
    # The named thinker belongs to this source, not every unrelated lesson's
    # policy. Keep attribution rules without hardcoding a past sample globally.
    assert '福柯' not in system and '福柯' in user
    assert '未勾选' in user and '不生成习题' in user
    assert '福柯讨论规训权力。' in user


@pytest.mark.parametrize('level', [1, 5])
def test_proof_generation_keeps_original_latest_request_and_depth_without_extra_calls(level):
    """Exercise the actual generator boundary, including the user's source limit."""
    from studyforge.teaching import REASONING_TEACHING_RULES, TEACHING_TASK_REMINDER
    from studyforge.learning_modes import mode_task_reminder
    calls = []
    original = '[来源 证明.txt｜行 1-2]\n一个固定区间内积分为零，所以被积函数处处为零。'
    latest = '只使用给定资料，先指出这一步的问题，简明解释反例，不写发现史。'
    class Client:
        async def complete(self, system, user):
            calls.append((system, user))
            return '# 证明条件\n固定区间结论不充分。\n\n## 系统性总结\n检查条件。'
    result = asyncio.run(generate_explanation(Client(), original, level, [], requirements=latest))
    assert len(calls) == 1
    system, user = calls[0]
    assert system.count(REASONING_TEACHING_RULES) == 1
    assert original in user and latest in user and '最新要求优先' in system
    assert user.endswith(TEACHING_TASK_REMINDER + mode_task_reminder('deep', role='lecture'))
    assert '未勾选' in user and '不生成习题' in user
    assert ('凝练方法选择' if level == 1 else '关键构造的启发') in user
    assert result.endswith('检查条件。')


def test_exercise_removal_preserves_source_quotes_code_and_final_summary():
    document = ('# 讲解\n> ## 思考题\n> 原文作者提出的问题。\n\n'
                '```python\n# 练习\nprint(1)\n```\n\n'
                '## 课后练习\n待作答问题。\n### 习题\n嵌套待作答问题。\n'
                '## 系统性总结\n原文的论证链条。')
    cleaned = remove_assigned_exercises(document)
    assert '> ## 思考题' in cleaned and '# 练习\nprint(1)' in cleaned
    assert '待作答问题' not in cleaned
    assert cleaned.endswith('原文的论证链条。')


def test_complete_summary_is_last_without_extra_model_calls():
    class Client:
        async def complete(self, *_):
            raise AssertionError('完整单篇讲解不应多调用模型')
    document = LESSON + '\n\n## 待核实事项\n一个背景细节尚待确认。'
    result = asyncio.run(finish_explanation(Client(), document))
    assert result.index('待核实事项') < result.index('系统性总结')
    assert result.endswith('理解概念、行为塑造过程和理论边界。')


@pytest.mark.parametrize('heading', ['**系统性总结**', '__系统性总结__', '系统性总结', '## 5）系统性总结'])
def test_alternative_summary_titles_do_not_generate_duplicate_summaries(heading):
    class Client:
        async def complete(self, *_):
            raise AssertionError('已有完整总结时不应重复调用模型')
    document = '# 原文讲解\n讲清原文。\n\n' + heading + '\n全文主旨与边界。'
    result = asyncio.run(finish_explanation(Client(), document))
    assert result.count('系统性总结') == 1
    assert result.endswith('全文主旨与边界。')


@pytest.mark.parametrize('heading', ['## 小结', '# 全文小结', '**小结**'])
def test_real_closing_recap_reuses_text_without_another_summary_call(heading):
    class Client:
        async def complete(self, *_):
            raise AssertionError('已有结尾小结时不应重复生成更长的总结')
    result = asyncio.run(finish_explanation(Client(), '# 配方法\n先解释构造。\n\n'+heading+'\n关键条件：B>0；零向量另行处理。'))
    assert result.count('系统性总结')==1
    assert result.endswith('关键条件：B>0；零向量另行处理。')


def test_local_recap_and_quoted_recap_do_not_replace_missing_final_summary():
    calls = []
    class Client:
        async def complete(self, system, user):
            calls.append(user)
            assert '只总结本次用户消息提供的正文' in system
            return '## 系统性总结\n保留零向量边界。'
    document = '# 配方法\n## 小结\n局部结论。\n## 零向量\n另行处理。\n> ## 小结\n> 原资料结尾。'
    result = asyncio.run(finish_explanation(Client(), document))
    assert len(calls)==1 and result.endswith('保留零向量边界。')


def test_bold_exercise_title_is_removed_without_touching_quoted_summary():
    document = '# 原文讲解\n> **系统性总结**\n> 原作者的表述。\n\n**课后练习**\n请作答。\n\n**系统性总结**\n全文结论。'
    cleaned = remove_assigned_exercises(document)
    assert '请作答' not in cleaned
    assert '> **系统性总结**' in cleaned and cleaned.endswith('全文结论。')


def test_missing_summary_and_multi_part_integration():
    calls = []
    class Client:
        async def complete(self, system, user):
            calls.append(user)
            return '## 系统性总结\n串联甲、乙两个部分。'
    result = asyncio.run(finish_explanation(Client(), '# 原文讲解\n核心概念甲。'))
    assert result.endswith('串联甲、乙两个部分。')
    chapters = ['# 甲\n## 系统性总结\n核心概念甲。', '# 乙\n## 系统性总结\n核心概念乙。']
    result = asyncio.run(finish_explanation(Client(), '\n\n'.join(chapters), chapters=chapters))
    assert '核心概念甲' in calls[-1] and '核心概念乙' in calls[-1]
    assert result.endswith('串联甲、乙两个部分。')


def test_summary_receives_complete_written_chapters_and_explicit_whole_scope():
    class Client:
        async def complete(self, system, user):
            assert '不仅总结最后一部分' in user
            assert '已完成第 1/2 部分讲解' in user and '已完成第 2/2 部分讲解' in user
            assert '正文重要条件甲' in user and '正文重要条件乙' in user
            return '## 系统性总结\n甲与乙共同构成完整主线。'
    chapters = ['# 甲\n正文重要条件甲。\n## 系统性总结\n过短的小结。',
                '# 乙\n正文重要条件乙。\n## 系统性总结\n过短的小结。']
    result = asyncio.run(finish_explanation(Client(), '\n\n'.join(chapters), chapters=chapters))
    assert result.endswith('甲与乙共同构成完整主线。')


@pytest.mark.parametrize('title', ['## 系统性总结\n', '# 系统性总结\n', ''])
def test_generated_summary_keeps_all_subsections_and_nests_their_headings(title):
    class Client:
        async def complete(self, system, user):
            return title + ('整篇概览。\n\n## 基础结构\ntable 与 tabular。\n\n'
                            '## 方法与边界\nbooktabs 与 multicolumn。\n\n'
                            '```text\n# 代码原样保留\n```\n\n> ## 引用标题原样保留\n')
    result = asyncio.run(finish_explanation(Client(), '# 正文\n已经完成的讲解。'))
    assert '### 基础结构' in result and '### 方法与边界' in result
    assert 'booktabs 与 multicolumn' in result and '# 代码原样保留' in result
    assert '> ## 引用标题原样保留' in result
    start, end = summary_section(result)
    assert end == len(result) and result[start:].startswith('## 系统性总结')


def make_source(tmp_path):
    project = create_project('社会科学')
    source = tmp_path / '原文.txt'
    source.write_text('福柯讨论规训权力通过日常安排影响行为。', encoding='utf-8')
    source_id = uid()
    execute('INSERT INTO sources(id,project_id,name,kind,original_path,segments_json,status,error,created_at) VALUES (?,?,?,?,?,?,?,?,?)',
            (source_id, project['id'], source.name, '.txt', str(source), '[]', 'ready', None, now()))
    return project, source_id


@pytest.mark.parametrize('exercises,fmt', [(False, 'docx'), (True, 'pdf')])
def test_generation_delivers_only_selected_files_and_separate_opt_in_exercises(tmp_path, monkeypatch, exercises, fmt):
    project, source_id = make_source(tmp_path)
    exercise_calls = []
    class Client:
        model = 'test-model'
        effort = 'high'
        async def complete(self, system, user):
            assert '原资料全文' in user
            return LESSON + '\n\n## 课后练习\n请回答一个不应放在讲解中的问题。'
    async def generate_set(*args, **kwargs):
        exercise_calls.append(True)
        set_id = uid()
        execute('INSERT INTO exercise_sets VALUES (?,?,?,?,?,?,?,?,?)',
                (set_id, project['id'], j([source_id]), '配套习题', j([dict(index=0, question='解释规训权力。', topic='规训权力')]),
                 j([dict(index=0, solution='私有答案', rubric='评分标准')]), 1, 'assigned', now()))
        return set_id, [dict(stage='独立解题', result={'solution': '私有答案'})]
    monkeypatch.setattr(server, 'client_for', lambda *args, **kwargs: Client())
    monkeypatch.setattr(server, 'generate_set', generate_set)
    value = server.GenerateIn(project_id=project['id'], source_ids=[source_id], format=fmt, exercises=exercises)
    job_id = create_job(project['id'], 'generate', value.model_dump())
    asyncio.run(server._generate_job(job_id, value))
    job = one('SELECT * FROM jobs WHERE id=?', (job_id,))
    assert job['status'] == 'completed', job['error']
    artifacts = many('SELECT * FROM artifacts ORDER BY kind')
    assert len(artifacts) == (2 if exercises else 1)
    assert [a['format'] for a in artifacts if a['kind'] == 'explanation'] == [fmt]
    assert [a['format'] for a in artifacts if a['kind'] == 'exercise'] == (['docx'] if exercises else [])
    assert len(exercise_calls) == int(exercises)
    assert {path.suffix for path in artifacts_dir().iterdir()} == ({'.pdf', '.docx'} if exercises else {'.docx'})
    assert '私有答案' not in job['result_json']
    assert '请回答一个不应' not in (''.join(p.text for p in Document(artifacts[0]['path']).paragraphs)
                                  if fmt == 'docx' else ''.join(page.get_text() for page in fitz.open(next(a['path'] for a in artifacts if a['kind'] == 'explanation'))))


def test_failed_export_does_not_deliver_an_unrequested_markdown_file(tmp_path, monkeypatch):
    project, source_id = make_source(tmp_path)
    class Client:
        async def complete(self, *_): return LESSON
    def broken_export(*args): raise ValueError('排版失败')
    monkeypatch.setattr(server, 'client_for', lambda *args, **kwargs: Client())
    monkeypatch.setattr(server, 'export_document', broken_export)
    value = server.GenerateIn(project_id=project['id'], source_ids=[source_id])
    job = create_job(project['id'], 'generate', value.model_dump())
    asyncio.run(server._generate_job(job, value))
    assert one('SELECT status FROM jobs WHERE id=?', (job,))['status'] == 'failed'
    assert not many('SELECT * FROM artifacts')
    assert not list(artifacts_dir().iterdir())


def test_quotes_have_distinct_word_and_pdf_typefaces_and_colours():
    content = '> 原文引用清晰可读。\n\nAI讲解清晰可读。'
    word = Document(export_document(content, 'docx', '排版验收', 'styled.docx'))
    quote = next(p for p in word.paragraphs if p.text.startswith('原文引用'))
    body = next(p for p in word.paragraphs if p.text.startswith('AI讲解'))
    assert quote.style.font.color.rgb != body.style.font.color.rgb
    assert str(quote.style.font.color.rgb) == SOURCE_COLOR
    assert str(body.style.font.color.rgb) == EXPLANATION_COLOR
    assert quote.style.element.get_or_add_rPr().get_or_add_rFonts().get(qn('w:eastAsia')) == 'SimSun'
    assert body.style.font.name == 'Microsoft YaHei'
    with fitz.open(export_document(content, 'pdf', '排版验收', 'styled.pdf')) as pdf:
        spans = [span for page in pdf for block in page.get_text('dict')['blocks'] if 'lines' in block
                 for line in block['lines'] for span in line['spans']]
        quote = next(span for span in spans if '原文引用' in span['text'])
        body = next(span for span in spans if 'AI讲解' in span['text'])
        assert quote['font'] != body['font']
        assert quote['color'] == int(SOURCE_COLOR, 16)
        assert body['color'] == int(EXPLANATION_COLOR, 16)


def test_quote_styling_survives_other_rich_formats(tmp_path):
    content = LESSON + '\n\n```python\nprint(1)\n```'
    ppt = Presentation(export_document(content, 'pptx', '排版验收', 'styled.pptx'))
    paragraphs = [p for slide in ppt.slides for shape in slide.shapes if shape.has_text_frame for p in shape.text_frame.paragraphs]
    quoted = next(p for p in paragraphs if p.text.startswith('规训权力通过'))
    assert str(quoted.font.color.rgb) == SOURCE_COLOR
    assert quoted.font.name == 'Times New Roman'
    mindmap = ET.parse(export_document(content, 'mm', '排版验收', 'styled.mm'))
    quoted_node = next(n for n in mindmap.iter('node') if n.get('TEXT', '').startswith('规训权力通过'))
    assert quoted_node.get('COLOR') == '#' + SOURCE_COLOR
    assert quoted_node.find('font').get('NAME') == 'SimSun'
    notebook = json.loads(export_document(content, 'ipynb', '排版验收', 'styled.ipynb').read_text(encoding='utf-8'))
    assert '<blockquote' in ''.join(notebook['cells'][0]['source'])
    card_content = json.dumps({'cards':[{'question':'材料中规训权力通过哪些日常安排影响行为？','answer':LESSON,'source':'原文.txt｜行 1-1'}]},ensure_ascii=False)
    package = export_document(card_content, 'apkg', '排版验收', 'styled.apkg')
    with zipfile.ZipFile(package) as archive:
        archive.extract('collection.anki2', tmp_path)
    with sqlite3.connect(tmp_path / 'collection.anki2') as connection:
        answers = '\n'.join(row[0] for row in connection.execute('SELECT flds FROM notes'))
        assert '<blockquote' in answers and '#' + SOURCE_COLOR in answers


def test_quote_parser_preserves_source_lines_and_safe_html():
    parsed = list(blocks('> 第一行原文\n> 第二行原文\n\n解释。'))
    assert parsed[0] == ('quote', '第一行原文\n第二行原文')
    assert parsed[1] == ('paragraph', '解释。')
    output = rich_html('> <script>alert(1)</script>\n\n<img onerror="bad()">')
    assert '<script>' not in output and '<img onerror=' not in output
    assert '&lt;script&gt;' in output
