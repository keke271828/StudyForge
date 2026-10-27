"""Real UI/network races on isolated data; no cloud requests or file approvals executed."""
import asyncio
from pathlib import Path

from docx import Document
from scripts import smoke_library_ui as fixture
from studyforge import conversations,lesson_feedback
from studyforge.config import artifacts_dir
from studyforge.mobile_context import remote_device
from studyforge.db import create_job,execute,j,now,one,uid,update_job

fixture.enable_semantic(False)
second=conversations.create(fixture.p['id'],'deepseek')
conversations.activate(fixture.c['id'])
ids=[]
original=fixture.case_path.parent/'完整原文.txt';original.write_text('符号要与含义对应，推导要检查条件。',encoding='utf-8')
for number in range(40):
    ident=uid();ids.append(ident)
    execute('INSERT INTO sources(id,project_id,name,kind,original_path,segments_json,status,created_at) VALUES (?,?,?,?,?,?,?,?)',
        (ident,fixture.p['id'],f'学习资料 {number+1:02}.txt','.txt',str(original),
         j([{'locator':'原文位置','text':'LARGE_CACHED_ORIGINAL '*12000}]),'ready',now()))
doc_ids=[];questionnaires=[]
for title in ['Alpha 概念讲解','Zulu 方法讲解']:
    path=artifacts_dir()/(uid()+'.docx');document=Document();document.add_heading(title,0);document.add_paragraph('本机界面验收文档；不宣称模型教学质量。');document.save(path)
    job=create_job(fixture.p['id'],'generate',{'conversation_id':fixture.c['id'],'source_ids':[]})
    artifact=fixture.server.add_artifact(fixture.p['id'],[],'explanation','docx',title,path,{'generation_job_id':job,'conversation_id':fixture.c['id']})
    reply=uid();execute('INSERT INTO messages VALUES (?,?,?,?,?)',(reply,fixture.c['id'],'assistant','已生成讲解文档：\n\n# '+title+'\n概念与条件。',now()))
    execute('INSERT INTO generation_dialogue_outputs VALUES (?,?,?)',(job,fixture.c['id'],reply))
    update_job(job,'completed',100,{'artifacts':[artifact]})
    doc_ids.append(artifact['id']);questionnaires.append(lesson_feedback.ensure_questionnaire(job)['id'])
worksheet=uid()
execute('INSERT INTO exercise_sets VALUES (?,?,?,?,?,?,?,?,?)',(worksheet,fixture.p['id'],'[]','变量理解练习',
    j([{'index':0,'topic':'概念','question':'用自己的话说明符号与量的区别。'}]),
    j([{'index':0,'solution':'HIDDEN_REFERENCE_BEFORE_SUBMISSION','rubric':'准确区分符号与量'}]),2,'assigned',now()))
tasks=[]
for label in ['任务甲','任务乙']:
    ident=create_job(fixture.p['id'],'generate',{'project_id':fixture.p['id'],'conversation_id':fixture.c['id'],
        'source_ids':[ids[0]],'dialogue_snapshot':{'messages':['PRIVATE_TASK_CONTEXT '*100000]}})
    update_job(ident,'paused',35,{'stage':label});tasks.append(ident)
control={'sources':0,'mobile_sources':0,'artifacts':0,'mobile_artifacts':0,'pauses':0,'approval_gets':0,'mobile_approval_gets':0,'approval_posts':0,
    'hold_sources':False,'hold_sort':False,'hold_approval_get':False,'hold_approval_post':False,'hold_pause':False}
gates={name:asyncio.Event() for name in ['sources','sort','approval_get','approval_post','pause']}
approval_id=uid()


def replace(path,method,endpoint):
    routes=fixture.server.app.router.routes
    index=next(i for i,route in enumerate(routes) if getattr(route,'path',None)==path and method in getattr(route,'methods',set()))
    fixture.server.app.add_api_route(path,endpoint,methods=[method]);route=routes.pop();routes[index]=route


async def sources(project_id:str):
    result=fixture.server.sources(project_id);control['sources']+=1
    mobile=bool(remote_device.get());control['mobile_sources']+=int(mobile)
    if mobile and control['hold_sources'] and project_id==fixture.p['id']:await gates['sources'].wait()
    return result


async def artifacts(project_id:str,sort:str='newest'):
    result=fixture.server.artifacts(project_id,sort);control['artifacts']+=1
    mobile=bool(remote_device.get());control['mobile_artifacts']+=int(mobile)
    if mobile and control['hold_sort'] and sort=='newest' and project_id==fixture.p['id']:await gates['sort'].wait()
    return result


async def pause(job_id:str):
    control['pauses']+=1
    if control['hold_pause']:await gates['pause'].wait()
    return await fixture.server.pause_generation(job_id)


async def get_approvals(project_id:str,conversation_id:str):
    result=fixture.server.agent_pending_approvals(project_id,conversation_id);control['approval_gets']+=1
    mobile=bool(remote_device.get());control['mobile_approval_gets']+=int(mobile)
    if mobile and control['hold_approval_get'] and conversation_id==fixture.c['id']:await gates['approval_get'].wait()
    return result


async def post_approval(request_id:str,value:fixture.server.ApprovalDecision):
    control['approval_posts']+=1
    if control['hold_approval_post']:await gates['approval_post'].wait()
    return fixture.server.agent_approval_decision(request_id,value)


for path,method,endpoint in [('/api/projects/{project_id}/sources','GET',sources),('/api/projects/{project_id}/artifacts','GET',artifacts),
    ('/api/jobs/{job_id}/pause','POST',pause),('/api/agent/approvals','GET',get_approvals),
    ('/api/agent/approvals/{request_id}','POST',post_approval)]:replace(path,method,endpoint)


@fixture.server.app.get('/__workspace_test__/state')
def state():return control|{'approval':(one('SELECT status FROM agent_approvals WHERE id=?',(approval_id,)) or {}).get('status')}


@fixture.server.app.post('/__workspace_test__/hold/{name}')
def hold(name:str):
    key='hold_'+name
    if key not in control:raise ValueError(name)
    gates[name].clear();control[key]=True
    return state()


@fixture.server.app.post('/__workspace_test__/release/{name}')
def release(name:str):control['hold_'+name]=False;gates[name].set();return state()


@fixture.server.app.post('/__workspace_test__/arm')
def arm():
    update_job(tasks[0],'running',41,{'stage':'任务甲 · 原阶段'})
    details={'paths':[{'path':str(original),'operation':'读取'}],'preview':'用于审批等待与归属验收，不执行文件操作。'}
    if one('SELECT id FROM agent_approvals WHERE id=?',(approval_id,)):
        execute("UPDATE agent_approvals SET status='pending',resolved_at=NULL WHERE id=?",(approval_id,))
    else:execute('INSERT INTO agent_approvals VALUES (?,?,?,?,?,?,?,?,?,?,?)',
        (approval_id,fixture.p['id'],fixture.c['id'],tasks[0],'sf_read_file',uid(),j(details),'pending',None,now(),None))
    return state()


@fixture.server.app.post('/__workspace_test__/progress')
def progress():update_job(tasks[0],'running',55,{'stage':'任务甲 · 新阶段'});return state()


@fixture.server.app.post('/__workspace_test__/feedback')
def feedback():
    execute("UPDATE lesson_questionnaires SET reply_status='completed',updated_at=? WHERE id=?",(now(),questionnaires[0]))
    return state()


routes=fixture.server.app.router.routes
mount=next(route for route in routes if getattr(route,'name',None)=='web');routes.remove(mount);routes.append(mount)
fixture.case.update(second_conversation_id=second['id'],source_ids=ids,artifact_ids=doc_ids,worksheet_id=worksheet,
    task_ids=tasks,approval_id=approval_id,upload_path=str(original))
if __name__=='__main__':asyncio.run(fixture.main())
