import asyncio
import base64
import io
import json
from pathlib import Path

import httpx
import pytest
from fastapi.testclient import TestClient
from PIL import Image

from studyforge import agent_tools, server
from studyforge.ai import DeepSeekClient, save_settings
from studyforge.agent_tools import AgentRuntime
from studyforge.db import create_project, execute, j, many, now, one, uid


@pytest.fixture(autouse=True)
def isolate_application_root(tmp_path,monkeypatch):
    from studyforge import agent_permissions
    monkeypatch.setattr(agent_permissions,'application_root',lambda:tmp_path/'application')


def run(runtime, tool, args):
    return asyncio.run(runtime.execute(tool, args))


def test_file_scope_exact_edit_and_backup(tmp_path):
    project = create_project('文件测试'); runtime = AgentRuntime(project['id'])
    outside = tmp_path / 'outside.txt'; outside.write_text('secret', encoding='utf-8')
    assert 'error' in run(runtime, 'sf_read_file', {'path': str(outside)})
    assert 'error' in run(runtime, 'sf_write_text', {'path': '../escape.txt', 'text': 'bad'})
    write = run(runtime, 'sf_write_text', {'path': 'lesson.py', 'text': 'x = 1\n'})
    assert 'error' not in write
    edit = run(runtime, 'sf_edit_text', {'path': 'lesson.py', 'old': '1', 'new': '2'})
    assert 'error' not in edit
    from pathlib import Path
    assert Path(edit['path']).read_text(encoding='utf-8') == 'x = 2\n'
    assert Path(edit['backup']).read_text(encoding='utf-8') == 'x = 1\n'
    save_settings({'agent_write_roots': [str(tmp_path)]})
    assert 'error' not in run(runtime, 'sf_edit_text', {'path': str(outside), 'old': 'secret', 'new': 'changed'})
    save_settings({'agent_file_write_enabled': False})
    assert 'error' in run(runtime, 'sf_write_text', {'path': 'blocked.txt', 'text': 'no'})
    assert len(many('SELECT * FROM agent_tool_events')) == 6


def test_original_source_paginated_not_preview(tmp_path):
    project = create_project('全文'); source_id = uid(); path = tmp_path / 'original.txt'
    path.write_text('完整原文。' * 150 + '尾部关键论据', encoding='utf-8')
    execute('INSERT INTO sources(id,project_id,name,kind,original_path,segments_json,status,created_at) VALUES (?,?,?,?,?,?,?,?)',
            (source_id, project['id'], path.name, '.txt', str(path), j([{'text':'错误预览'}]), 'ready', now()))
    runtime = AgentRuntime(project['id'], [source_id]); text, offset = '', 0
    while True:
        result = run(runtime, 'sf_read_source', {'source_id': source_id, 'offset': offset, 'limit': 200})
        assert 'error' not in result; text += result['text']
        if result['next_offset'] is None: break
        offset = result['next_offset']
    assert '尾部关键论据' in text and '错误预览' not in text and '[来源' in text
    assert 'error' in run(runtime, 'sf_read_source', {'source_id': uid()})


def test_video_permission_blocks_before_downloader_and_revokes_cache(tmp_path, monkeypatch):
    calls = []
    path = tmp_path / 'video.mp4'; path.write_bytes(b'video')
    def download(url, key, **kwargs): calls.append(kwargs); return path
    monkeypatch.setattr(agent_tools, 'download_public_resource', download)
    monkeypatch.setattr(agent_tools, 'validate_public_url', lambda url: url)
    monkeypatch.setattr(agent_tools, 'parse_file', lambda path: [{'locator':'00:00', 'text':'视频字幕', 'meta':{}}])
    runtime = AgentRuntime(); url = 'https://www.bilibili.com/video/BV1xx411c7mD'
    assert 'error' in run(runtime, 'sf_read_url', {'url': url}) and calls == []
    save_settings({'agent_web_video_enabled': True})
    result = run(runtime, 'sf_read_url', {'url': url})
    assert result['method'] == 'video_transcript_and_sampled_frames' and '视频字幕' in result['text']
    assert calls == [{'allow_video': True}]
    save_settings({'agent_web_video_enabled': False})
    assert 'error' in run(runtime, 'sf_read_url', {'url': url}) and len(calls) == 1


def test_network_and_tool_budget_are_enforced_live(monkeypatch):
    runtime = AgentRuntime(); save_settings({'agent_network_enabled': False, 'agent_max_tool_calls': 1})
    assert 'error' in run(runtime, 'sf_search_web', {'query':'因果'})
    assert '额度' in run(runtime, 'sf_list_files', {})['error']
    save_settings({'agent_tools_enabled': False})
    assert '关闭' in run(runtime, 'sf_list_files', {})['error']


def test_actual_image_bytes_saved_and_only_owned_preview_served(monkeypatch):
    project = create_project('图片'); runtime = AgentRuntime(project['id'])
    monkeypatch.setattr(agent_tools, 'get_secret', lambda key: None)
    assert '尚未配置' in run(runtime, 'sf_generate_image', {'prompt':'负反馈图'})['error']
    monkeypatch.setattr(agent_tools, 'get_secret', lambda key: 'test-key')
    buffer = io.BytesIO(); Image.new('RGB', (20, 20), 'blue').save(buffer, 'PNG')
    original = httpx.AsyncClient
    transport = httpx.MockTransport(lambda request: httpx.Response(200, json={'data':[{'b64_json':base64.b64encode(buffer.getvalue()).decode()}]}))
    monkeypatch.setattr(httpx, 'AsyncClient', lambda *a, **kw: original(transport=transport))
    result = run(runtime, 'sf_generate_image', {'prompt':'负反馈图'})
    assert 'error' not in result and result['artifact']['format'] == 'png'
    response = TestClient(server.app).get(result['image_url'])
    assert response.status_code == 200 and response.content == buffer.getvalue()
    assert TestClient(server.app).get('/api/agent/assets/' + uid() + '/image').status_code == 404


def test_deepseek_real_protocol_tool_loop_and_private_reasoning(monkeypatch):
    monkeypatch.setattr('studyforge.ai.get_secret', lambda key: 'test-key')
    requests = []; original = httpx.AsyncClient
    def response(request):
        body = json.loads(request.content); requests.append(body)
        delta = {'reasoning_content':'private reasoning', 'tool_calls':[{'index':0, 'id':'call-1', 'type':'function',
                 'function':{'name':'sf_write_text', 'arguments':j({'path':'demo.py','text':'print(42)\n'})}}]} if len(requests) == 1 else {'content':'文件已创建。'}
        packet = {'choices':[{'delta':delta, 'finish_reason':'tool_calls' if len(requests)==1 else 'stop'}]}
        return httpx.Response(200, text='data: '+j(packet)+'\n\ndata: {"choices":[],"usage":{"total_tokens":1}}\n\ndata: [DONE]\n\n', headers={'content-type':'text/event-stream'})
    monkeypatch.setattr(httpx, 'AsyncClient', lambda *a, **kw: original(transport=httpx.MockTransport(response)))
    runtime = AgentRuntime(); client = DeepSeekClient('test', agent_runtime=runtime)
    async def collect(): return [event async for event in client.agent_stream([{'role':'system','content':'导师'}, {'role':'user','content':'写示例'}])]
    events = asyncio.run(collect())
    assert events[-1] == ('final', '文件已创建。')
    assert len(requests) == 2 and requests[1]['messages'][-1]['role'] == 'tool'
    assert requests[1]['messages'][-2]['reasoning_content'] == 'private reasoning'
    assert not any('private reasoning' in str(value) for kind, value in events)
    assert (runtime.workspace/'demo.py').read_text(encoding='utf-8') == 'print(42)\n'


def test_generation_tool_cannot_export_unrequested_documents():
    runtime = AgentRuntime(publishing=False)
    assert '统一导出' in run(runtime, 'sf_create_document', {'content':'正文'})['error']


def source_for(project, path, segments=None):
    source_id = uid()
    execute('INSERT INTO sources(id,project_id,name,kind,original_path,segments_json,status,created_at) VALUES (?,?,?,?,?,?,?,?)',
            (source_id, project['id'], path.name, path.suffix, str(path), j(segments or []), 'ready', now()))
    return source_id


def test_document_images_are_readable_by_path_and_locator_and_stale_images_revoked(tmp_path, monkeypatch):
    from docx import Document
    project = create_project('图文资料'); image = tmp_path / 'blue.png'
    Image.new('RGB', (20, 20), 'blue').save(image)
    original = tmp_path / 'article.docx'; doc = Document(); doc.add_paragraph('图片中的信息需要视觉核查。'); doc.add_picture(str(image)); doc.save(original)
    source_id = source_for(project, original); runtime = AgentRuntime(project['id'], [source_id])
    calls = []
    async def vision(path, question): calls.append(path); return '图片为蓝色。'
    monkeypatch.setattr(runtime, 'vision', vision)
    read = run(runtime, 'sf_read_source', {'source_id': source_id})
    frame = read['media_frames'][0]
    direct = run(runtime, 'sf_inspect_media', {'path': frame['path']})
    by_locator = run(runtime, 'sf_inspect_media', {'path': str(original), 'locator': frame['locator']})
    assert direct['source'] == str(original) and direct['locator'] == frame['locator']
    assert by_locator['description'] == '图片为蓝色。' and by_locator['method'] == 'document_image'
    doc.add_paragraph('此文档已更新。'); doc.save(original)
    assert '原资料已变化' in run(runtime, 'sf_inspect_media', {'path': frame['path']})['error']
    fresh = run(runtime, 'sf_read_source', {})['media_frames'][0]
    assert fresh['path'] != frame['path']
    runtime.source_ids.clear()
    assert 'error' in run(runtime, 'sf_inspect_media', {'path': fresh['path']})
    assert len(calls) == 2


def test_same_named_originals_never_overwrite_each_others_images(tmp_path):
    from docx import Document
    from studyforge.parsers import parse_file
    parts = []
    for directory, color in [('first', 'red'), ('second', 'blue')]:
        folder = tmp_path / directory; folder.mkdir(); image = folder / 'picture.png'
        Image.new('RGB', (20, 20), color).save(image)
        original = folder / 'same-name.docx'; doc = Document(); doc.add_picture(str(image)); doc.save(original)
        parts.append(parse_file(original)[0]['meta']['image_path'])
    assert parts[0] != parts[1]
    with Image.open(parts[0]) as first, Image.open(parts[1]) as second:
        assert first.getpixel((0, 0)) == (255, 0, 0) and second.getpixel((0, 0)) == (0, 0, 255)


def test_web_video_permission_covers_frame_and_cached_file_routes(monkeypatch):
    from studyforge.config import temp_dir
    path = temp_dir() / 'video.mp4'; path.write_bytes(b'video')
    image = temp_dir() / 'frame.png'; Image.new('RGB', (20, 20), 'blue').save(image)
    monkeypatch.setattr(agent_tools, 'download_public_resource', lambda *a, **kw: path)
    monkeypatch.setattr(agent_tools, 'validate_public_url', lambda url: url)
    monkeypatch.setattr(agent_tools, 'parse_file', lambda path: [{'locator': '约 10 秒画面', 'text': '板书', 'meta': {'image_path': str(image), 'time': 10}}])
    runtime = AgentRuntime(); calls = []
    async def vision(path, question): calls.append(path); return '视频板书。'
    monkeypatch.setattr(runtime, 'vision', vision)
    save_settings({'agent_web_video_enabled': True})
    read = run(runtime, 'sf_read_url', {'url': 'https://www.bilibili.com/video/BV1xx411c7mD'})
    assert run(runtime, 'sf_inspect_media', {'path': read['media_frames'][0]['path']})['frame_time_seconds'] == 10
    save_settings({'agent_web_video_enabled': False})
    frame=read['media_frames'][0]['path'];original=read['original_path']
    for name, target in [('sf_inspect_media', frame), ('sf_inspect_media', original), ('sf_read_file', original), ('sf_read_file', frame)]:
        assert '网页视频权限已关闭' in run(runtime, name, {'path': str(target)})['error']
    assert len(calls) == 1


def test_untrusted_extracted_image_path_cannot_read_unrelated_files(tmp_path, monkeypatch):
    project = create_project('图像范围'); original = tmp_path / 'source.txt'; original.write_text('资料', encoding='utf-8')
    other = tmp_path / 'other.png'; Image.new('RGB', (20, 20), 'blue').save(other)
    source = source_for(project, original); runtime = AgentRuntime(project['id'], [source])
    monkeypatch.setattr(agent_tools, 'parse_file', lambda path: [{'locator':'图 1', 'text':'', 'meta': {'image_path': str(other)}}])
    assert '尚未授权' in run(runtime, 'sf_inspect_media', {'path': str(original)})['error']


def test_diagram_listing_does_not_silently_omit_images_after_first_page(tmp_path, monkeypatch):
    original = tmp_path / 'large.txt'; original.write_text('正文', encoding='utf-8')
    project = create_project('长图文'); source = source_for(project, original)
    monkeypatch.setattr('studyforge.materials.parse_file', lambda path: [{'locator': f'第 {i} 页', 'text': '内容', 'meta': {'image_path': f'page-{i}.png'}} for i in range(121)])
    runtime = AgentRuntime(project['id'], [source])
    first = run(runtime, 'sf_read_source', {})
    last = run(runtime, 'sf_read_source', {'media_offset': first['next_media_offset']})
    assert first['media_frame_count'] == 121 and len(first['media_frames']) == 120
    assert last['media_frames'][0]['locator'] == '第 120 页' and last['next_media_offset'] is None


def test_owned_generated_images_remain_readable_in_later_turns(monkeypatch):
    project = create_project('图片回看'); foreign = create_project('其他项目')
    conversation=uid()
    execute('INSERT INTO conversations VALUES (?,?,?,?,?,?,?,?)',(conversation,project['id'],'deepseek',None,now(),'[]','test',None))
    runtime = AgentRuntime(project['id'],conversation_id=conversation); monkeypatch.setattr(agent_tools, 'get_secret', lambda key: 'test-key')
    buffer = io.BytesIO(); Image.new('RGB', (20, 20), 'blue').save(buffer, 'PNG')
    original_client = httpx.AsyncClient
    monkeypatch.setattr(httpx, 'AsyncClient', lambda *a, **kw: original_client(transport=httpx.MockTransport(
        lambda request: httpx.Response(200, json={'data': [{'b64_json': base64.b64encode(buffer.getvalue()).decode()}]}))))
    image = run(runtime, 'sf_generate_image', {'prompt':'负反馈图'})
    later = AgentRuntime(project['id'],conversation_id=conversation)
    async def vision(path, question): return '可以回看生成图。'
    monkeypatch.setattr(later, 'vision', vision)
    assert run(later, 'sf_inspect_media', {'path': image['path']})['description'] == '可以回看生成图。'
    assert json.loads(later.context())['generated_images'][0]['path'] == image['path']
    assert '尚未授权' in run(AgentRuntime(foreign['id']), 'sf_inspect_media', {'path': image['path']})['error']
    assert '尚未授权' in run(later, 'sf_edit_text', {'path': image['path'], 'old':'x', 'new':'y'})['error']


@pytest.mark.parametrize('suffix', ['.bmp', '.tiff'])
def test_supported_original_image_formats_use_vision_directly(tmp_path, monkeypatch, suffix):
    image = tmp_path / ('picture' + suffix); Image.new('RGB', (20, 20), 'blue').save(image)
    save_settings({'agent_read_roots': [str(tmp_path)]})
    runtime = AgentRuntime()
    async def vision(path, question): return '直接读取图像。'
    monkeypatch.setattr(runtime, 'vision', vision)
    assert run(runtime, 'sf_inspect_media', {'path': str(image)})['method'] == 'vision'


def test_original_video_frame_inspection_reuses_completed_transcript(monkeypatch, tmp_path):
    from studyforge.config import temp_dir
    project = create_project('视频资料'); video = tmp_path / 'video.mp4'; video.write_bytes(b'original-video')
    image = temp_dir() / 'frame.png'; Image.new('RGB', (20, 20), 'blue').save(image)
    source = source_for(project, video, [{'locator':'0-10 秒', 'text':'完整字幕', 'meta':{}},
                                        {'locator':'约 10 秒画面', 'text':'板书', 'meta':{'time':10, 'image_path':str(image)}}])
    runtime = AgentRuntime(project['id'], [source])
    def forbidden(path): raise AssertionError('已经有完整转录，不应重复解码整个视频')
    monkeypatch.setattr(agent_tools, 'parse_file', forbidden)
    async def vision(path, question): return '板书图形。'
    monkeypatch.setattr(runtime, 'vision', vision)
    assert '完整字幕' in run(runtime, 'sf_read_source', {})['text']
    result = run(runtime, 'sf_inspect_media', {'path':str(video), 'time_seconds':10})
    assert result['description'] == '板书图形。' and result['frame_time_seconds'] == 10


def test_document_media_scope_survives_later_turn_but_does_not_restore_revoked_roots(tmp_path, monkeypatch):
    from docx import Document
    project = create_project('多轮图像'); image = tmp_path / 'blue.png'; Image.new('RGB', (20, 20), 'blue').save(image)
    original = tmp_path / 'article.docx'; doc = Document(); doc.add_picture(str(image)); doc.save(original)
    source = source_for(project, original); conversation = uid()
    execute('INSERT INTO conversations VALUES (?,?,?,?,?,?,?,?)', (conversation, project['id'], 'deepseek', None, now(), j([source]), 'test', None))
    runtime = AgentRuntime(project['id'], [source], conversation_id=conversation)
    frame = run(runtime, 'sf_read_source', {})['media_frames'][0]
    later = AgentRuntime(project['id'], [source], conversation_id=conversation)
    async def vision(path, question): return '后续讨论继续核对同一插图。'
    monkeypatch.setattr(later, 'vision', vision)
    assert 'description' in run(later, 'sf_inspect_media', {'path':frame['path']})
    revoked = AgentRuntime(project['id'], [], conversation_id=conversation)
    assert '尚未授权' in run(revoked, 'sf_inspect_media', {'path':frame['path']})['error']


def test_web_video_revoked_during_download_stops_before_transcription(monkeypatch):
    from studyforge.config import temp_dir
    video = temp_dir() / 'web-video.mp4'; video.write_bytes(b'video')
    def download(*a, **kw):
        save_settings({'agent_web_video_enabled': False}); return video
    def forbidden(path): raise AssertionError('权限撤销后不得开始转录或抽帧')
    monkeypatch.setattr(agent_tools, 'download_public_resource', download)
    monkeypatch.setattr(agent_tools, 'parse_file', forbidden)
    monkeypatch.setattr(agent_tools, 'validate_public_url', lambda url: url)
    save_settings({'agent_web_video_enabled': True})
    result = run(AgentRuntime(), 'sf_read_url', {'url':'https://www.bilibili.com/video/BV1xx411c7mD'})
    assert '网页视频权限已关闭' in result['error']


def test_read_root_revoked_during_vision_does_not_publish_image_description(tmp_path, monkeypatch):
    image = tmp_path / 'image.png'; Image.new('RGB', (20, 20), 'blue').save(image)
    save_settings({'agent_read_roots': [str(tmp_path)]}); runtime = AgentRuntime()
    async def vision(path, question):
        save_settings({'agent_read_roots': []}); return '不再获准读取的图像描述'
    monkeypatch.setattr(runtime, 'vision', vision)
    result = run(runtime, 'sf_inspect_media', {'path': str(image)})
    assert '尚未授权' in result['error'] and 'description' not in result


def test_image_service_connection_failure_has_readable_error(monkeypatch):
    monkeypatch.setattr(agent_tools, 'get_secret', lambda key: 'test-key')
    original = httpx.AsyncClient
    def fail(request): raise httpx.ConnectError('', request=request)
    monkeypatch.setattr(httpx, 'AsyncClient', lambda *a, **kw: original(transport=httpx.MockTransport(fail)))
    result = run(AgentRuntime(), 'sf_generate_image', {'prompt': '学习示意图'})
    assert '网络连接失败' in result['error'] and result['type'] == 'ConnectError'


@pytest.mark.parametrize('suffix,text', [('.ipynb', '{"cells":[{"cell_type":"code","source":["x=1"]}],"metadata":{},"nbformat":4,"nbformat_minor":0}'),
                                       ('.cs', 'class Lesson { public const int Value = 1; }'),
                                       ('.sql', 'SELECT 1 AS value;'), ('.ps1', '$value = 1')])
def test_agent_can_write_and_edit_supported_source_languages_and_notebooks(suffix, text):
    runtime = AgentRuntime()
    created = run(runtime, 'sf_write_text', {'path':'lesson'+suffix, 'text':text})
    assert 'error' not in created
    edited = run(runtime, 'sf_edit_text', {'path':created['path'], 'old':'1', 'new':'2'})
    assert 'error' not in edited and Path(edited['backup']).read_text(encoding='utf-8') == text
    read = run(runtime, 'sf_read_file', {'path':created['path']})
    assert 'error' not in read and '2' in read['text']
