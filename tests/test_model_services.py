import asyncio
import json
from pathlib import Path

import httpx
import pytest
from docx import Document
from fastapi.testclient import TestClient

from studyforge import ai, conversations, server, providers
from studyforge.agent_permissions import decide, list_requests
from studyforge.agent_tools import AgentRuntime
from studyforge.codex_bridge import CodexBridge
from studyforge.config import data_dir
from studyforge.db import create_project, create_job, execute, j, many, now, one, uid
from studyforge.provider_clients import GeminiClient, ServiceClient, gemini_body, service_client


@pytest.fixture
def credentials(monkeypatch):
    saved = {}
    monkeypatch.setattr(providers, 'get_secret', saved.get)
    monkeypatch.setattr(providers, 'set_secret', lambda key, value: saved.__setitem__(key, value))
    monkeypatch.setattr(providers, 'delete_secret', lambda key: saved.pop(key, None))
    monkeypatch.setattr('studyforge.provider_clients.get_secret', saved.get)
    return saved


def service(protocol='openai', **extra):
    return providers.save_service({'name': '测试 '+protocol, 'protocol': protocol, 'base_url': 'https://fixture.example/v1',
                                  'default_model': 'selected-model', 'models': ['second-model'], 'api_key': 'private-key', **extra})


def transport(monkeypatch, handler):
    original = httpx.AsyncClient
    monkeypatch.setattr(httpx, 'AsyncClient', lambda *a, **kw: original(*a, **(kw | {'transport': httpx.MockTransport(handler)})))


def reply(protocol, text='', calls=None, finish='STOP', streaming=True):
    if protocol == 'gemini':
        parts = ([{'text': text}] if text else []) + (calls or [])
        body = {'candidates': [{'content': {'role':'model', 'parts': parts}, 'finishReason': finish}]}
    else:
        body = {'choices': [{'delta': {'content': text, 'tool_calls': calls or []},
                             'finish_reason': 'tool_calls' if calls else 'stop' if finish == 'STOP' else finish}]}
    if not streaming:
        if protocol!='gemini':body['choices'][0]['message']=body['choices'][0].pop('delta')
        return httpx.Response(200, json=body)
    return httpx.Response(200, text='data: '+j(body)+'\n\n', headers={'content-type': 'text/event-stream'})


def tool_call(protocol, name, args):
    if protocol == 'gemini': return [{'thoughtSignature': 'opaque-signature', 'functionCall': {'id':'native-call','name': name, 'args': args}}]
    return [{'index':0, 'id':'call', 'type':'function', 'function': {'name': name, 'arguments': j(args)}}]


def cli_pipes(child):
    """Expose real pipe semantics around a scripted CLI response."""
    child.stdout, child.stderr = asyncio.StreamReader(), asyncio.StreamReader()
    class Input:
        def __init__(self): self.data = bytearray(); self.task = None
        def write(self, data): self.data.extend(data)
        async def drain(self): pass
        def close(self):
            async def finish():
                out, err = await child.communicate(bytes(self.data))
                child.stdout.feed_data(out); child.stdout.feed_eof()
                child.stderr.feed_data(err); child.stderr.feed_eof()
            self.task = asyncio.create_task(finish())
    child.stdin = Input()
    original_kill = child.kill
    def kill():
        original_kill()
        if child.stdin.task and not child.stdin.task.done(): child.stdin.task.cancel()
    child.kill = kill
    return child


def test_service_crud_never_serializes_secrets_and_preferences_are_per_service(credentials):
    row = service(); other = service('gemini')
    assert row['id'] != other['id'] and row['configured']
    assert 'private-key' not in (data_dir()/'model-services.json').read_text()
    with TestClient(server.app) as api:
        assert 'private-key' not in api.get('/api/providers').text
        assert 'private-key' not in api.get('/api/settings').text
        response = api.put('/api/providers/'+row['id'], json={'name':'重命名服务'})
        assert response.status_code == 200 and response.json()['configured']
        assert credentials[providers.secret_name(row['id'])] == 'private-key'
        response = api.put('/api/settings', json={'chat_provider':other['id'], 'provider_model_preferences':{row['id']:'second-model'}})
        assert response.status_code == 200
        response = api.put('/api/settings', json={'provider_model_preferences':{other['id']:'my-gemini'}})
        assert response.json()['provider_model_preferences'] == {row['id']:'second-model',other['id']:'my-gemini'}
        assert api.put('/api/providers/'+row['id'], json={'clear_key':True}).json()['configured'] is False
        assert api.delete('/api/providers/'+other['id']).status_code == 200
        assert providers.secret_name(other['id']) not in credentials
        assert ai.settings()['chat_provider'] == 'deepseek'


@pytest.mark.parametrize('change', [
    {'name':42},{'name':''},{'protocol':'unknown'},{'base_url':'http://cloud.example/v1'},
    {'base_url':'https://user:pass@example.com/v1'},{'base_url':'https://example.com/v1?key=secret'},
    {'base_url':'https://example.com/v1/chat/completions'},{'models':'one-model'},
    {'supports_tools':'yes'},{'default_model':'model\nname'},{'api_key':'key\r\nAuthorization: bad'},
    {'clear_key':True,'api_key':'another'},{'headers':{'Authorization':'bad'}}])
def test_invalid_configuration_is_rejected_before_writing(credentials, change):
    with pytest.raises((ValueError, TypeError)): service(**change)
    assert not (data_dir()/'model-services.json').exists()
    assert not credentials


def test_local_openai_compatible_service_can_omit_key(credentials):
    row = service(base_url='http://127.0.0.1:11434/v1', requires_key=False, api_key='')
    assert row['configured'] and isinstance(service_client(row['id']), ServiceClient)


@pytest.mark.parametrize('protocol', ['openai','gemini'])
def test_catalog_reads_service_endpoint_and_merges_manual_models(credentials, monkeypatch, protocol):
    row = service(protocol); requests = []
    def handler(request):
        requests.append(request)
        if protocol == 'gemini':
            return httpx.Response(200,json={'models':[{'name':'models/live-gemini','displayName':'在线 Gemini','supportedGenerationMethods':['generateContent']},
                                                     {'name':'models/embedding','supportedGenerationMethods':['embedContent']}]})
        return httpx.Response(200,json={'data':[{'id':'online-model'}]})
    transport(monkeypatch, handler)
    result = asyncio.run(providers.model_catalog(row['id']))
    assert result['source']=='api' and result['listed_count']==1
    assert len(result['models'])==3 and not any(m['id']=='embedding' for m in result['models'])
    request = requests[0]
    assert request.url.path.endswith('/models') and 'private-key' not in str(request.url)
    assert request.headers['x-goog-api-key' if protocol=='gemini' else 'Authorization'] == ('private-key' if protocol=='gemini' else 'Bearer private-key')


def test_catalog_failure_preserves_manual_models_and_does_not_leak_response(credentials, monkeypatch):
    row = service()
    transport(monkeypatch, lambda r:httpx.Response(401,json={'message':'private-key credential rejected'}))
    result=asyncio.run(providers.model_catalog(row['id']))
    assert result['source']=='manual' and result['warning']
    assert 'private-key' not in j(result) and len(result['models'])==2


def test_native_gemini_retains_all_thought_signatures_and_parallel_function_ids():
    native={'role':'model','parts':[{'text':'公开内容','thoughtSignature':'text-signature'},
        {'functionCall':{'id':'native-one','name':'first','args':{}},'thoughtSignature':'signature-one'},
        {'functionCall':{'id':'native-two','name':'second','args':{}},'thoughtSignature':'signature-two'}]}
    body=gemini_body({'messages':[{'role':'system','content':'教学要求'}, {'role':'user','content':'问题'},
        {'role':'assistant','content':'公开内容','_gemini_content':native,'tool_calls':[
            {'id':'one','function':{'name':'first','arguments':'{}'}},{'id':'two','function':{'name':'second','arguments':'{}'}}]},
        {'role':'tool','tool_call_id':'one','content':'{"text":"result-one"}'},
        {'role':'tool','tool_call_id':'two','content':'{"text":"result-two"}'}]})
    assert body['contents'][1] == native
    assert len(body['contents'][2]['parts'])==2
    assert body['contents'][2]['parts'][0]['functionResponse']['id']=='native-one'
    assert body['contents'][2]['parts'][1]['functionResponse']['id']=='native-two'


@pytest.mark.parametrize('protocol', ['openai','gemini'])
@pytest.mark.parametrize('approved', [True, False])
def test_real_chat_suspends_for_approval_then_continues_using_selected_service(credentials, monkeypatch, tmp_path, protocol, approved):
    row=service(protocol); ai.save_settings({'chat_provider':row['id'],'background_learning_enabled':False,'agent_permission_mode':'approval'})
    project=create_project('第三方审批'); chat=conversations.create(project['id'],row['id'],'selected-model')
    path=tmp_path/'target.txt';path.write_text('before',encoding='utf-8');requests=[]
    def handler(request):
        body=json.loads(request.content);requests.append(body)
        if len(requests)==1:return reply(protocol,'等待你审批。',tool_call(protocol,'sf_write_text',{'path':str(path),'text':'after','overwrite':True}))
        if protocol=='gemini':
            saved=body['contents'][-2]['parts'][-1]
            assert saved['thoughtSignature']=='opaque-signature'
            assert body['contents'][-1]['parts'][0]['functionResponse']['id']=='native-call'
        return reply(protocol,'操作已完成。' if approved else '你拒绝了此次操作。')
    transport(monkeypatch,handler)
    async def run():
        response=await server.chat_stream(server.ChatIn(project_id=project['id'],conversation_id=chat['id'],provider=row['id'],source_ids=[],text='请修改指定文件'))
        events=[]
        async def consume():
            async for event in response.body_iterator:events.append(json.loads(event))
        task=asyncio.create_task(consume())
        for _ in range(150):
            pending=list_requests(project['id'],chat['id'])
            if pending:break
            await asyncio.sleep(.01)
        assert pending
        await asyncio.sleep(.05)
        assert not task.done() and path.read_text()=='before' and len(requests)==1
        assert any(e['type']=='tool' and e['value']['status']=='waiting_approval' for e in events)
        decide(pending[0]['id'],project['id'],chat['id'],approved)
        await asyncio.wait_for(task,4)
        assert not any(e['type']=='error' for e in events) and events[-1]['type']=='done'
    asyncio.run(run())
    assert path.read_text()==('after' if approved else 'before') and len(requests)==2
    assert one('SELECT protocol_version FROM agent_dialogue_state WHERE conversation_id=?',(chat['id'],))['protocol_version']==(3 if protocol=='gemini' else 2)


@pytest.mark.parametrize('protocol', ['openai','gemini'])
def test_read_chat_and_actual_docx_generation_never_try_codex(credentials, monkeypatch, tmp_path, protocol):
    row=service(protocol,supports_tools=False)
    ai.save_settings({'chat_provider':row['id'],'background_learning_enabled':False})
    async def forbidden(*a,**kw):raise AssertionError('第三方服务不得启动 Codex')
    monkeypatch.setattr(CodexBridge,'start',forbidden)
    requests=[]
    def handler(request):
        body=json.loads(request.content);requests.append(body)
        assert request.url.host=='fixture.example'
        if protocol=='openai': assert body['model'] in ('selected-model','second-model')
        return reply(protocol,'# 相关与因果\n共同变化不足以证明因果，必须排除共同原因。\n\n## 系统性总结\n实验还需要考虑外部效度。')
    transport(monkeypatch,handler)
    project=create_project('第三方完整学习流程'); sid=uid()
    original=tmp_path/'original.txt';original.write_text('相关不等于因果。尾部论据：实验也有外部效度限制。',encoding='utf-8')
    execute('INSERT INTO sources(id,project_id,name,kind,original_path,segments_json,status,created_at) VALUES (?,?,?,?,?,?,?,?)',
            (sid,project['id'],original.name,'.txt',str(original),j([{'text':'错误的界面摘要','locator':'摘要'}]),'ready',now()))
    async def run():
        read=server.ReadIn(project_id=project['id'],source_ids=[sid]);await server.choose_generation_model(read)
        job=create_job(project['id'],'read',read.model_dump());await server._read_job(job,read)
        assert one('SELECT * FROM jobs WHERE id=?',(job,))['status']=='completed'
        response=await server.chat_stream(server.ChatIn(project_id=project['id'],source_ids=[sid],text='解释共同原因',model='second-model'))
        events=[json.loads(event) async for event in response.body_iterator]
        assert not any(e['type']=='error' for e in events)
        cid=next(e['value'] for e in events if e['type']=='conversation')
        value=server.GenerateIn(project_id=project['id'],source_ids=[sid],conversation_id=cid,model='old-writer')
        started=await server.generate(value)
        for _ in range(400):
            job=one('SELECT * FROM jobs WHERE id=?',(started['job_id'],))
            if job['status'] in ('completed','failed'):break
            await asyncio.sleep(.01)
        assert job['status']=='completed',job['error']
    asyncio.run(run())
    artifact=one("SELECT * FROM artifacts WHERE kind='explanation'")
    text='\n'.join(p.text for p in Document(artifact['path']).paragraphs)
    assert '共同原因' in text
    assert json.loads(artifact['metadata_json'])['dialogue_provider']==row['id']
    assert len(requests)==3 and all('尾部论据' in j(body) for body in requests)
    assert all('错误的界面摘要' not in j(body) for body in requests)


@pytest.mark.parametrize('protocol', ['openai','gemini'])
def test_json_exercises_are_reviewed_and_exported_without_answers(credentials, monkeypatch, protocol):
    from studyforge.exercises import generate_set, student_document
    from studyforge.exporters import export_document
    import studyforge.exercises as exercises
    row=service(protocol,supports_tools=False)
    async def no_search(*a,**kw):return []
    monkeypatch.setattr(exercises,'search_exercises',no_search)
    questions=[{'question':f'题{i+1}：相关为什么不足以证明因果？','solution':'私有解析：可能有共同原因','rubric':'指出共同原因','topic':'因果推理','source_url':''} for i in range(4)]
    sequence=iter([{'questions':questions},{'answers':[{'index':i,'solution':'共同原因','issues':[]} for i in range(4)]},
                   {'passed':True,'issues':[]},{'passed':True,'issues':[]}])
    bodies=[]
    def handler(request):
        body=json.loads(request.content);bodies.append(body)
        if protocol=='gemini':assert body['generationConfig']['responseMimeType']=='application/json'
        else:assert body['response_format']=={'type':'json_object'}
        return reply(protocol,j(next(sequence)),streaming=False)
    transport(monkeypatch,handler)
    project=create_project('第三方习题')
    client=service_client(row['id']);sid,reports=asyncio.run(generate_set(client,project['id'],[],'相关不等于因果',1))
    path=export_document(student_document(sid),'docx','题单','model-exercises.docx')
    text='\n'.join(p.text for p in Document(path).paragraphs)
    assert '题1' in text and '私有解析' not in text and len(bodies)==4
    assert reports and one('SELECT status FROM exercise_sets WHERE id=?',(sid,))['status']=='assigned'


def test_disabled_removed_and_unknown_services_cannot_launch_models_but_history_is_readable(credentials):
    row=service();project=create_project('服务生命周期');chat=conversations.create(project['id'],row['id'],'selected-model')
    with TestClient(server.app) as api:
        assert api.put('/api/providers/'+row['id'],json={'enabled':False}).status_code==200
        assert api.post('/api/chat/stream',json={'project_id':project['id'],'provider':row['id'],'text':'测试'}).status_code==400
        assert api.delete('/api/providers/'+row['id']).status_code==200
        assert api.get(f"/api/conversations/{chat['id']}?project_id={project['id']}").status_code==200
        assert api.post(f"/api/projects/{project['id']}/conversations",json={'provider':'unknown'}).status_code==400


def test_service_with_active_job_cannot_change_endpoint_or_delete_credentials(credentials):
    row=service();project=create_project('保存任务配置');create_job(project['id'],'generate',{'provider':row['id']})
    with TestClient(server.app) as api:
        assert api.put('/api/providers/'+row['id'],json={'base_url':'https://other.example/v1'}).status_code==409
        assert api.delete('/api/providers/'+row['id']).status_code==409
    assert providers.secret_name(row['id']) in credentials


def test_native_gemini_image_body_is_inline_and_truncated_output_is_not_exported(credentials, monkeypatch, tmp_path):
    from PIL import Image
    row=service('gemini');requests=[]
    def handler(request):
        body=json.loads(request.content);requests.append(body)
        return httpx.Response(200,json={'candidates':[{'content':{'parts':[{'text':'只输出一半'}]},'finishReason':'MAX_TOKENS'}]})
    transport(monkeypatch,handler)
    image=tmp_path/'source.png';Image.new('RGB',(20,20)).save(image)
    with pytest.raises(RuntimeError,match='长度上限'):asyncio.run(service_client(row['id']).analyze_image(image,'原图'))
    assert requests[0]['contents'][0]['parts'][1]['inlineData']['mimeType']=='image/jpeg'


def test_gemini_cli_uses_dedicated_auth_profile_no_native_tools_or_inherited_keys(credentials, monkeypatch):
    from studyforge.gemini_cli import prepare_profile, status, cli_packet
    row=service('gemini_cli')
    monkeypatch.setenv('GEMINI_API_KEY','foreign-secret')
    root,env=prepare_profile(row['id'])
    assert 'GEMINI_API_KEY' not in env
    assert str(root).startswith(str(data_dir()))
    config=json.loads((root/'.gemini'/'settings.json').read_text())
    assert config['tools']['core']==['__studyforge_no_native_tools__'] and config['mcp']['allowed']==[]
    assert 'decision = "deny"' in (root/'deny-native-tools.toml').read_text()
    assert config['security']['auth']['selectedType']=='oauth-personal'
    assert not status(row['id'])['configured']
    parsed=cli_packet(j({'content':'有效正文','tool_calls':[{'name':'sf_read_learning_profile','arguments':{}}]}))
    assert parsed['choices'][0]['finish_reason']=='tool_calls'
    with pytest.raises(RuntimeError):cli_packet('{"content":""')


def test_gemini_cli_missing_installation_is_an_actionable_error_not_empty_content(credentials, monkeypatch):
    from studyforge import gemini_cli
    row=service('gemini_cli');monkeypatch.setattr(gemini_cli,'cli_command',lambda:None)
    with pytest.raises(RuntimeError,match='安装'):asyncio.run(service_client(row['id']).complete('规则','提问'))


@pytest.mark.parametrize('protocol',['openai','gemini'])
def test_json_and_tool_capability_switches_change_real_requests(credentials,monkeypatch,protocol):
    row=service(protocol,supports_json=False,supports_tools=False);seen=[]
    def handler(request):
        body=json.loads(request.content);seen.append(body)
        assert 'tools' not in body
        if protocol=='gemini':assert 'responseMimeType' not in body['generationConfig']
        else:assert 'response_format' not in body
        return reply(protocol,'{"answer":"有效内容"}')
    transport(monkeypatch,handler)
    result=asyncio.run(ai.client_for(row['id']).complete('要求','输入',json_mode=True))
    assert json.loads(result)['answer']=='有效内容' and len(seen)==1


def test_native_gemini_second_chat_turn_reattaches_only_its_own_signed_parts(credentials,monkeypatch):
    row=service('gemini');pid=create_project('上下文隔离')['id'];cid=conversations.create(pid,row['id'],'selected-model')['id']
    other=conversations.create(pid,row['id'],'selected-model')['id'];execute('INSERT INTO messages VALUES (?,?,?,?,?)',(uid(),other,'user','OTHER_CHAT_RAW_SECRET',now()))
    requests=[]
    def handler(request):
        body=json.loads(request.content);requests.append(body)
        if len(requests)==1:
            return httpx.Response(200,json={'candidates':[{'content':{'role':'model','parts':[{'text':'首轮讲解','thoughtSignature':'signed-final-text'}]},'finishReason':'STOP'}]})
        assert any(part.get('thoughtSignature')=='signed-final-text' for content in body['contents'] for part in content['parts'])
        assert 'OTHER_CHAT_RAW_SECRET' not in j(body)
        return reply('gemini','后续讲解')
    transport(monkeypatch,handler)
    async def run():
        for text in ['第一轮知识问题','继续讲解原知识']:
            response=await server.chat_stream(server.ChatIn(project_id=pid,conversation_id=cid,provider=row['id'],text=text))
            events=[json.loads(line) async for line in response.body_iterator]
            assert not any(event['type']=='error' for event in events)
    asyncio.run(run());assert len(requests)==2


def test_cli_account_executes_shared_tools_after_approval_and_keeps_long_prompt_off_command_line(credentials,monkeypatch,tmp_path):
    from studyforge import gemini_cli
    row=service('gemini_cli');root,env=gemini_cli.prepare_profile(row['id'])
    (root/'.gemini'/'oauth_creds.json').write_text('{}')
    monkeypatch.setattr(gemini_cli,'cli_command',lambda:['fixture-node','fixture-cli'])
    target=tmp_path/'file.txt';target.write_text('before',encoding='utf-8')
    ai.save_settings({'chat_provider':row['id'],'background_learning_enabled':False,'agent_permission_mode':'approval'})
    pid=create_project('Gemini账号审批')['id'];cid=conversations.create(pid,row['id'],'selected-model')['id'];calls=[]
    class Child:
        returncode=None
        async def communicate(self,prompt):
            packet=json.loads(prompt.decode().split('任务输入：\n',1)[1]);calls.append(packet)
            finished=packet['messages'][-1]['role']=='tool'
            result={'content':'已保存' if finished else '等待审批','tool_calls':[] if finished else [{'name':'sf_write_text','arguments':{'path':str(target),'text':'after','overwrite':True}}]}
            self.returncode=0
            return (j({'type':'message','role':'assistant','content':j(result)})+'\n'+j({'type':'result','status':'success'})+'\n').encode(),b''
        async def wait(self):return self.returncode
        def kill(self):self.returncode=-1
    async def spawn(*args,**kwargs):
        assert all('LONG_SOURCE_SENTINEL' not in arg for arg in args)
        assert kwargs['env']['GEMINI_FORCE_ENCRYPTED_FILE_STORAGE']=='false'
        assert kwargs['cwd'].startswith(str(root/'runs'))
        assert '--approval-mode' in args and 'plan' in args
        return cli_pipes(Child())
    monkeypatch.setattr(asyncio,'create_subprocess_exec',spawn)
    async def run():
        response=await server.chat_stream(server.ChatIn(project_id=pid,conversation_id=cid,provider=row['id'],text='LONG_SOURCE_SENTINEL 修改文件'))
        events=[]
        async def consume():
            async for data in response.body_iterator:events.append(json.loads(data))
        task=asyncio.create_task(consume())
        for _ in range(150):
            pending=list_requests(pid,cid)
            if pending:break
            await asyncio.sleep(.01)
        assert pending and not task.done() and target.read_text()=='before'
        decide(pending[0]['id'],pid,cid,True);await asyncio.wait_for(task,3)
        assert not any(e['type']=='error' for e in events)
    asyncio.run(run());assert len(calls)==2 and target.read_text()=='after'
    assert one('SELECT protocol_version FROM agent_dialogue_state WHERE conversation_id=?',(cid,))['protocol_version']==4


def test_cli_model_protocol_timeout_kills_only_its_child(credentials,monkeypatch):
    from studyforge import gemini_cli
    row=service('gemini_cli');root,_=gemini_cli.prepare_profile(row['id']);(root/'.gemini'/'oauth_creds.json').write_text('{}')
    monkeypatch.setattr(gemini_cli,'cli_command',lambda:['fixture-cli']);killed=[]
    class Child:
        returncode=None
        async def communicate(self,data):await asyncio.Event().wait()
        def kill(self):self.returncode=-1;killed.append(True)
        async def wait(self):return self.returncode
    async def spawn(*a,**kw):return cli_pipes(Child())
    monkeypatch.setattr(asyncio,'create_subprocess_exec',spawn)
    with pytest.raises(RuntimeError,match='超时'):
        asyncio.run(service_client(row['id']).run_cli({'messages':[]},timeout=.01))
    assert killed==[True]


def test_cli_login_and_logout_are_explicit_and_do_not_delete_student_history(credentials,monkeypatch):
    from studyforge import gemini_cli
    row=service('gemini_cli');pid=create_project('账号管理')['id'];cid=conversations.create(pid,row['id'])['id']
    monkeypatch.setattr(gemini_cli,'cli_command',lambda:['fixture-cli'])
    launched=[]
    class Process:
        def poll(self):return 0
    monkeypatch.setattr(gemini_cli.subprocess,'Popen',lambda *a,**kw:launched.append((a,kw)) or Process())
    with TestClient(server.app) as api:
        assert api.post('/api/providers/'+row['id']+'/login').status_code==200
        profile=gemini_cli.profile_dir(row['id']);(profile/'.gemini'/'oauth_creds.json').write_text('{}')
        assert providers.get_service(row['id']) and gemini_cli.status(row['id'])['configured']
        assert api.post('/api/providers/'+row['id']+'/logout').status_code==200
        assert not profile.exists() and one('SELECT id FROM conversations WHERE id=?',(cid,))
    assert len(launched)==1


def test_official_cli_entry_is_read_from_package_bin_instead_of_old_hardcoded_path(credentials,monkeypatch):
    from studyforge import gemini_cli
    package=data_dir()/'gemini-cli-runtime'/'node_modules'/'@google'/'gemini-cli'
    (package/'bundle').mkdir(parents=True);(package/'bundle'/'gemini.js').write_text('//fixture')
    (package/'package.json').write_text(j({'bin':{'gemini':'bundle/gemini.js'}}))
    monkeypatch.setattr(gemini_cli.shutil,'which',lambda name:'fixture-node' if name=='node' else None)
    assert gemini_cli.cli_command()==['fixture-node',str(package/'bundle'/'gemini.js')]
    (package/'package.json').write_text(j({'bin':{'gemini':'../../outside.js'}}))
    assert gemini_cli.cli_command() is None


def test_connection_detection_blocks_configuration_changes_until_it_finishes(credentials,monkeypatch):
    from fastapi import HTTPException
    from studyforge import provider_routes
    row=service();started=asyncio.Event();finish=asyncio.Event()
    class Client:
        model='selected-model'
        async def complete(self,*a,**kw):started.set();await finish.wait();return '连接正常'
        async def close(self):pass
    monkeypatch.setattr('studyforge.provider_clients.service_client',lambda *a,**kw:Client())
    async def run():
        task=asyncio.create_task(provider_routes.test(row['id'],{}));await started.wait()
        with pytest.raises(HTTPException) as result:provider_routes.inactive(row['id'])
        assert result.value.status_code==409
        with pytest.raises(HTTPException) as duplicate:await provider_routes.test(row['id'],{})
        assert duplicate.value.status_code==409
        finish.set();assert (await task)['ok']
        provider_routes.inactive(row['id'])
    asyncio.run(run())
