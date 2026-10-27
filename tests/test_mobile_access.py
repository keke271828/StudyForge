import asyncio
import json
import secrets
import time
from datetime import datetime, timedelta, timezone
from pathlib import Path

from fastapi.testclient import TestClient
from starlette.responses import StreamingResponse
import pytest

from studyforge import server, conversations
from studyforge.db import create_project, execute, many, one, unj, now
from studyforge.mobile_access import MobileManager, COOKIE, hash_token, allowed, private_address
from studyforge.mobile_context import remote_device


def mobile_client():
    manager = MobileManager(server.app)
    manager.addresses = ['127.0.0.1']
    client = TestClient(manager.gateway, base_url='https://127.0.0.1', client=('127.0.0.1', 15000))
    return manager, client


def pair(manager, client):
    code = manager.new_code()
    response = client.post('/mobile/pair', json={'code': code, 'name': 'Android验收手机'})
    assert response.status_code == 200, response.text
    request = response.json()
    poll = '/mobile/pair/' + request['device_id']
    headers = {'X-Mobile-Pair-Ticket': request['ticket']}
    assert client.get(poll, headers=headers).json()['status'] == 'pending'
    assert client.get('/api/projects').status_code == 401
    manager.decide(request['device_id'], True)
    response = client.get(poll, headers=headers)
    assert response.json()['status'] == 'active', response.text
    assert 'Secure' in response.headers['set-cookie'] and 'HttpOnly' in response.headers['set-cookie']
    assert 'SameSite=strict' in response.headers['set-cookie']
    row = one('SELECT * FROM mobile_devices WHERE id=?', (request['device_id'],))
    secret = client.cookies.get(COOKIE)
    assert row['token_hash'] == hash_token(secret) and secret not in json.dumps(row)
    assert client.get(poll, headers=headers).status_code == 403
    return request['device_id']


def test_phone_needs_confirmation_and_has_no_desktop_settings_access():
    manager, client = mobile_client()
    with client:
        ident = pair(manager, client)
        assert client.get('/').status_code == 200
        for path in ('/api/settings', '/api/desktop/status', '/api/mobile/start', '/api/mobile/certificate', '/server.py', '/app.js/../server.py'):
            response = client.get(path)
            assert response.status_code in (403, 404), (path, response.status_code)
        assert client.get('/api/projects', headers={'Origin': 'https://evil.invalid'}).status_code == 403
        assert client.post('/api/projects', json={'name': 'x'}, headers={'Sec-Fetch-Site': 'cross-site'}).status_code == 403
        assert client.get('/api/projects', headers={'Host': 'evil.invalid'}).status_code == 403
        assert client.get('/api/projects').status_code == 200
        project = create_project('手机创建对话不抢电脑导航')
        desktop_chat = conversations.create(project['id'], 'deepseek')
        current = conversations.workspace()
        created = client.post(f"/api/projects/{project['id']}/conversations", json={'provider': 'deepseek'})
        assert created.status_code == 200, created.text
        assert conversations.workspace() == current
        assert one('SELECT conversation_id FROM project_workspaces WHERE project_id=?', (project['id'],))['conversation_id'] == desktop_chat['id']
        assert client.post('/mobile/logout').status_code == 200
        assert one('SELECT status FROM mobile_devices WHERE id=?', (ident,))['status'] == 'revoked'
        assert client.get('/api/projects').status_code == 401


def test_gateway_streams_upload_and_download_without_parsing():
    from docx import Document
    from studyforge.exporters import export_document
    manager, client = mobile_client()
    with client:
        pair(manager, client)
        response = client.post('/api/projects', json={'name': '手机项目'})
        assert response.status_code == 200, response.text
        project = response.json()
        upload = client.post('/api/import/file', data={'project_id': project['id']}, files={'file': ('phone.txt', '手机上传的原文'.encode())})
        assert upload.status_code == 200, upload.text
        assert upload.json()['job_id'] is None
        row = one('SELECT * FROM sources WHERE id=?', (upload.json()['source_id'],))
        assert row['status'] == 'stored' and Path(row['original_path']).read_text(encoding='utf-8') == '手机上传的原文'
        selected = client.post('/api/sources/' + row['id'] + '/parse', json={'mode': 'text', 'selected_text': '只读这句原文', 'locator': '用户选段'})
        assert selected.status_code == 200, selected.text
        path = export_document('# 手机验收成果\n\n下载的讲解正文', 'docx', '手机验收成果', 'phone.docx')
        # Match the actual artifact schema without invoking a paid model.
        from studyforge.db import uid, j
        artifact = uid()
        execute('INSERT INTO artifacts(id,project_id,source_ids_json,kind,format,title,path,created_at,metadata_json) VALUES (?,?,?,?,?,?,?,?,?)',
                (artifact, project['id'], '[]', 'explanation', 'docx', '手机验收成果', str(path), now(), '{}'))
        response = client.get('/api/artifacts/' + artifact + '/download')
        assert response.status_code == 200, response.text
        assert response.content == path.read_bytes()
        assert 'filename' in response.headers['content-disposition']


def test_mobile_chat_is_durable_idempotent_and_scoped_to_device(monkeypatch):
    manager, client = mobile_client()
    async def fake_chat(request):
        assert remote_device.get()
        async def stream():
            for value in ('回复第一段', '；回复第二段'):
                yield json.dumps({'type': 'delta', 'value': value}) + '\n'
                await asyncio.sleep(.1)
            yield json.dumps({'type': 'done', 'value': ''}) + '\n'
        return StreamingResponse(stream())
    monkeypatch.setattr(server, 'chat_stream', fake_chat)
    with client:
        device = pair(manager, client)
        project = create_project('手机持久回复')
        chat = conversations.create(project['id'], 'deepseek')
        payload = {'request_id': secrets.token_hex(16), 'request': {'project_id': project['id'], 'conversation_id': chat['id'], 'source_ids': [], 'text': '测试'}}
        started = client.post('/api/mobile/chat', json=payload)
        assert started.status_code == 200, started.text
        ident = started.json()['run_id']
        retry = client.post('/api/mobile/chat', json=payload).json()
        assert retry['run_id'] == ident and retry['reused']
        mismatch = {**payload, 'request': {**payload['request'], 'text': '不同问题'}}
        assert client.post('/api/mobile/chat', json=mismatch).status_code == 409
        # The original POST has ended; the server continues producing the reply.
        for _ in range(100):
            row = client.get('/api/mobile/chat/' + ident).json()
            if row['status'] == 'completed':
                break
            time.sleep(.03)
        assert row['status'] == 'completed' and row['content'] == '回复第一段；回复第二段'
        assert client.get('/api/mobile/chat/' + ('f' * 32)).status_code == 404
        client.post('/mobile/logout')
        assert client.get('/api/mobile/chat/' + ident).status_code == 401


def test_phone_clamps_full_access_and_rechecks_revocation(monkeypatch):
    from studyforge.agent_permissions import configuration
    monkeypatch.setattr('studyforge.ai.settings', lambda: {'agent_permission_mode': 'full_access'})
    manager, client = mobile_client()
    with client:
        ident = pair(manager, client)
        token = remote_device.set(ident)
        try:
            assert configuration() == 'approval'
            execute("UPDATE mobile_devices SET status='revoked' WHERE id=?", (ident,))
            with pytest.raises(PermissionError):
                configuration()
        finally:
            remote_device.reset(token)


def test_tls_root_is_persistent_and_has_host_names(tmp_path):
    from cryptography import x509
    from cryptography.hazmat.primitives import hashes
    from studyforge.mobile_tls import certificates
    first = certificates(tmp_path, ['192.168.1.23'])
    root_before = Path(first['ca']).read_bytes()
    second = certificates(tmp_path, ['192.168.1.24'])
    assert root_before == Path(second['ca']).read_bytes()
    leaf = x509.load_pem_x509_certificate(Path(second['certificate']).read_bytes())
    root = x509.load_der_x509_certificate(root_before)
    leaf.verify_directly_issued_by(root)
    names = leaf.extensions.get_extension_for_class(x509.SubjectAlternativeName).value
    assert '192.168.1.24' in str(names) and 'localhost' in str(names)
    assert first['fingerprint'] == second['fingerprint']


def test_gateway_allowlist_and_lan_boundary():
    assert private_address('192.168.1.2') and private_address('127.0.0.1')
    assert not private_address('8.8.8.8') and not private_address('100.64.0.1')
    assert allowed('POST', '/api/sources/' + 'a' * 32 + '/parse')
    assert not allowed('DELETE', '/api/projects/' + 'a' * 32)
    assert not allowed('PUT', '/api/settings')
    for path in ('/api/homework/'+'a'*32+'/draft','/api/coding/'+'a'*32+'/preview','/api/coding/'+'a'*32+'/submit'):
        assert allowed('PUT' if path.endswith('/draft') else 'POST',path)
    assert not allowed('DELETE','/api/homework/'+'a'*32)


def phone_worksheet(monkeypatch, coding=False):
    from studyforge import ai
    from studyforge.db import uid, j
    ai.save_settings({'deepseek_model':'phone-homework','agent_tools_enabled':False,
                      'background_learning_enabled':False})
    project=create_project('手机作业项目')
    chat=conversations.create(project['id'],'deepseek','phone-homework')
    ident=uid()
    question={'index':0,'topic':'变量与计算','question':'说明变量的含义，并计算2+3。'}
    solution={'index':0,'solution':'PHONE_PRIVATE_SOLUTION：2+3=5；变量表示指定的量。','rubric':'说明量的含义并正确计算'}
    if coding:
        question.update(topic='函数加法',question='用 JavaScript 实现 add(a,b)，返回两数之和。',language='javascript')
        solution.update(solution='function add(a,b){return a+b;} // PHONE_PRIVATE_CODE',test_code='检查 add(2,3)===5')
    execute('INSERT INTO exercise_sets VALUES (?,?,?,?,?,?,?,?,?)',
            (ident,project['id'],'[]','手机独立作业',j([question]),j([solution]),2,'assigned',now()))

    class Teacher:
        model='phone-homework'
        effort=None
        def __init__(self):self.calls=[];self.closed=0
        async def close(self):self.closed+=1
        async def complete(self,system,user,**kwargs):
            self.calls.append(system)
            assert remote_device.get(), 'Phone-originated background work lost its permission identity'
            await asyncio.sleep(.06)
            if '从学生作答文件中提取' in system:
                return j({'answers':['我的完整作答'],'uncertainties':[]})
            if '逐题诊断错因并补讲' in system:
                data=json.JSONDecoder().raw_decode(user)[0];causes=data['student_causes']
                known=bool(causes) or coding
                return j({'items':[{'index':0,'kind':'implementation' if coding else 'terminology',
                    'confidence':'self_report' if causes else 'supported' if known else 'uncertain',
                    'cause':'运算符与要求不符' if coding else '需确认变量的含义',
                    'evidence':[{'source':'questionnaire' if causes else 'answer',
                                 'quote':causes[0]['text'] if causes else data['student_answers'][0]}],
                    'clarification':'作答时哪个量的含义不清楚？',
                    'explanation':'这里需要先说明变量所表示的量，再根据题意选择运算。两数之和使用加法，2加3可以逐次加三个1得到5。代码要返回a+b，再用简单数值检查返回值，不能把求和写成求差。',
                    'check':'逐个说明变量的含义，并用2和3核对计算。'}]})
            if '编程教师' in system:
                return j({'score':50,'feedback':'把求和写成求差','issues':['return a-b 与题意不符']})
            return j({'scores':[50],'feedback':['变量含义需核对'],'uncertainties':[]})
    teacher=Teacher()
    monkeypatch.setattr(server,'client_for',lambda *args,**kwargs:teacher)
    async def choose(*args,**kwargs):pass
    monkeypatch.setattr(server,'choose_generation_model',choose)
    return project,chat,ident,teacher


def await_phone_attempt(client, attempt, expected):
    for _ in range(200):
        response=client.get('/api/homework/'+attempt['id']+'?project_id='+attempt['project_id'])
        assert response.status_code==200,response.text
        value=response.json()
        if value['status'] in expected:return value
        if value['status']=='failed':pytest.fail(value['error'])
        time.sleep(.025)
    pytest.fail('Phone homework did not reach '+str(expected))


def test_phone_submission_unlocks_docx_and_questionnaire_resumes_one_grader(monkeypatch):
    from io import BytesIO
    from docx import Document
    from studyforge import homework
    project,chat,worksheet,teacher=phone_worksheet(monkeypatch)
    other=create_project('另一个手机项目')
    manager,client=mobile_client()
    with client:
        device=pair(manager,client)
        assert client.get('/homework.js').status_code==200
        before=client.get('/api/exercises/'+worksheet)
        assert 'solutions' not in before.json() and 'PHONE_PRIVATE_SOLUTION' not in before.text
        payload={'answers':[''],'provider':'deepseek','model':teacher.model,
                 'conversation_id':chat['id'],'background':True,'detailed_feedback':True}
        assert client.post('/api/exercises/'+worksheet+'/submit',json=payload).status_code==400
        assert not many('SELECT id FROM homework_attempts') and not teacher.calls
        payload['answers']=['我不知道这个变量表示什么，只写了2+3=4']
        submitted=client.post('/api/exercises/'+worksheet+'/submit',json=payload)
        assert submitted.status_code==200,submitted.text
        attempt=submitted.json()
        assert attempt['solution_artifact'] and not attempt['feedback_artifact']
        assert unj(one('SELECT payload_json FROM jobs WHERE id=?',(attempt['job_id'],))['payload_json'])['_mobile_device_id']==device
        document=client.get(attempt['solution_artifact']['download_url'])
        assert document.status_code==200
        assert 'PHONE_PRIVATE_SOLUTION' in '\n'.join(p.text for p in Document(BytesIO(document.content)).paragraphs)
        attempt=await_phone_attempt(client,attempt,{'awaiting_causes'})
        assert not attempt['feedback_artifact'] and len(attempt['questions'])==1
        response={'responses':[{'index':0,'kinds':['terminology'],'text':'我分不清变量表示的量和计算结果'}]}
        route='/api/homework/'+attempt['id']
        assert client.put(route+'/draft?project_id='+other['id'],json=response).status_code==404
        saved=client.put(route+'/draft?project_id='+project['id'],json=response)
        assert saved.status_code==200,saved.text
        assert client.get(route+'?project_id='+project['id']).json()['cause_draft']==response['responses']
        resumed=client.post(route+'/causes?project_id='+project['id'],json=response)
        assert resumed.status_code==200,resumed.text
        attempt=await_phone_attempt(client,attempt,{'completed'})
        report=client.get(attempt['feedback_artifact']['download_url'])
        text='\n'.join(p.text for p in Document(BytesIO(report.content)).paragraphs)
        assert '错因分析' in text and '针对性讲解' in text and '我分不清变量' in text
        assert len(many('SELECT id FROM submissions'))==1
        retry=client.post('/api/exercises/'+worksheet+'/submit',json=payload)
        assert retry.status_code==200 and retry.json()['reused']
        assert client.get('/api/projects/'+other['id']+'/homework').json()==[]
        assert client.get('/api/projects/'+other['id']+'/artifacts').json()==[]
        assert len([call for call in teacher.calls if '批改者' in call or '评分复核者' in call])==2
        assert homework.public_attempt(homework.get_attempt(attempt['id']))['status']=='completed'


def test_phone_code_preview_keeps_language_encoding_and_no_fake_final_score(monkeypatch):
    from studyforge import coding
    project,chat,worksheet,teacher=phone_worksheet(monkeypatch,coding=True)
    monkeypatch.setattr(coding,'_run_assignment_tests',lambda *args:{
        'ran':False,'score':0,'reason':'测试中未运行Docker，仅静态审查'})
    manager,client=mobile_client()
    code='// 中文代码说明\nfunction add(a,b){return a-b;}'
    with client:
        pair(manager,client)
        preview=client.post('/api/coding/'+worksheet+'/preview',
                            files={'file':('answer.js',code.encode('utf-16'))})
        assert preview.status_code==200,preview.text
        assert preview.json()['code']==code
        assert not many('SELECT id FROM homework_attempts') and not teacher.calls
        assert client.post('/api/coding/'+worksheet+'/preview',
                           files={'file':('answer.py',b'print(1)')}).status_code==400
        submitted=client.post('/api/coding/'+worksheet+'/submit',
            data={'provider':'deepseek','model':teacher.model,'conversation_id':chat['id'],
                  'background':'true','detailed_feedback':'true'},
            files={'file':('answer.js',code.encode('utf-8'))})
        assert submitted.status_code==200,submitted.text
        attempt=await_phone_attempt(client,submitted.json(),{'completed'})
        assert attempt['score'] is None and attempt['provisional']
        assert attempt['solution_artifact'] and attempt['feedback_artifact']
        original=unj(one('SELECT answers_json FROM homework_attempts WHERE id=?',(attempt['id'],))['answers_json'])
        assert original==[code]


def test_phone_answer_file_uses_its_file_limit_instead_of_json_limit(monkeypatch):
    project,chat,worksheet,teacher=phone_worksheet(monkeypatch)
    manager,client=mobile_client()
    with client:
        pair(manager,client)
        response=client.post('/api/exercises/'+worksheet+'/extract-answers',
            data={'provider':'deepseek','model':teacher.model},
            files={'file':('answers.txt',b'answer '*450000)})
        assert response.status_code==200,response.text
        assert response.json()['answers']==['我的完整作答']
        assert teacher.closed==1
        assert not many('SELECT id FROM homework_attempts')
        # A multipart code file just below 2 MB is valid despite its envelope.
        from studyforge.db import j,uid
        code=uid()
        execute('INSERT INTO exercise_sets VALUES (?,?,?,?,?,?,?,?,?)',
            (code,project['id'],'[]','JavaScript大代码预览',j([{'index':0,'topic':'编码','question':'实现函数','language':'javascript'}]),j([]),1,'assigned',now()))
        payload=b'//' + b'a'*(2*1024*1024-4) + b'\n'
        preview=client.post('/api/coding/'+code+'/preview',files={'file':('answer.js',payload)})
        assert preview.status_code==200,preview.text[:200]
        assert len(preview.json()['code'])==len(payload)
