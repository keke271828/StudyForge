import asyncio
import base64
import io

import httpx
import pytest
from PIL import Image

from studyforge import agent_tools
from studyforge.agent_tools import AgentRuntime
from studyforge.db import create_project, many, uid
from studyforge.exporters import export_document


@pytest.mark.parametrize('format', ['docx', 'pdf', 'pptx'])
def test_generated_illustration_is_embedded_in_the_selected_document(monkeypatch, format):
    project = create_project('图文成果'); runtime = AgentRuntime(project['id'], publishing=False)
    monkeypatch.setattr(agent_tools, 'get_secret', lambda key: 'test-key')
    buffer = io.BytesIO(); Image.new('RGB', (80, 40), 'blue').save(buffer, 'PNG')
    original_client = httpx.AsyncClient
    def response(request): return httpx.Response(200, json={'data': [{'b64_json': base64.b64encode(buffer.getvalue()).decode()}]})
    monkeypatch.setattr(httpx, 'AsyncClient', lambda *a, **kw: original_client(transport=httpx.MockTransport(response)))
    picture = asyncio.run(runtime.execute('sf_generate_image', {'prompt':'学习示意图'}))
    assert 'error' not in picture and 'artifact' not in picture and many('SELECT * FROM artifacts') == []
    content = '# 图文讲解\n图前解释。\n\n![蓝色示意图](' + picture['image_url'] + ')\n\n图后解释。\n\n## 系统性总结\n图像辅助理解。'
    path = export_document(content, format, '图文讲解', uid() + '.' + format)
    if format == 'docx':
        from docx import Document
        doc = Document(path)
        assert len(doc.inline_shapes) == 1 and '蓝色示意图' in '\n'.join(p.text for p in doc.paragraphs)
    elif format == 'pdf':
        import fitz
        with fitz.open(path) as pdf: assert sum(len(page.get_images()) for page in pdf) == 1
    else:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
        prs = Presentation(path)
        pictures = [shape for slide in prs.slides for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
        assert len(pictures) == 1 and pictures[0].image.blob == buffer.getvalue()


@pytest.mark.parametrize('format', ['docx', 'pdf', 'pptx'])
def test_missing_illustration_is_reported_before_exporting_a_fake_image(format):
    with pytest.raises(ValueError, match='图片不存在'):
        export_document('![缺失图](/api/agent/assets/' + uid() + '/image)', format, '不能伪造', uid() + '.' + format)
