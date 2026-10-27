from zipfile import ZipFile

from docx import Document
from lxml import etree

from studyforge.document_render import inline_parts, rich_html, table_row, blocks
from studyforge.exporters import export_document


def test_bold_math_and_table_absolute_values_are_actual_equations(tmp_path):
    content = r'''# **极限与 $N=\max\{N_1,N_2\}$**

**选择 $\varepsilon=\frac{b-a}{2}$，因为两个开邻域不相交。**

| 对象 | 含义 |
| --- | --- |
| **$|x_n-a|<\varepsilon$** | **严格小于** |

代码 `$x$` 保持代码。
'''
    assert list(inline_parts(r'**取 $\frac{b-a}{2}$ 即可**')) == [
        ('bold', '取 '), ('math', r'\frac{b-a}{2}'), ('bold', ' 即可')]
    assert table_row(r'| **$|x_n-a|$** | 含义 |') == [r'**$|x_n-a|$**', '含义']
    path = export_document(content, 'docx', '加粗公式验收', 'nested-math.docx')
    document = Document(path)
    assert len(document.tables[0].columns) == 2
    assert any(run.bold for p in document.paragraphs for run in p.runs if '选择' in run.text)
    with ZipFile(path) as package:
        root = etree.fromstring(package.read('word/document.xml'))
    ns = {'m': 'http://schemas.openxmlformats.org/officeDocument/2006/math',
          'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    assert len(root.xpath('//m:oMath', namespaces=ns)) == 3
    assert root.xpath('//m:f', namespaces=ns)
    ordinary = ''.join(root.xpath('//w:t/text()', namespaces=ns))
    assert r'\frac' not in ordinary and r'\max' not in ordinary
    assert '$x$' in ordinary  # code is deliberately literal
    html = rich_html(content)
    assert r'\(\varepsilon=\frac{b-a}{2}\)' in html
    assert '<strong>选择 </strong>' in html and '<code>$x$</code>' in html


def test_code_pipes_and_thematic_breaks_do_not_create_or_extend_tables():
    assert list(blocks('用 `{|c|c|}` 指定列。\n---')) == [('paragraph', '用 `{|c|c|}` 指定列。')]
    content = '对象 | 含义\n--- | ---\n`{|c|}` | 一列\n后面解释 `{c|c}`。\n---'
    parsed = list(blocks(content))
    assert parsed[0] == ('table', [['对象', '含义'], ['`{|c|}`', '一列']])
    assert parsed[1] == ('paragraph', '后面解释 `{c|c}`。')
    one_column = list(blocks('| 标题 |\n|---|\n| 内容 |\n\n结束。'))
    assert one_column[0] == ('table', [['标题'], ['内容']])


def test_native_word_table_repeats_header_and_keeps_first_row_with_it(tmp_path):
    path = export_document('| 对象 | 含义 |\n|---|---|\n| 甲 | 内容 |', 'docx', '表头验收', 'header.docx')
    document = Document(path)
    table = document.tables[0]
    assert table.rows[0]._tr.xpath('./w:trPr/w:tblHeader')
    assert all(cell.paragraphs[0].paragraph_format.keep_with_next for cell in table.rows[0].cells)


def test_word_tables_keep_uneven_row_padding_and_separate_native_cells():
    content = r'''| 节点 | 条件 |
|---|---|
| 甲 | **$r>0$** | 第三列补充 |
| 乙 |

表格后的解释。

| 独立小表 |
|---|
| 丙 |
'''
    document = Document(export_document(content, 'docx', '不等长表格', 'uneven-tables.docx'))
    first,second=document.tables
    assert [[cell.text for cell in row.cells] for row in first.rows]==[
        ['节点','条件',''],['甲','','第三列补充'],['乙','','']]
    assert len(first.rows[1].cells[1]._tc.xpath('.//m:oMath'))==1
    assert [[cell.text for cell in row.cells] for row in second.rows]==[['独立小表'],['丙']]
    assert any(p.text=='表格后的解释。' for p in document.paragraphs)
    assert first.rows[0]._tr.xpath('./w:trPr/w:tblHeader') and second.rows[0]._tr.xpath('./w:trPr/w:tblHeader')
    assert not first.rows[1]._tr.xpath('./w:trPr/w:tblHeader')


def test_independent_ordered_lists_restart_and_preserve_their_start_in_word_and_pdf():
    content = '## Earlier\n7. Alpha\n1. Beta\n\n## New section\n1. Gamma\n1. Delta'
    document = Document(export_document(content, 'docx', 'List check', 'lists.docx'))
    items = [p for p in document.paragraphs if p.text in ('Alpha', 'Beta', 'Gamma', 'Delta')]
    numbering = document.part.numbering_part.element
    number_ids = [p._p.pPr.numPr.numId.val for p in items]
    assert number_ids[0] == number_ids[1] and number_ids[2] == number_ids[3]
    assert number_ids[0] != number_ids[2]
    starts = [numbering.num_having_numId(index).lvlOverride_lst[0].startOverride.val
              for index in (number_ids[0], number_ids[2])]
    assert starts == [7, 1]
    import fitz
    with fitz.open(export_document(content, 'pdf', 'List check', 'lists.pdf')) as pdf:
        rendered = '\n'.join(page.get_text() for page in pdf)
    assert all(value in rendered for value in ('7. Alpha', '8. Beta', '1. Gamma', '2. Delta'))
