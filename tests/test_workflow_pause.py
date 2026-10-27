"""Real workflow boundaries with controlled model replies and isolated storage."""
import asyncio
from pathlib import Path
import threading

from docx import Document
from fastapi import HTTPException
import pytest

from studyforge import ai, coding, conversations, exercises, knowledge, learning, server, workflows
from studyforge.db import create_job, create_project, execute, j, many, now, one, uid, unj, update_job
from studyforge.desktop_lifecycle import recover_interrupted_jobs
from studyforge.generation_control import active_generations, control_for
from studyforge.job_control import active_jobs, cancel_job


def material(tmp_path):
    ai.save_settings({'deepseek_model':'pause-model','background_learning_enabled':False})
    project=create_project('恢复测试')
    chat=conversations.create(project['id'],'deepseek','pause-model')['id']
    source=uid();path=tmp_path/'原文.txt';path.write_text('相关不能推出因果。需要检查共同原因。',encoding='utf-8')
    execute('INSERT INTO sources(id,project_id,name,kind,original_path,segments_json,status,created_at) VALUES (?,?,?,?,?,?,?,?)',
            (source,project['id'],path.name,'.txt',str(path),'[]','ready',now()))
    return project['id'],chat,source,path


def worksheet(project,source,status='assigned',language=None):
    set_id=uid();question={'index':0,'question':'解释为什么相关不足以证明因果。','topic':'因果'}
    if language:question.update(language=language,starter_code='def check(): pass')
    execute('INSERT INTO exercise_sets VALUES (?,?,?,?,?,?,?,?,?)',
            (set_id,project,j([source]),'迁移题单',j([question]),j([{'index':0,'solution':'PRIVATE_REFERENCE','rubric':'检查共同原因','test_code':'PRIVATE_TESTS'}]),2,status,now()))
    if status!='assigned':
        execute('INSERT INTO submissions VALUES (?,?,?,?,?,?,?)',
                (uid(),set_id,j(['已完成作答']),j({'scores':[79],'feedback':['还缺论证']}),79,'{}',now()))
    return set_id


async def wait_status(job_id,status):
    for _ in range(600):
        row=one('SELECT * FROM jobs WHERE id=?',(job_id,))
        if row['status']==status:return row
        if row['status']=='failed':pytest.fail(row['error'])
        await asyncio.sleep(.005)
    pytest.fail(f'未到达 {status}: {row}')


async def stop_worker(job_id,task=None):
    task=task or active_jobs.get(job_id)
    task.cancel()
    with pytest.raises(asyncio.CancelledError):await task
    await asyncio.sleep(0)
    assert job_id not in active_generations


class Replies:
    model='pause-model'
    def __init__(self,reply):
        self.reply=reply;self.calls=[];self.entered=asyncio.Event();self.release=asyncio.Event();self.hold=True
    async def complete(self,system,user,**kwargs):
        self.calls.append((system,user))
        if self.hold:
            self.hold=False;self.entered.set();await self.release.wait()
        return self.reply(system,user)
    async def close(self):pass


async def pause_request(job,client):
    await asyncio.wait_for(client.entered.wait(),3)
    await server.pause_generation(job)
    assert one('SELECT status FROM jobs WHERE id=?',(job,))['status']=='pausing'
    client.release.set();await wait_status(job,'paused')


def test_read_restart_reuses_parts_and_keeps_chat_scope(tmp_path,monkeypatch):
    project,chat,source,_=material(tmp_path)
    monkeypatch.setattr(learning,'explanation_batches',lambda digest:['part-A','part-B','part-C'])
    async def run():
        client=Replies(lambda s,u:'# 阅读笔记\n'+u)
        monkeypatch.setattr(server,'client_for',lambda *a,**k:client)
        value=server.ReadIn(project_id=project,source_ids=[source],provider='deepseek',model='pause-model',conversation_id=chat)
        start=await server.read_material(value);job=start['job_id']
        await pause_request(job,client)
        assert len(client.calls)==1 and many('SELECT * FROM reading_sessions')==[]
        assert (await server.read_material(value))['job_id']==job
        await stop_worker(job)
        await server.resume_generation(job);result=await wait_status(job,'completed')
        assert len(client.calls)==3 and len(many('SELECT * FROM reading_sessions'))==1
        assert len(unj(result['result_json'])['source_reading'])==1
        assert server.jobs(project)[0]['conversation_id']==chat
        assert len(unj(one('SELECT notes_json FROM reading_sessions')['notes_json']))==3
    asyncio.run(run())


def test_knowledge_restart_saves_only_once_without_new_model_request(tmp_path,monkeypatch):
    project,chat,source,_=material(tmp_path)
    async def run():
        client=Replies(lambda s,u:'# 通读笔记');client.hold=False
        monkeypatch.setattr(server,'client_for',lambda *a,**k:client)
        value=server.ReadIn(project_id=project,source_ids=[source],provider='deepseek',model='pause-model',conversation_id=chat)
        reading=(await wait_status((await server.read_material(value))['job_id'],'completed'))
        reading_id=unj(reading['result_json'])['reading_id']
        original=server.read_original_sources(server.source_records([source],project))[0]
        graph={'concepts':[{'id':'cause','name':'因果','domain':'逻辑','description':'相关并不足够。','basis':'source',
              'evidence':[{'source_id':source,'locator':original['segments'][0]['locator'],'quote':'相关不能推出因果'}]}],'relations':[]}
        client=Replies(lambda s,u:j(graph));monkeypatch.setattr(server,'client_for',lambda *a,**k:client)
        value=server.KnowledgeIn(**value.model_dump(),reading_id=reading_id)
        job=(await server.build_knowledge(value))['job_id'];await pause_request(job,client)
        assert many('SELECT * FROM knowledge_maps')==[]
        await stop_worker(job);await server.resume_generation(job);await wait_status(job,'completed')
        assert len(client.calls)==1 and len(many('SELECT * FROM knowledge_maps'))==1
        assert server.jobs(project)[0]['reading_id']==reading_id
    asyncio.run(run())


def test_quiz_pause_restart_reuses_review_and_scores_once(tmp_path,monkeypatch):
    project,chat,source,_=material(tmp_path);set_id=worksheet(project,source)
    async def run():
        client=Replies(lambda s,u:j({'scores':[80],'feedback':['解释正确'],'uncertainties':[]}))
        async def factory(*a,**k):return client
        monkeypatch.setattr(server,'assessment_client',factory)
        value=server.SubmitIn(answers=['可能存在共同原因。'],background=True,provider='deepseek',model='pause-model',conversation_id=chat)
        job=(await server.submit_answers(set_id,value))['job_id'];await pause_request(job,client)
        assert 'solutions' not in exercises.public_set(set_id,reveal=True)
        with pytest.raises(HTTPException,match='409'):await server.submit_answers(set_id,value)
        await stop_worker(job);await server.resume_generation(job);await wait_status(job,'completed')
        assert len(client.calls)==2 and len(many('SELECT * FROM submissions'))==1
        assert one('SELECT status FROM exercise_sets WHERE id=?',(set_id,))['status']=='passed'
        artifact=one("SELECT * FROM artifacts WHERE kind='solutions'")
        assert unj(artifact['metadata_json'])['conversation_id']==chat
        assert 'PRIVATE_REFERENCE' in '\n'.join(p.text for p in Document(artifact['path']).paragraphs)
    asyncio.run(run())


def test_pause_after_grade_commit_does_not_regrade_on_restart(tmp_path,monkeypatch):
    project,chat,source,_=material(tmp_path);set_id=worksheet(project,source)
    async def run():
        client=Replies(lambda s,u:j({'scores':[79],'feedback':['缺少论证']}));client.hold=False
        async def factory(*a,**k):return client
        monkeypatch.setattr(server,'assessment_client',factory)
        export=workflows.export;first=True
        async def hold_export(*a,**k):
            nonlocal first
            if first:
                first=False;workflows.current_control.get().pause();await workflows.checkpoint()
            return await export(*a,**k)
        monkeypatch.setattr(workflows,'export',hold_export)
        job=(await server.submit_answers(set_id,server.SubmitIn(answers=['只知道相关'],provider='deepseek',model='pause-model',background=True)))['job_id']
        await wait_status(job,'paused')
        assert one('SELECT status FROM exercise_sets')['status']=='feynman_required'
        assert len(many('SELECT * FROM submissions'))==1 and len(client.calls)==2
        await stop_worker(job);await server.resume_generation(job);await wait_status(job,'completed')
        assert len(client.calls)==2 and len(many('SELECT * FROM submissions'))==1
        assert len(many("SELECT * FROM artifacts WHERE kind='solutions'"))==1
    asyncio.run(run())


def test_coding_generation_restart_keeps_snapshot_and_hides_reference(tmp_path,monkeypatch):
    project,chat,source,_=material(tmp_path)
    draft={'title':'线性关系','problem':'输入实数x，返回2x；使用MATLAB。','starter_code':'function y = twice(x)\n% TODO\nend',
           'reference_code':'PRIVATE_REFERENCE','topic':'函数','test_code':''}
    def reply(system,user):
        if '独立解题员' in system:return j({'approach':'核对输入输出','passed':True,'blocking_issues':[],'notes':[]})
        if '审查员' in system:return j({'passed':True,'blocking_issues':[],'notes':[]})
        return j(draft)
    async def run():
        client=Replies(reply)
        async def factory(*a,**k):return client
        monkeypatch.setattr(server,'assessment_client',factory)
        job=(await server.coding_generate(server.CodingIn(project_id=project,source_ids=[source],language='matlab',
                provider='deepseek',model='pause-model',background=True,conversation_id=chat)))['job_id']
        await pause_request(job,client);await stop_worker(job)
        # Memory may change while paused; the draft must still reuse its saved input.
        monkeypatch.setattr('studyforge.learning_memory.profile_context',lambda digest:'后来变化的学习画像')
        await server.resume_generation(job);await wait_status(job,'completed')
        assert len(client.calls)==4 and len(many('SELECT * FROM exercise_sets'))==1
        artifact=one("SELECT * FROM artifacts WHERE kind='exercise'")
        text='\n'.join(p.text for p in Document(artifact['path']).paragraphs)
        assert 'TODO' in text and 'PRIVATE_REFERENCE' not in text
        assert unj(artifact['metadata_json'])['conversation_id']==chat
    asyncio.run(run())


def test_container_step_finishes_before_pause_and_is_not_run_twice(tmp_path,monkeypatch):
    project,chat,source,_=material(tmp_path);set_id=worksheet(project,source,language='python')
    entered=threading.Event();release=threading.Event();runs=[]
    def container(*args):
        runs.append(args);entered.set();assert release.wait(4)
        return {'ran':True,'passed':True,'score':100,'output':'3 tests passed'}
    monkeypatch.setattr(coding,'_run_assignment_tests',container)
    async def run():
        client=Replies(lambda s,u:j({'score':90,'feedback':'代码正确','issues':[]}));client.hold=False
        async def factory(*a,**k):return client
        monkeypatch.setattr(server,'assessment_client',factory)
        value=server.AssessmentAIIn(background=True,provider='deepseek',model='pause-model',conversation_id=chat)
        job=(await server.enqueue_assessment('coding_grade',value,set_id,'def check(): return True'))['job_id']
        assert await asyncio.to_thread(entered.wait,3)
        await server.pause_generation(job)
        assert one('SELECT status FROM jobs WHERE id=?',(job,))['status']=='pausing'
        release.set();await wait_status(job,'paused');assert client.calls==[]
        await stop_worker(job);await server.resume_generation(job);await wait_status(job,'completed')
        assert len(runs)==1 and len(client.calls)==1 and len(many('SELECT * FROM submissions'))==1
        assert one('SELECT score FROM submissions')['score']==98
    asyncio.run(run())


@pytest.mark.parametrize('kind',['feynman','repair','advance'])
def test_progression_restart_creates_one_child_set_and_one_artifact(tmp_path,monkeypatch,kind):
    project,chat,source,_=material(tmp_path)
    status='feynman_required' if kind=='feynman' else 'repair_required' if kind=='repair' else 'retry_ready'
    parent=worksheet(project,source,status)
    if kind=='repair':execute('INSERT INTO practice_options VALUES (?,?,?)',(parent,'adaptive',now()))
    async def run():
        client=Replies(lambda s,u:j({'score':85,'passed':True,'feedback':'理解过关','gaps':[]}));client.hold=False
        async def factory(*a,**k):return client
        monkeypatch.setattr(server,'assessment_client',factory)
        entered=asyncio.Event();release=asyncio.Event();generated=[]
        async def new_set(client,project_id,sources,digest,title,difficulty=1,**kwargs):
            entered.set();await release.wait()
            await client.complete('新增迁移题','明确题意')
            child=worksheet(project_id,source);execute('UPDATE exercise_sets SET difficulty=? WHERE id=?',(difficulty,child))
            if spec:=kwargs.get('repair_spec'):
                # A generated repair set must satisfy the real target contract;
                # this test controls scheduling, not the repair validation.
                questions=[];solutions=[]
                for index,target in enumerate(spec['targets']):
                    questions.append({'index':index,'question':f'新情境{index}：观察到同伴讨论与成绩相关，怎样判断共同原因？请写出理由。','topic':target['topic']})
                    solutions.append({'index':index,'solution':'应区分关联与因果，并检查原有动机等共同原因。',
                        'rubric':'检查共同原因与论证依据。','repair_target_index':target['index'],
                        'repair_focus':'用新的学习情境检查对应失分知识点，不预先确定学生的错因。'})
                execute('UPDATE exercise_sets SET questions_json=?,solutions_json=? WHERE id=?',(j(questions),j(solutions),child))
            generated.append(child);return child,[]
        monkeypatch.setattr(server,'generate_set',new_set)
        values=server.AssessmentAIIn(provider='deepseek',model='pause-model',background=True,conversation_id=chat)
        if kind=='feynman':start=await server.feynman(parent,server.FeynmanIn(**values.model_dump(),explanation='先分清相关和因果，再检查共同原因、选择偏差与反向因果，只有符合必要条件才能进一步论证。'*4))
        elif kind=='repair':start=await server.repair_practice(parent,values)
        else:start=await server.continue_progression(parent,values)
        job=start['job_id'];await asyncio.wait_for(entered.wait(),3)
        await server.pause_generation(job);release.set();await wait_status(job,'paused')
        # The Feynman grade is already committed before child generation pauses.
        initial_calls=len(client.calls)
        if kind=='feynman':assert one('SELECT status FROM exercise_sets WHERE id=?',(parent,))['status']=='retry_ready'
        await stop_worker(job);await server.resume_generation(job);await wait_status(job,'completed')
        assert len(client.calls)==initial_calls+1 and len(generated)==1
        assert len(many('SELECT * FROM exercise_sets'))==2
        assert len(many("SELECT * FROM artifacts WHERE kind='exercise'"))==1
        assert one('SELECT difficulty FROM exercise_sets WHERE id=?',(generated[0],))['difficulty']==(2 if kind=='repair' else 3)
        if kind=='feynman':assert sum('费曼学习法评审' in s for s,u in client.calls)==1
    asyncio.run(run())


@pytest.mark.parametrize('changed',['source','worksheet'])
def test_resume_rejects_mutated_inputs_before_model_calls(tmp_path,monkeypatch,changed):
    project,chat,source,path=material(tmp_path)
    async def run():
        client=Replies(lambda s,u:'# 阅读笔记')
        monkeypatch.setattr(server,'client_for',lambda *a,**k:client)
        if changed=='source':
            job=(await server.read_material(server.ReadIn(project_id=project,source_ids=[source],provider='deepseek',model='pause-model')))['job_id']
        else:
            set_id=worksheet(project,source)
            async def factory(*a,**k):return client
            monkeypatch.setattr(server,'assessment_client',factory)
            client.reply=lambda s,u:j({'scores':[80],'feedback':['正确']})
            job=(await server.submit_answers(set_id,server.SubmitIn(answers=['完整作答'],provider='deepseek',model='pause-model',background=True)))['job_id']
        await pause_request(job,client)
        if changed=='source':path.write_text('修改了论证依据。',encoding='utf-8')
        else:execute('UPDATE exercise_sets SET difficulty=5 WHERE id=?',(set_id,))
        with pytest.raises(HTTPException) as exc:await server.resume_generation(job)
        assert exc.value.status_code==409 and len(client.calls)==1
        assert one('SELECT status FROM jobs WHERE id=?',(job,))['status']=='paused'
        await cancel_job(job)
        assert job not in active_generations
        assert many('SELECT * FROM reading_sessions')==[] and many('SELECT * FROM submissions')==[]
    asyncio.run(run())


def test_shutdown_recovers_supported_jobs_and_lists_all_paused_jobs(tmp_path):
    project,chat,source,_=material(tmp_path)
    jobs=[]
    for _ in range(35):
        payload={'project_id':project,'source_ids':[source],'conversation_id':chat,'provider':'deepseek'}
        job=create_job(project,'read',payload);jobs.append(job)
    for _ in range(32):
        job=create_job(project,'import',{});update_job(job,'completed',100,{})
    counts=recover_interrupted_jobs()
    assert counts['paused_generations']==35
    listed=server.jobs(project)
    assert len([x for x in listed if x['status']=='paused'])==35
    assert len([x for x in listed if x['status']=='completed'])==30
    assert all(x['can_pause'] for x in listed if x['status']=='paused')


def test_actual_restart_releases_only_interrupted_task_submission_locks(tmp_path):
    project,chat,source,_=material(tmp_path);parent=worksheet(project,source)
    payload={'workflow_version':1,'set_id':parent,'values':{'answers':['作答']}}
    job=create_job(project,'quiz_grade',payload);update_job(job,'paused',30,{})
    # Simulate a killed process: its finally block never releases the lock.
    execute('INSERT INTO practice_submission_locks VALUES (?,?,?)',(parent,job+':crashed-owner','2099-01-01T00:00:00+00:00'))
    other=worksheet(project,source)
    execute('INSERT INTO practice_submission_locks VALUES (?,?,?)',(other,'unrelated-mcp-owner','2099-01-01T00:00:00+00:00'))
    recover_interrupted_jobs()
    assert one('SELECT 1 FROM practice_submission_locks WHERE exercise_set_id=?',(parent,)) is None
    assert one('SELECT token FROM practice_submission_locks WHERE exercise_set_id=?',(other,))['token']=='unrelated-mcp-owner'


@pytest.mark.parametrize('entry',['coding','hint'])
def test_mcp_grading_and_hints_cannot_bypass_paused_task(tmp_path,entry):
    project,chat,source,_=material(tmp_path);parent=worksheet(project,source,language='python' if entry=='coding' else None)
    payload={'workflow_version':1,'set_id':parent,'values':{'provider':'deepseek'}}
    job=create_job(project,'coding_grade' if entry=='coding' else 'quiz_grade',payload);update_job(job,'paused',30,{})
    if entry=='coding':
        with pytest.raises(ValueError,match='暂停的任务'):asyncio.run(coding.grade_code(None,parent,'def check(): return True'))
    else:
        from studyforge import practice
        with pytest.raises(ValueError,match='批改任务'):practice.get_hint(parent,0,'motivation',1)
    assert many('SELECT * FROM submissions')==[] and many('SELECT * FROM practice_hints')==[]
