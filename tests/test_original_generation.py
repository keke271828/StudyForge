import asyncio
import json
import zipfile
from pathlib import Path

import fitz
import httpx
import pytest
from docx import Document
from fastapi.testclient import TestClient

from studyforge import model_catalog, server, urls
from studyforge.ai import CodexClient, explanation_batches, source_chunks
from studyforge.codex_bridge import CodexBridge
from studyforge.db import create_job, create_project, execute, j, now, one, uid
from studyforge.exporters import export_document
from studyforge.materials import read_original_sources
from studyforge.parsers import parse_file


@pytest.mark.parametrize('encoding', ['utf-8-sig', 'utf-16', 'utf-16-be', 'gb18030'])
def test_windows_text_encodings_preserve_chinese(tmp_path, encoding):
    original = tmp_path / 'source.txt'
    original.write_bytes('弹簧振子：质量与能量守恒。 E=constant'.encode(encoding))
    assert parse_file(original)[0]['text'] == '弹簧振子：质量与能量守恒。 E=constant'


def test_generation_rereads_current_original_not_stale_preview(tmp_path):
    original = tmp_path / 'source.txt'
    original.write_text('完整原文\n' + '中间推导\n' * 5000 + '末尾重要条件：无阻尼。', encoding='utf-8')
    record = dict(id='source', name='力学.txt', original_path=str(original),
                  segments=[dict(locator='预览', text='旧概括内容')])
    fresh = read_original_sources([record])[0]
    full = '\n'.join(item['text'] for item in fresh['segments'])
    assert '旧概括内容' not in full and full.endswith('末尾重要条件：无阻尼。')
    assert fresh['read_manifest']['bytes'] == original.stat().st_size
    original.unlink()
    with pytest.raises(ValueError, match='原文件不存在'):
        read_original_sources([record])


def test_long_original_batches_preserve_all_content_and_provenance():
    text = ''.join(f'不可遗漏的推导{i:05d}。' for i in range(9000))
    chunks = source_chunks([dict(name='教材', segments=[dict(locator='第 5 章', text=text)])])
    batches = explanation_batches(''.join(chunks), max_chars=20000)
    assert len(batches) > 2
    assert all(len(batch) <= 20000 for batch in batches)
    for i in (0, 1000, 4500, 8999):
        assert ''.join(batches).count(f'不可遗漏的推导{i:05d}。') == 1
    assert all('[来源 教材｜第 5 章]' in batch for batch in batches)


def test_word_export_has_native_formulas_and_rereads_in_source_order(tmp_path):
    content = '# 能量守恒\n\n质量 $m$ 与速度 $v$。\n\n$$E=\\frac{1}{2}mv^{2}$$\n\n|量|单位|\n|---|---|\n|质量|kg|\n\n表后重要条件：无阻尼。'
    output = export_document(content, 'docx', '物理讲解', 'formula.docx')
    with zipfile.ZipFile(output) as archive:
        xml = archive.read('word/document.xml').decode('utf-8')
        styles = archive.read('word/styles.xml').decode('utf-8')
    assert '<m:f>' in xml and '<m:sSup>' in xml
    assert '\\frac' not in xml and '$' not in xml and '\ufffd' not in xml
    assert 'w:eastAsia="Microsoft YaHei"' in styles
    text = '\n'.join(part['text'] for part in parse_file(output))
    assert '\\frac{1}{2}' in text
    assert text.index('kg') < text.index('表后重要条件')
    assert Document(output).tables[0].cell(1, 0).text == '质量'


def test_aligned_formulas_render_in_word_without_raw_markup():
    output = export_document('$$\n\\begin{aligned}x&=\\frac12mv^2\\\\y&=\\sqrt{x}\\end{aligned}\n$$',
                             'docx', '多行推导', 'aligned.docx')
    with zipfile.ZipFile(output) as archive:
        xml = archive.read('word/document.xml').decode('utf-8')
    assert '<m:m>' in xml and '<m:f>' in xml and 'aligned' not in xml


def test_pdf_has_readable_chinese_and_formula_images():
    output = export_document('# 振动\n中文变量说明：质量 $m$。\n$$\n\\omega=\\sqrt{\\frac{k}{m}}\n\\qquad E=\\frac12mv^2\n$$\n|公式|中文|\n|---|---|\n|$x(t)=x_0\\cos(\\omega t)+\\frac{v_0}{\\omega}\\sin(\\omega t)$|运动公式|',
                             'pdf', '中文物理讲解', 'formula.pdf')
    with fitz.open(output) as pdf:
        assert '中文物理讲解' in ''.join(page.get_text() for page in pdf)
        assert pdf[0].get_images()
        assert '\\sqrt' not in pdf[0].get_text()


def test_notebook_includes_original_outputs(tmp_path):
    notebook = tmp_path / 'model.ipynb'
    notebook.write_text(json.dumps({'cells': [{'cell_type': 'code', 'source': ['print(7)'],
                                              'outputs': [{'output_type': 'stream', 'text': ['7\n']}]}]}), encoding='utf-8')
    parts = parse_file(notebook)
    assert [part['text'] for part in parts] == ['print(7)', '7']


def test_codex_completed_answer_excludes_streamed_commentary():
    class Bridge:
        async def stream_turn(self, prompt, **kwargs):
            yield 'delta', '我将先阅读资料。'
            yield 'delta', '最终讲解'
            yield 'final', '最终讲解'
    assert asyncio.run(CodexClient(Bridge(), model='gpt-6-astra').complete('讲解', '原文')) == '最终讲解'


def test_bridge_uses_local_image_and_authoritative_final(monkeypatch):
    bridge = CodexBridge()
    calls = []
    async def start():
        pass
    async def request(method, params):
        calls.append((method, params))
        if method == 'thread/start':
            return {'thread': {'id': 't', 'ephemeral': True}}
        if method == 'turn/start':
            for event in [
                {'method': 'item/completed', 'params': {'threadId': 't', 'item': {'type': 'agentMessage', 'phase': 'commentary', 'text': '处理过程'}}},
                {'method': 'item/completed', 'params': {'threadId': 't', 'item': {'type': 'agentMessage', 'phase': 'final_answer', 'text': '正确讲解'}}},
                {'method': 'turn/completed', 'params': {'threadId': 't', 'turn': {'status': 'completed'}}}]:
                await bridge.events.put(event)
        return {}
    monkeypatch.setattr(bridge, 'start', start)
    monkeypatch.setattr(bridge, 'request', request)
    async def run():
        return [item async for item in bridge.stream_turn('识图', image_paths=['C:/original.png'])]
    assert asyncio.run(run())[-1] == ('final', '正确讲解')
    assert calls[1][1]['input'][1] == {'type': 'localImage', 'path': 'C:/original.png'}
    assert calls[0][1]['ephemeral'] is True


def test_generation_uses_original_and_records_actual_model(tmp_path, monkeypatch):
    project = create_project('完整原文件')
    source_id = uid()
    path = tmp_path / 'original.txt'
    path.write_text('原文中的关键边界条件：无阻尼。', encoding='utf-8')
    execute('INSERT INTO sources(id,project_id,name,kind,original_path,segments_json,status,error,created_at) VALUES (?,?,?,?,?,?,?,?,?)',
            (source_id, project['id'], 'original.txt', '.txt', str(path), j([dict(locator='preview', text='错误预览')]), 'ready', None, now()))
    class Client:
        model = 'gpt-6-astra'
        effort = 'high'
        async def complete(self, system, user):
            assert '原文中的关键边界条件' in user and '错误预览' not in user
            return '# 振子\n无阻尼。 [original.txt｜行 1-1]\n\n## 系统性总结\n理想模型的边界条件是无阻尼。'
    monkeypatch.setattr(server, 'client_for', lambda *args, **kwargs: Client())
    value = server.GenerateIn(project_id=project['id'], source_ids=[source_id], provider='codex', format='docx', review_mode='express')
    job_id = create_job(project['id'], 'generate', value.model_dump())
    asyncio.run(server._generate_job(job_id, value))
    assert one('SELECT * FROM jobs WHERE id=?', (job_id,))['status'] == 'completed'
    metadata = json.loads(one("SELECT * FROM artifacts WHERE kind='explanation'")['metadata_json'])
    assert metadata['model'] == 'gpt-6-astra' and metadata['source_reading'][0]['method'] == 'original_file'


def test_quality_recommendation_never_falls_back_to_luna():
    models = [dict(id='gpt-6-luna', name='Luna', efforts=['high']),
              dict(id='gpt-6-astra', name='Astra', efforts=['low', 'high'])]
    assert model_catalog.recommended_generation(models)['model'] == 'gpt-6-astra'
    with pytest.raises(ValueError):
        model_catalog.recommended_generation(models[:1])


def test_endpoint_preserves_selected_dialogue_model(monkeypatch):
    project = create_project('模型策略')
    captured = []
    async def catalog():
        return {'models': [dict(id='gpt-6-astra', name='Astra', efforts=['high'])]}
    async def validate(model, effort):
        assert (model, effort) == ('gpt-6-astra', 'high')
    async def job(job_id, value):
        captured.append(value.model)
    monkeypatch.setattr(model_catalog, 'get_models', catalog)
    monkeypatch.setattr(model_catalog, 'validate_effort', validate)
    monkeypatch.setattr(server, 'source_records', lambda *args, **kwargs: [{}])
    monkeypatch.setattr(server, '_generate_job', job)
    async def run():
        result = await server.generate(server.GenerateIn(project_id=project['id'], source_ids=['s'],
                                       provider='codex', model='gpt-6-luna', generation_quality='quality'))
        await asyncio.sleep(0)
        return result
    assert asyncio.run(run())['model'] == 'gpt-6-luna'
    assert captured == ['gpt-6-luna']


def test_403_gives_browser_import_action(monkeypatch):
    monkeypatch.setattr(urls, 'validate_public_url', lambda url: url)
    real_client = httpx.Client
    transport = httpx.MockTransport(lambda request: httpx.Response(403, request=request))
    monkeypatch.setattr(urls.httpx, 'Client', lambda **kwargs: real_client(transport=transport, **kwargs))
    with pytest.raises(urls.BrowserImportRequired, match='浏览器导入'):
        urls.download_public_resource('https://zhuanlan.zhihu.com/p/718779138', 'test')


def test_browser_page_preserves_full_body_formula_and_url(monkeypatch):
    monkeypatch.setattr(server, 'validate_public_url', lambda url: url)
    project = create_project('浏览器')
    with TestClient(server.app) as client:
        response = client.post('/api/import/browser-page', json=dict(project_id=project['id'],
            url='https://example.org/article', title='文章', html='<article><p>' + '完整文章内容。' * 100 + '</p><img alt="E=mc^2"><p>末尾细节</p><script>bad()</script></article>'))
        assert response.status_code == 200
        source = one('SELECT * FROM sources WHERE id=?', (response.json()['source_id'],))
        original = Path(source['original_path']).read_text(encoding='utf-8')
        assert 'bad()' not in original and 'https://example.org/article' in original and '末尾细节' in original
        extracted = '\n'.join(part['text'] for part in parse_file(Path(source['original_path'])))
        assert 'E=mc^2' in extracted
        assert client.post('/api/import/browser-page', json=dict(project_id=project['id'],
            url='https://example.org/article', title='安全验证', html='<article>'+'验证页面'*100+'</article>')).status_code == 422
