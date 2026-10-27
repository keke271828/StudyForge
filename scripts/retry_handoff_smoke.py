"""Retry only generation in an existing isolated live learning smoke fixture."""
import asyncio
import json
import os
import sys
from pathlib import Path

from docx import Document
from studyforge import server
from studyforge.codex_bridge import bridge
from studyforge.db import one, j


async def main():
    root = Path(sys.argv[1]).resolve()
    build = Path(__file__).resolve().parents[1] / 'build'
    if not root.is_relative_to(build) or not root.name.startswith('learning-live-'):
        raise ValueError('Only isolated live smoke fixtures are allowed')
    os.environ['STUDYFORGE_DATA_DIR'] = str(root)
    job = one("SELECT * FROM jobs WHERE kind='generate' ORDER BY created_at DESC LIMIT 1")
    value = server.GenerateIn.model_validate(json.loads(job['payload_json']))
    print(j({'stage': 'retry_handoff', 'root': str(root), 'model': value.model}), flush=True)
    try:
        await server._generate_job(job['id'], value)
        result = one('SELECT * FROM jobs WHERE id=?', (job['id'],))
        assert result['status'] == 'completed', result['error']
        artifact = one("SELECT * FROM artifacts WHERE kind='explanation'")
        handoff = json.loads(artifact['metadata_json'])['mentor_handoff']
        assert handoff['status'] == 'completed' and handoff['mentor']['model'] == value.model
        assert handoff['mentor_thread_id'] == one('SELECT remote_thread_id FROM conversations WHERE id=?', (value.conversation_id,))['remote_thread_id']
        misconceptions = [item for item in handoff['packet']['understanding'] if item['status'] == 'misconception']
        assert misconceptions and any('冰淇淋' in evidence['quote'] for item in misconceptions for evidence in item['evidence'])
        document = Document(artifact['path'])
        text = '\n'.join(p.text for p in document.paragraphs) + '\n'.join(c.text for table in document.tables for row in table.rows for c in row.cells)
        assert '冰淇淋' in text and '外部效度' in text and document.tables and '系统性总结' in text
        (root / '导师交接自检.json').write_text(j(handoff), encoding='utf-8')
        print(j({'ok': True, 'model': value.model, 'handoff_verified': True, 'misconceptions': len(misconceptions),
                 'characters': len(text), 'tables': len(document.tables), 'artifact': artifact['path']}), flush=True)
    finally:
        await bridge.close()


if __name__ == '__main__':
    asyncio.run(main())
