"""Live proof pedagogy through production chat and generation, in isolated data.

Uses the configured DeepSeek credential without printing it. No extra model
review is requested; explicit checks and human inspection assess the outputs.
"""
from __future__ import annotations

import asyncio
import json
import os
import re
import tempfile
import time
from pathlib import Path
from unittest.mock import patch

root = Path(__file__).resolve().parents[1]
target = Path(tempfile.mkdtemp(prefix='proof-teaching-', dir=root / 'build'))
os.environ['STUDYFORGE_DATA_DIR'] = str(target / 'data')

from docx import Document
from studyforge import server
from studyforge.ai import save_settings
from studyforge.codex_bridge import CodexBridge
from studyforge.db import create_job, create_project, execute, j, many, now, one, uid

NOETHER = r'''这是一份有待核对的诺特定理证明草稿。
实际轨道使作用量 S=\int L(q,\dot q,t)dt 最小，所以 \delta S=0。
考虑 t 不变而 q_i 变成 q_i+\epsilon\xi_i 的连续对称变换。
利用欧拉—拉格朗日方程，\delta S=\epsilon\int d(\sum_i p_i\xi_i)/dt\,dt，其中 p_i=\partial L/\partial\dot q_i。
由于这个固定区间的积分为零，被积函数处处为零，所以 Q=\sum_i p_i\xi_i 守恒。
即使 \delta L/\epsilon=dB/dt，上述 Q 仍然不需要修改。
例如一维自由粒子 L=m\dot q^2/2，取 \xi=t 的伽利略推进，也应有 Q=m\dot q t 守恒。
场论中局域规范变换 A_\mu\to A_\mu+\partial_\mu\chi 对应电荷守恒，其证明与全局连续对称的第一定理完全一样。
草稿末尾称物理定律的本质只能是对称性，没有给出论证。
核对资料：Noether, Invariant Variation Problems，https://arxiv.org/abs/physics/0503066 。
'''

SCHWARZ = r'''这是实数向量的柯西—施瓦茨不等式证明草稿。
设 A=\sum_i a_i^2，B=\sum_i b_i^2，C=\sum_i a_i b_i。
从 \sum_i(a_i-tb_i)^2\geq 0 得到 A-2tC+t^2B\geq 0，所有实数 t 都成立。
所以这个二次式判别式非正，即 C^2\leq AB。
等号条件写成 a=tb，没有单独讨论 b=0。
'''


def source(project, name, content):
    path = target / name
    path.write_text(content, encoding='utf-8')
    source_id = uid()
    execute('INSERT INTO sources(id,project_id,name,kind,original_path,segments_json,status,created_at) VALUES (?,?,?,?,?,?,?,?)',
            (source_id, project['id'], name, '.txt', str(path), '[]', 'ready', now()))
    return source_id


async def chat(project, selected, text):
    started = time.perf_counter()
    response = await server.chat_stream(server.ChatIn(project_id=project['id'], source_ids=[selected],
        provider='deepseek', model=model, text=text))
    events = []
    async for item in response.body_iterator:
        event = json.loads(item)
        events.append(event)
        if event['type'] == 'status':
            print(event['value'], flush=True)
        elif event['type'] == 'tool':
            print(f"  {event['value']['tool']} · {event['value']['status']}", flush=True)
    failures = [item['value'] for item in events if item['type'] == 'error']
    if failures:
        raise RuntimeError(str(failures))
    return (next(item['value'] for item in events if item['type']=='conversation'),
            next(item['value'] for item in reversed(events) if item['type']=='replace'),
            round(time.perf_counter()-started, 1))


async def main():
    save_settings({'chat_provider': 'deepseek', 'deepseek_model': model, 'review_mode': 'express'})
    project = create_project('证明思路真实验收')
    noether = source(project, '诺特证明草稿.txt', NOETHER)
    print(f'验收目录：{target}', flush=True)
    print('1/3：真实导师讲解诺特定理，核对草稿与实际原论文网页。', flush=True)
    conversation, reply, seconds = await chat(project, noether,
        '请把原文件的证明讲透。我不懂为什么突然出现分部积分，也没看懂末尾的电荷推论。'
        '请实际读取资料里给出的原论文网页，核对第一、第二定理及论文年份，不推测作者心理。')
    (target/'noether-chat.md').write_text(reply, encoding='utf-8')
    print('2/3：同一对话生成详讲 DOCX，不增加独立审查。', flush=True)
    value = server.GenerateIn(project_id=project['id'], source_ids=[noether], conversation_id=conversation,
        provider='deepseek', model=model, format='docx', summary_level=5, exercises=False,
        review_mode='express', generation_quality='custom',
        instructions='延续刚刚的讲解，写给知道微积分但初学分析力学的学生。讲清证明和自由粒子例子，不布置习题。')
    job = create_job(project['id'], 'generate', value.model_dump())
    started = time.perf_counter()
    await server._generate_job(job, value)
    result = one('SELECT * FROM jobs WHERE id=?', (job,))
    if result['status'] != 'completed':
        raise RuntimeError(result['error'])
    generation_seconds = round(time.perf_counter()-started, 1)
    artifact = one("SELECT * FROM artifacts WHERE kind='explanation'")
    meta = json.loads(artifact['metadata_json'])
    lecture_message = one('SELECT m.content FROM generation_dialogue_outputs g JOIN messages m ON m.id=g.message_id WHERE g.job_id=?', (job,))['content']
    (target/'noether-lecture.md').write_text(lecture_message, encoding='utf-8')
    doc = Document(artifact['path'])
    text = '\n'.join(p.text for p in doc.paragraphs if p.text.strip())
    equations = len(doc.element.xpath('.//m:oMath'))
    print('3/3：纯数学证明，检查方法启发、零向量边界及不强行关联物理。', flush=True)
    schwarz = source(project, '柯西证明草稿.txt', SCHWARZ)
    _, math_reply, math_seconds = await chat(project, schwarz,
        '请讲透这份证明。我会展开多项式，但不懂为什么要构造含 t 的平方和；只围绕这个数学问题解释，不联网，不写历史。')
    (target/'schwarz-chat.md').write_text(math_reply, encoding='utf-8')
    tools = many('SELECT tool,status FROM agent_tool_events WHERE conversation_id=?', (conversation,))
    checks = {
        'same_dialogue': meta['conversation_id']==conversation and meta['dialogue_model']==model,
        'one_selected_artifact': len(many('SELECT * FROM artifacts'))==1 and artifact['format']=='docx',
        'no_independent_review': meta['reviews'][0]['passed'] is None and len(meta['reviews'])==1,
        'webpage_read': any(t['tool']=='sf_read_url' and t['status']=='completed' for t in tools),
        'word_math_native': equations>0,
        'source_quote_style': any(p.style.name=='Source Quote' for p in doc.paragraphs),
        'summary_last': '系统性总结' in '\n'.join(p.text for p in doc.paragraphs[-12:]),
        'no_assignments': not re.search(r'(?m)^#{1,6}\s*(?:课后练习|配套习题|思考题|作业)', lecture_message),
        'noether_corrections_present': all(w in reply + text for w in ('驻', '第一', '第二', '边界', '分部积分')),
        'method_explanation_present': any(w in math_reply for w in ('构造', '投影', '逼近')) and any(w in math_reply for w in ('想到', '为什么', '动机', '思路')),
        'zero_vector_case': bool(re.search(r'(?:B|b|\\mathbf\{b\}|\\boldsymbol\{b\})\s*(?:=|为|是|是一个)\s*(?:0|零)', math_reply)),
        'no_forced_physics': not re.search('诺特|拉格朗日|楞次|量子|牛顿', math_reply),
        'codex_disabled': not codex_attempts,
    }
    report = {'model': model, 'checks': checks, 'chat_seconds': seconds, 'generation_seconds': generation_seconds,
              'math_seconds': math_seconds, 'docx': artifact['path'], 'word_equations': equations,
              'tools': tools, 'automatic_checks_only': True}
    (target/'result.json').write_text(j(report), encoding='utf-8')
    print(j(report), flush=True)
    assert all(checks.values()), checks


model = os.environ.get('STUDYFORGE_SMOKE_MODEL', 'deepseek-flash')
codex_attempts = []


async def forbidden(*args, **kwargs):
    codex_attempts.append(True)
    raise AssertionError('DeepSeek 独立教学验收禁止调用 Codex')


if __name__ == '__main__':
    with patch.object(CodexBridge, 'start', forbidden), patch.object(server.model_catalog, 'get_models', forbidden), \
         patch.object(server.model_catalog, 'validate_effort', forbidden):
        asyncio.run(main())
