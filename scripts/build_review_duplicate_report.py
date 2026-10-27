"""Saved real quick reviews, a rejected general candidate and a bounded fix."""
import base64
import copy
import hashlib
from html import escape
import json
from pathlib import Path
import re
import shutil
from xml.etree import ElementTree as ET

from scripts.build_teaching_comparison import TEMPLATE,inline_script

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'build/review-duplicate-acceptance-20261021'
TITLE='快速核查与重复段落修订实际验收'
NOTES={
 'gauss':[
  '旧流程发现了“区域无电荷+无穷远趋零就零场”的范围歧义，把所讨论区域改为全空间并写入。主例和通量计算保留；这次核查仍未全面处理改写引文等问题，修订也未再次独立复核。',
  '普遍提供段落编号的候选返回passed=true、issues=[]，没有修正原有范围歧义，且本次更慢。未采用这个普遍方案。没有报错不代表没有错误；这个视图保留漏检的真实交付。'],
 'rate':[
  '旧流程发现总结中“只有单步基元反应才能得到该形式”的必要性过强，逐字局部修正成功。但未报告正文“B表观零级”等待核对内容；不能把这一项修正视作全篇正确。',
  '普遍编号候选发现同一总结问题，仍选择逐字before，没有使用编号；本次更慢，也没有稳定扩大问题覆盖。未采用。正确的单位、0.03计算及其他未核查内容均原样保留。'],
 'duplicate':[
  '旧流程通过把未改变的标题包含在before里，成功区分正文和总结，两处都写入；不能声称旧流程本次失败。修订说“N必须且只需≥max”时，应明确指从n>N保证两组下标条件，而不是断言N1、N2一定是实际估计的最小门槛。',
  '采用的窄方案只对重复普通段落给出两个定位编号，模型实际使用编号将正文与总结两处建议一起写入，正确的N1=5、N2=8、N=10例子保留。表述回到任意N≥max都可用；“最小统一／充分门槛”仍应按下标条件理解。不是完整证明认证，每臂只有一次。'],
}


def main():
    OUT.mkdir(parents=True,exist_ok=True)
    from studyforge.dialogue_scope import CONTEXT_BOUNDARY
    from studyforge.deepseek_protocol import thinking_parameters
    from docx import Document
    groups=[];table=[];matched=0;word_files=[]
    configs=[('review-passages-real-20261021-1',['gauss','rate'],('baseline','registered'),False),
             ('review-duplicate-real-20261021-1',['duplicate'],('baseline','narrow'),True)]
    for dirname,ids,phases,adopted in configs:
        folder=ROOT/'build'/dirname
        cases=json.loads((folder/'inputs.json').read_text(encoding='utf-8'))['cases']
        for case in cases:
            if case['id'] not in ids:continue
            rows=[];summaries=[]
            for number,phase in enumerate(phases):
                record=next(r for r in json.loads((folder/(phase+'-results.json')).read_text(encoding='utf-8')) if r['case']==case['id'])
                assert record['status']=='completed' and record['model_requests']==1
                raw=(folder/record['answer_file']).read_bytes();text=raw.decode('utf-8').replace('\r\n','\n')
                assert hashlib.sha256(text.encode()).hexdigest()==record['answer_sha256']
                requests=json.loads((folder/(case['id']+'-'+phase+'-records.json')).read_text(encoding='utf-8'))
                assert len(requests)==1 and requests[0]['json_mode']
                body={'model':'deepseek-flash','messages':[{'role':'system','content':requests[0]['system']+CONTEXT_BOUNDARY},
                    {'role':'user','content':requests[0]['user']}],'stream':True,**thinking_parameters('high'),
                    'response_format':{'type':'json_object'},'stream_options':{'include_usage':True}}
                trace=record['http_attempts'][0]
                assert hashlib.sha256(json.dumps(body,sort_keys=True,ensure_ascii=False).encode()).hexdigest()==trace['body_sha256'];matched+=1
                label='旧快速核查' if number==0 else '重复段落编号（已采用）' if adopted else '普遍提供编号（未采用）'
                name={'gauss':'高斯','rate':'速率','duplicate':'重复门槛'}[case['id']]+'_'+label+'.docx'
                shutil.copyfile(folder/record['docx_file'],OUT/name)
                word=Document(OUT/name);assert word.paragraphs
                word_files.append({'file':name,'sha256':hashlib.sha256((OUT/name).read_bytes()).hexdigest(),
                    'paragraphs':len(word.paragraphs),'tables':len(word.tables),'structurally_reopened':True})
                if case['id']=='duplicate':
                    words='\n'.join(p.text for p in word.paragraphs)
                    assert 'N1=5、N2=8' in words and 'N=10' in words and '任何更大的N都不能使用' not in words
                    assert 'passage_id' not in words and '尚未再次独立复核' in words
                rows.append({'case':case['id'],'arm':phase,'label':label,'text':raw.decode('utf-8'),
                    'sha256':hashlib.sha256(raw).hexdigest(),'review':NOTES[case['id']][number],
                    'seconds':round(record['seconds'],2),'input_chars':trace['input_characters'],'output_chars':record['characters'],
                    'response_model':trace['returned_model'],'usage':[trace['usage']],
                    'request':{'actual_public_request_and_response':requests,'reports':record['reports'],
                        'actual_http_body_sha256':trace['body_sha256'],'draft_sha256':record['draft_sha256']}})
                summaries.append(f"{record['seconds']:.2f}秒／{len(record['reports'][0].get('issues',[]))}项报告问题")
            groups.append({'round':'窄方案：已采用' if adopted else '普遍方案：未采用','case':case['id'],'title':case['title'],
                'question':'对同一份待审草稿作一次快速核查，保留学生原要求：'+case['requirements'],
                'source':case['source']+'\n\n【同一份待审草稿，含刻意保留的问题】\n'+case['draft']+
                    '\n\n核对范围：\n'+'\n'.join(case['checks']),'rows':rows})
            table.append('<tr><td>'+escape(case['title'])+'</td><td>'+summaries[0]+'</td><td>'+summaries[1]+'</td><td>'+('仅重复定位采用' if adopted else '普遍编号拒绝')+'</td></tr>')
    assert matched==6 and len(groups)==3 and len(word_files)==6
    # Open on the adopted feature; the rejected experiments remain visible.
    groups.sort(key=lambda group:group['case']!='duplicate')
    unchanged=json.loads((ROOT/'build/review-passages-real-20261021-1/narrow-unchanged-request-checks.json').read_text(encoding='utf-8'))
    replay=json.loads((ROOT/'build/review-duplicate-real-20261021-1/narrow-final-code-replay.json').read_text(encoding='utf-8'))
    assert unchanged['passed'] and replay['passed']
    perf=json.loads((ROOT/'build/review-duplicate-local-overhead-20261021.json').read_text(encoding='utf-8'))
    tests=ET.parse(ROOT/'build/review-duplicate-final-regression-20261021.xml').getroot().findall('.//testcase')
    assert len(tests)==452 and all(t.find(k) is None for t in tests for k in ('failure','error','skipped'))
    word_links=' · '.join('<a href="'+escape(r['file'],quote=True)+'">'+escape(r['file'])+'</a>' for r in word_files)
    intro='''<header><div class="eyebrow">STUDYFORGE · QUICK REVIEW · 2026-10-21</div>
<h1>让正文和总结中的重复修订准确落地</h1><p>只为逐字定位存在歧义的重复普通段落提供编号；用真实结果拒绝无收益的普遍编号方案。</p>
<p><a href="../../验收最新版.vbs">启动最新版源码</a> · <a href="../../docs/快速核查与重复段落修订验收.md">详细验收</a>。本轮未打包。</p>
<div class="notice"><strong>已采用的是重复段落的定位备用路径，未更换导师共同教学规则，也未增加审查轮次。</strong>
本页6份视图是“草稿经过快速核查后的实际交付”，不是6份新撰写的讲解。普遍方案更慢且有漏检，未采用；重复段落组旧、新两版均完成两处修订。所有原始结果及其不足保留。</div></header>'''
    evidence='''<details class="prompt"><summary>查看采用范围、实际结果、Word与验证限制</summary>
<h2>采用：只帮助定位重复普通段落</h2><p>当正文和总结出现完全相同的普通段落时，单独抄这段before可能无法唯一定位。现在只对这种重复提供绑定当前草稿和精确位置的编号，模型可分别指定段落，原有逐字定位仍可用。无重复时不提供编号及附加规则，两份旧样本的系统提示和用户输入已逐字核对保持相同，零新增外部请求。</p>
<p>预览只是定位开头，完整待审讲解始终提供。原文引用、来源、网址、代码、列表、表格及已有核查提示不走这条编号路径。草稿改变、记录损坏、未知编号、冲突或一组修订某处无法应用时保留疑点；不做公式等价或模糊匹配。</p>
<p>登记最多24个位置、定位JSON最多3200字符；同一段的全部出现位置都能放下时才登记，超过8处的一组不登记。修订依然最多8处，同一问题整组应用。成功记录registered_passage与精确Unicode位置，不冒充模型逐字抄过原句；应用建议不等于已重新核实正确。</p>
<h2>没有采用普遍编号方案</h2><p>先在高斯和速率两份真实旧讲解上作对照。普遍提供编号让两组都更慢；速率组仍用逐字before修同一项，未稳定扩大问题覆盖；高斯组候选反而返回passed=true，漏掉旧流程发现的条件范围问题。因此没有把这个方案放进每次核查。</p>
<div class="report-table"><table><thead><tr><th>材料</th><th>旧流程</th><th>候选／窄方案</th><th>决定</th></tr></thead><tbody>__ROWS__</tbody></table></div>
<p>重复门槛组是明确构造的含错样例：同一句“必须且只能取最大值，更大N不能使用”出现在正文和总结，旁边保留N=10同样成立的正确例子。旧流程本次利用带标题的逐字锚点也完成了修订；新流程实际使用两个编号，保留例子并修订两处。没有声称旧流程必定失败、编号能提高数学判断能力或整体提速。</p>
<p>6次真实请求均为deepseek-flash／high，每个视图一次快速核查，无工具、联网、对话历史或再次独立复核。每组原文、草稿、要求、风格、模式和详略固定。实际HTTP请求指纹6份逐一核对；公开请求、原始JSON和交付保留，未保存私有思考正文。最终源码回放同一保存响应，核对窄方案输入和交付相同，零新增模型请求。</p>
<h2>教学上仍有不足</h2><p>高斯普遍候选漏检，是本次直接看到的失败。速率两版都只报告总结必要性问题，正文其他待核对内容未因此解决。重复门槛的“最小统一门槛”也应限定为从n>N保证两组下标条件，不能误认为N1和N2一定是实际估计的最小门槛。问题说明、修订和核查通过标记都不是金标准。</p>
<p>本机唯一定位只能避免改错位置，不能证明建议本身正确。默认讲解风格、学习模式、全文原文输入、三种其他审查档位及习题规则继续保留；没有给学生增加练习或掌握分数。</p>
<h2>验证与打开</h2><p>452项相关检查通过。其中两项受控完整原生生成→快速核查→Word及公开聊天流程，验证正文和总结同时写入，以及第二个编号无效时整组保留疑点。每项仍只有一次撰写与一次原有快速核查，无新增习题或画像分数。另验证原有格式别名、严格匹配、引用／来源保护、并行、暂停、书库输入与指南。</p>
<p>本机定位准备的样本中位耗时约0.37–6.76毫秒，临时分配约29KB–544KB；只是约2.3千／4.7万字符的4份合成草稿，不代表全应用RAM或整项生成耗时。没有把额外定位数据说成免费：重复场景会增加少量输入；没有重复的两份真实旧样本保持原请求逐字不变。</p>
<p>__WORDS__</p><p>6份Word均实际导出并重读结构；重复组的原数值例子和修订说明保留。它们是核查对照，不是定稿教材；本轮没有在Microsoft Word中做分页与视觉验收。页面公式、链接、电脑和390像素宽度另经浏览器核对。未新增工作台交互，不把报告预览当作物理手机或全应用UI验收。</p>
<p>双击顶部入口使用最新版源码，旧安装EXE不含本轮修改。已有历史成果不会自动重写；本轮未打包。</p></details>'''
    page=re.sub(r'<header>.*?</header>',lambda _:intro,TEMPLATE,flags=re.S).replace('<main>','<main>'+evidence.replace('__ROWS__',''.join(table)).replace('__WORDS__',word_links),1)
    page=page.replace('<title>StudyForge · 讲解质量真实对照</title>',f'<title>{TITLE} · StudyForge</title>')
    page=page.replace('grid-template-columns:repeat(3,minmax(0,1fr))','grid-template-columns:repeat(2,minmax(0,1fr))')
    page=page.replace('</style>','.report-table{overflow-x:auto}table{border-collapse:collapse;width:100%;font-size:14px}td,th{padding:10px;border:1px solid #d2e0d6;text-align:left}a{color:#236b5b;overflow-wrap:anywhere}.prompt{overflow-wrap:anywhere}.lesson-math.display-math{display:block;overflow-x:auto;overflow-y:hidden;padding:6px 0}.katex-display{overflow:visible}</style>',1)
    css=(ROOT/'web/vendor/katex.css').read_text(encoding='utf-8')
    css=re.sub(r'url\(fonts/([^)]*\.woff2)\)',lambda m:'url(data:font/woff2;base64,'+base64.b64encode((ROOT/'web/vendor/fonts'/m[1]).read_bytes()).decode()+')',css)
    css=re.sub(r',url\(fonts/[^)]*\) format\("(?:woff|truetype)"\)','',css)
    page=page.replace('__KATEX_CSS__',css).replace('__PAYLOAD__',json.dumps({'groups':groups,'count':6,'failures':[]},ensure_ascii=False).replace('<','\\u003c'))
    page=page.replace('__LIBRARIES__','\n'.join(inline_script((ROOT/'web/vendor'/name).read_text(encoding='utf-8')) for name in ('marked.js','purify.js','katex.js')))
    page=page.replace('__RENDERER__',inline_script((ROOT/'web/lesson_render.js').read_text(encoding='utf-8')))
    page=page.replace('份原始回答','份实际核查交付').replace('查看本次提供的完整合成原资料','查看完整原资料、同一待审草稿与核对范围')
    target=OUT/(TITLE+'.html');target.write_text(page,encoding='utf-8')
    result={'answers':6,'groups':3,'report_table_rows':[3],'program_checks':len(tests),'real_model_requests':6,
        'request_fingerprints_matched':matched,'broad_registry_adopted':False,'duplicate_registry_adopted':True,
        'unchanged_saved_requests':len(unchanged['cases']),'final_code_saved_response_replay':replay['passed'],
        'replay_new_external_requests':0,'word_files':word_files,'local_preparation':perf,'package_created':False,
        'output_sha256':hashlib.sha256(target.read_bytes()).hexdigest()}
    (OUT/'manifest.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({'path':str(target),'views':6,'checks':len(tests),'requests':matched,'word_files':len(word_files)},ensure_ascii=True))


if __name__=='__main__':main()
