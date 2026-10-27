from __future__ import annotations

import ast
import hashlib
import json
import re
import subprocess
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path
from datetime import datetime
from typing import Callable

from .config import artifacts_dir
from .teaching import summary_section
from .deliverables import cards, bundle_files, json_object, static_check
from .document_render import (blocks, inline_parts, math_image, validate_document_text, word_math,
                              rich_html, owned_image, image_size, EXPLANATION_COLOR, SOURCE_COLOR, EXPLANATION_FONT, SOURCE_FONT)

Exporter = Callable[[str, str, Path], None]
EXPORTERS: dict[str, Exporter] = {}


def register(*formats: str):
    def decorate(func: Exporter) -> Exporter:
        for fmt in formats:
            EXPORTERS[fmt] = func
        return func
    return decorate


def _sections(content: str) -> list[tuple[str, str]]:
    blocks: list[tuple[str, str]] = []
    title = "概述"
    lines: list[str] = []
    fence = None
    for line in content.splitlines():
        marker = re.match(r'^\s*(`{3,}|~{3,})', line)
        was_fenced = fence is not None
        if marker:
            if fence is None: fence = marker[1]
            elif marker[1][0] == fence[0] and len(marker[1]) >= len(fence): fence = None
        if not was_fenced and not marker and (line.startswith("# ") or line.startswith("## ")):
            if lines:
                blocks.append((title, "\n".join(lines).strip()))
            title = line.lstrip("# ").strip()
            lines = []
        else:
            lines.append(line)
    if lines:
        blocks.append((title, "\n".join(lines).strip()))
    return blocks or [("内容", content)]


@register("docx", "word")
def export_docx(content: str, title: str, path: Path) -> None:
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.enum.style import WD_STYLE_TYPE
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    doc = Document()
    doc.add_heading(title, 0)
    for name in ('Normal', 'Title', 'Heading 1', 'Heading 2', 'Heading 3', 'List Bullet', 'List Number'):
        style = doc.styles[name]
        style.font.name = EXPLANATION_FONT
        style.font.color.rgb = RGBColor.from_string(EXPLANATION_COLOR)
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), EXPLANATION_FONT)
    normal = doc.styles['Normal']
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.space_after = Pt(7)
    quoted = doc.styles.add_style('Source Quote', WD_STYLE_TYPE.PARAGRAPH)
    quoted.base_style = normal
    quoted.font.name = 'Times New Roman'
    quoted.font.color.rgb = RGBColor.from_string(SOURCE_COLOR)
    quoted.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), SOURCE_FONT)
    quoted.paragraph_format.left_indent = Cm(.4)
    quoted.paragraph_format.right_indent = Cm(.3)
    quoted.paragraph_format.space_before = Pt(8)
    quoted.paragraph_format.space_after = Pt(12)
    properties = quoted.element.get_or_add_pPr()
    shading = OxmlElement('w:shd'); shading.set(qn('w:fill'), 'F3F5F7'); properties.append(shading)
    borders = OxmlElement('w:pBdr')
    border = OxmlElement('w:left')
    for key, value in {'val': 'single', 'sz': '12', 'space': '8', 'color': '9BAAB8'}.items():
        border.set(qn('w:' + key), value)
    borders.append(border); properties.append(borders)
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Cm(2)
        section.left_margin = section.right_margin = Cm(2.2)

    def write(paragraph, text):
        for kind, part in inline_parts(text):
            if kind == 'math':
                try:
                    paragraph._p.append(word_math(part))
                except Exception as exc:
                    raise ValueError(f'无法排版 Word 公式：{part[:100]}') from exc
            else:
                run = paragraph.add_run(part)
                run.bold = kind == 'bold'
                if kind == 'code':
                    run.font.name = 'Consolas'

    numbering = doc.part.numbering_part.element
    base_number_id = doc.styles['List Number'].element.pPr.numPr.numId.val
    number_abstract = numbering.num_having_numId(base_number_id).abstractNumId.val
    active_number_id = None
    for kind, value in blocks(content, preserve_list_numbers=True):
        if kind != 'number':
            active_number_id = None
        image = owned_image(value) if kind == 'paragraph' else None
        if image:
            width, height = image_size(image[0], 16.5, 21)
            paragraph = doc.add_paragraph(); paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.add_run().add_picture(str(image[0]), width=Cm(width), height=Cm(height))
            if image[1]: write(doc.add_paragraph(), image[1])
        elif kind == 'heading':
            write(doc.add_heading('', level=min(value[0], 3)), value[1])
        elif kind in ('math','quote_math'):
            paragraph = doc.add_paragraph(style='Source Quote' if kind=='quote_math' else None)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph._p.append(word_math(value))
        elif kind == 'table':
            columns = max(len(row) for row in value)
            table = doc.add_table(rows=len(value), cols=columns)
            table.style = 'Table Grid'
            # New tables are rectangular and unmerged. python-docx rebuilds
            # its entire cell grid for every table.cell() lookup; reuse that
            # grid within this table while preserving the same XML nodes.
            cells = table._cells
            header = OxmlElement('w:tblHeader')
            table.rows[0]._tr.get_or_add_trPr().append(header)
            for row_index, row in enumerate(value):
                for column, cell in enumerate(row):
                    paragraph = cells[row_index * columns + column].paragraphs[0]
                    write(paragraph, cell)
                    if row_index == 0:
                        paragraph.paragraph_format.keep_with_next = True
        elif kind == 'code':
            paragraph = doc.add_paragraph()
            run = paragraph.add_run(value)
            run.font.name = 'Consolas'; run.font.size = Pt(9)
        elif kind == 'quote':
            write(doc.add_paragraph(style='Source Quote'), value)
        else:
            style = {'bullet': 'List Bullet', 'number': 'List Number'}.get(kind)
            paragraph = doc.add_paragraph(style=style)
            if kind == 'number':
                match = re.match(r'^(\d+)[.)]\s+(.*)$', value)
                if active_number_id is None:
                    number = numbering.add_num(number_abstract)
                    number.add_lvlOverride(0).add_startOverride(int(match[1]))
                    active_number_id = number.numId
                properties = paragraph._p.get_or_add_pPr().get_or_add_numPr()
                properties.get_or_add_ilvl().val = 0
                properties.get_or_add_numId().val = active_number_id
                value = match[2]
            write(paragraph, value)
    doc.save(path)


@register("pdf")
def export_pdf(content: str, title: str, path: Path) -> None:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image, Table, TableStyle, Preformatted
    font_path = Path("C:/Windows/Fonts/simhei.ttf")
    if font_path.exists() and "StudyForgeChinese" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("StudyForgeChinese", str(font_path)))
    if not font_path.exists() and 'STSong-Light' not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'))
    font = "StudyForgeChinese" if font_path.exists() else 'STSong-Light'
    source_font_path = Path('C:/Windows/Fonts/simsun.ttc')
    source_font = 'StudyForgeSourceChinese' if source_font_path.exists() else 'STSong-Light'
    if source_font_path.exists() and source_font not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(source_font, str(source_font_path), subfontIndex=0))
    elif not source_font_path.exists() and source_font not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(UnicodeCIDFont(source_font))
    styles = {
        "title": ParagraphStyle("sf-title", fontName=font, fontSize=18, leading=26, spaceAfter=16, keepWithNext=True),
        "heading": ParagraphStyle("sf-heading", fontName=font, fontSize=14, leading=21, spaceBefore=12, spaceAfter=7, keepWithNext=True),
        "body": ParagraphStyle("sf-body", fontName=font, fontSize=10.5, leading=18, spaceAfter=7,
                               wordWrap='CJK', textColor=colors.HexColor('#' + EXPLANATION_COLOR), alignment=TA_LEFT),
    }
    styles['mixed'] = ParagraphStyle('sf-mixed', parent=styles['body'], wordWrap='LTR')
    styles['quote'] = ParagraphStyle('sf-quote', parent=styles['body'], fontName=source_font,
                                     textColor=colors.HexColor('#' + SOURCE_COLOR),
                                     backColor=colors.HexColor('#F3F5F7'), borderPadding=9,
                                     borderWidth=.5, borderColor=colors.HexColor('#CCD5DE'),
                                     leftIndent=12, rightIndent=12, spaceBefore=8, spaceAfter=12)

    def rich(text, max_width=450, color=EXPLANATION_COLOR):
        result = []
        for kind, part in inline_parts(text):
            if kind == 'math':
                from PIL import Image as PILImage
                image = math_image(part, color=color)
                with PILImage.open(image) as bitmap:
                    width, height = bitmap.size
                height_pt = height * .21
                width_pt = width * .21
                if width_pt > max_width:
                    height_pt *= max_width / width_pt
                    width_pt = max_width
                result.append(f'<img src="{image.as_posix()}" width="{width_pt:.2f}" height="{height_pt}" valign="middle"/>')
            else:
                escaped = _xml_escape(part)
                result.append(f'<b>{escaped}</b>' if kind == 'bold' else escaped)
        return ''.join(result)

    def paragraph(text, style, max_width=450, breaks=False, prefix=''):
        markup = prefix + rich(text, max_width,color=SOURCE_COLOR if style is styles['quote'] else EXPLANATION_COLOR)
        if breaks: markup=markup.replace('\n','<br/>')
        height = max([float(x) for x in re.findall(r'height="([\d.]+)"', markup)] or [0])
        sized = ParagraphStyle('sf-equation-line', parent=style, leading=max(style.leading,height+5))
        return Paragraph(markup,sized)

    story = [Paragraph(_xml_escape(title), styles["title"])]
    list_number = None
    for kind, value in blocks(content, preserve_list_numbers=True):
        if kind != 'number':
            list_number = None
        picture = owned_image(value) if kind == 'paragraph' else None
        if picture:
            width, height = image_size(picture[0], 490, 600)
            story.extend([Image(str(picture[0]), width=width, height=height), Spacer(1, 9)])
            if picture[1]: story.append(paragraph(picture[1], styles['body']))
        elif kind == 'heading':
            story.append(paragraph(value[1], styles['heading']))
        elif kind=='quote_math':
            story.append(paragraph('$'+re.sub(r'\s+',' ',value)+'$',styles['quote']))
        elif kind == 'math':
            image = Image(str(math_image(value)))
            scale = min(.24, 480 / image.imageWidth, 150 / image.imageHeight)
            if scale < .12: raise ValueError('PDF公式过大，缩放后难以阅读；请分行或选择Word格式')
            image.drawWidth = image.imageWidth * scale
            image.drawHeight = image.imageHeight * scale
            story.extend([image, Spacer(1, 9)])
        elif kind == 'table':
            columns = max(len(row) for row in value)
            data = []
            for row in value:
                cells = []
                for cell in row:
                    cells.append(paragraph(cell, styles['mixed'] if any(k=='math' for k,_ in inline_parts(cell)) else styles['body'], 490/columns-14))
                data.append(cells + [''] * (columns-len(row)))
            table = Table(data, colWidths=[490 / columns] * columns, repeatRows=1)
            table.setStyle(TableStyle([('GRID', (0, 0), (-1, -1), .4, colors.HexColor('#d0ded9')),
                                       ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#edf5f1')),
                                       ('VALIGN', (0, 0), (-1, -1), 'TOP')]))
            story.extend([table, Spacer(1, 10)])
        elif kind == 'code':
            # Escape code before passing it through the reportlab paragraph parser.
            story.append(Paragraph(_xml_escape(value).replace('\n', '<br/>'), styles['body']))
        elif kind == 'quote':
            story.append(paragraph(value, styles['quote'],breaks=True))
        else:
            prefix = '• ' if kind == 'bullet' else ''
            if kind == 'number':
                match = re.match(r'^(\d+)[.)]\s+(.*)$', value)
                list_number = int(match[1]) if list_number is None else list_number+1
                prefix = f'{list_number}. '
                value = match[2]
            story.append(paragraph(value, styles['mixed'] if any(k=='math' for k,_ in inline_parts(value)) else styles['body'],prefix=prefix))
    SimpleDocTemplate(str(path), pagesize=A4, leftMargin=48, rightMargin=48,
                      topMargin=48, bottomMargin=48).build(story)


def _xml_escape(value: str) -> str:
    return value.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


@register("pptx", "ppt")
def export_pptx(content: str, title: str, path: Path) -> None:
    from .slide_export import export_slides
    return export_slides(content, title, path)


@register("mm", "mindmap")
def export_mindmap(content: str, title: str, path: Path) -> None:
    root = ET.Element("map", version="1.0.1")
    top = ET.SubElement(root, "node", TEXT=title)
    stack = [(0,top)]
    for kind,value in blocks(content):
        if kind=='heading':
            depth,label=value
            while stack[-1][0]>=depth: stack.pop()
            child=ET.SubElement(stack[-1][1],'node',TEXT=label)
            stack.append((depth,child));continue
        if kind=='table': value='\n'.join(' | '.join(row) for row in value)
        if value:
            node=ET.SubElement(stack[-1][1],'node',TEXT=value,
                               COLOR='#'+(SOURCE_COLOR if kind in ('quote','quote_math') else EXPLANATION_COLOR))
            ET.SubElement(node,'font',NAME=SOURCE_FONT if kind in ('quote','quote_math') else EXPLANATION_FONT,SIZE='14')
    ET.ElementTree(root).write(path, encoding="utf-8", xml_declaration=True)


@register("apkg", "anki")
def export_anki(content: str, title: str, path: Path) -> None:
    import genanki
    identifier = int(hashlib.sha1(title.encode()).hexdigest()[:8], 16)
    model = genanki.Model(1607392319, "StudyForge QA", fields=[{"name": "Question"}, {"name": "Answer"}],
                          templates=[{"name": "Card 1", "qfmt": "{{Question}}", "afmt": "{{FrontSide}}<hr>{{Answer}}"}],
                          css=f'.card {{font-family:Microsoft YaHei,Arial,sans-serif;color:#{EXPLANATION_COLOR};font-size:20px;line-height:1.65;text-align:left;background:white;padding:20px;}}')
    deck = genanki.Deck(identifier, title)
    if not content.lstrip().startswith('{'):
        raise ValueError('Anki需要聚焦单一问题的JSON卡片：{"cards":[{"question":"问题","answer":"答案","source":"来源"}]}；不能将整章直接作为背面')
    questions = set()
    payload=json_object(content)
    for index,card in enumerate(cards(payload, max_count=5000, max_question=10000 if payload.get('purpose')=='student_exercises' else 400)):
        if card['question'] in questions: raise ValueError('卡片题面重复，请合并或区分条件后再导出')
        questions.add(card['question'])
        answer = card['answer'] + ('\n\n来源：' + card['source'] if card['source'] else '\n\n来源待核对。')
        warning=payload['cards'][index].get('_source_warning')
        if isinstance(warning,str): answer+='\n\n来源定位需核对：'+warning[:2000]
        deck.add_note(genanki.Note(model=model, guid=genanki.guid_for(title, card['question']),
                                  fields=[rich_html(card['question']), rich_html(answer)]))
    genanki.Package(deck).write_to_file(str(path))


def _fenced_code(content: str, language: str) -> str:
    aliases = {"py": "python", "m": "matlab", "wl": "wolfram", "js": "javascript",
               "ts": "typescript", "cpp": "c++"}
    wanted = {language, aliases.get(language, language)}
    content = _runnable_content(content)
    blocks = re.findall(r"```([A-Za-z0-9_+-]*)\s*\n(.*?)\n```", content, re.S)
    selected = [body for label, body in blocks if label.lower() in wanted]
    if not selected:
        raise ValueError(f"讲解没有可导出的 {language} 代码块；请先生成该格式的代码内容")
    return "\n\n".join(selected)


def _runnable_content(content):
    from .teaching import headings
    matches = [start for start, _, label in headings(content)
               if re.fullmatch(r'(?:可运行模型或示例(?:\s*[/｜|]\s*Runnable model or example)?|Runnable model or example)', label, re.I)]
    return content[matches[-1]:] if matches else content


def _execution_notice(content):
    heading = _runnable_content(content).partition('\n')[0].strip().lstrip('#').strip()
    zh = '运行状态：仅导出；尚未执行验证，依赖、数值结果和适用条件需在目标环境核对。'
    en = 'Execution status: exported only; not executed. Verify dependencies, numerical results and applicability in the target environment.'
    if re.fullmatch('Runnable model or example', heading, re.I): return en
    if re.fullmatch(r'可运行模型或示例\s*[/｜|]\s*Runnable model or example', heading, re.I): return zh+' / '+en
    return zh


@register("py", "js", "ts", "cpp", "m", "wl")
def export_code(content: str, title: str, path: Path) -> None:
    language = path.suffix.lstrip(".")
    code = _fenced_code(content, language)
    static_check(path.name, code)
    prefix = '// ' if language in ('js','ts','cpp') else '% ' if language == 'm' else '# '
    notice = _execution_notice(content)
    header = '(* '+notice+' *)\n' if language == 'wl' else prefix+notice+'\n'
    if code.startswith('#!'):
        first,_,remaining=code.partition('\n')
        code=first+'\n'+header+remaining
    else: code=header+code
    path.write_text(code + "\n", encoding="utf-8")


@register("ipynb")
def export_ipynb(content: str, title: str, path: Path) -> None:
    _fenced_code(content, 'py')  # Require at least one Python cell.
    runnable = _runnable_content(content)
    start = len(content)-len(runnable)
    regions = list(re.finditer(r'```(?:python|py)\s*\n(.*?)\n```', content[start:], re.S))
    cells, cursor = [], 0
    for match in regions:
        left, right = start+match.start(), start+match.end()
        if content[cursor:left].strip():
            cells.append({'cell_type':'markdown','metadata':{},'source':[rich_html(content[cursor:left])]})
        code = match[1]; ast.parse(code)
        cells.append({'cell_type':'code','metadata':{},'execution_count':None,'outputs':[], 'source':code.splitlines(keepends=True)})
        cursor = right
    if content[cursor:].strip(): cells.append({'cell_type':'markdown','metadata':{},'source':[rich_html(content[cursor:])]})
    notice = f'# {title}\n\n{_execution_notice(content)}\n\n'
    if cells[0]['cell_type'] == 'markdown': cells[0]['source'].insert(0, notice)
    else: cells.insert(0, {'cell_type':'markdown','metadata':{},'source':[notice]})
    for index, cell in enumerate(cells): cell['id'] = hashlib.sha256((str(index)+''.join(cell['source'])).encode()).hexdigest()[:16]
    notebook = {
        "cells": cells,
        "metadata": {"kernelspec": {"display_name": "Python 3", "language": "python", "name": "python3"}},
        "nbformat": 4, "nbformat_minor": 5,
    }
    path.write_text(json.dumps(notebook, ensure_ascii=False, indent=2), encoding="utf-8")


@register("zip")
def export_zip(content: str, title: str, path: Path) -> None:
    if content.lstrip().startswith('{'):
        files = bundle_files(json_object(content))
    else:
        files = [{'path':'README.md','content':f'# {title}\n\n{content}'}]
        for index, (language, code) in enumerate(re.findall(r"```([A-Za-z0-9_+-]*)\s*\n(.*?)\n```", content, re.S), 1):
            extension = {"python": "py", "javascript": "js", "matlab": "m", "wolfram": "wl", 'c++':'cpp', 'typescript':'ts'}.get(language, language or "txt")
            files.append({'path':f'example_{index}.{extension}', 'content':code})
        files = bundle_files({'files':files})
    checks = [row for file in files if (row := static_check(file['path'], file['content']))]
    if any(file['path'].casefold() == 'studyforge-validation.json' for file in files):
        raise ValueError('StudyForge-validation.json为保留的验证记录文件名')
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for file in files: archive.writestr(file['path'], file['content'])
        archive.writestr('StudyForge-validation.json', json.dumps({'executed':False,'checks':checks,
                        'note':'未执行程序，跨文件协作、依赖和模型结果需自行核对。'}, ensure_ascii=False, indent=2))


@register('md','txt')
def export_text(content: str, title: str, path: Path) -> None:
    path.write_text(content,encoding='utf-8',newline='')


def export_document(content: str, fmt: str, title: str, filename: str) -> Path:
    validate_document_text(content)
    fmt = fmt.lower()
    exporter = EXPORTERS.get(fmt)
    if not exporter:
        raise ValueError(f"不支持导出格式 {fmt}")
    path = artifacts_dir() / filename
    if fmt=='m':
        from .deliverables import matlab_entry_name
        base=matlab_entry_name(_fenced_code(content,'m')) or datetime.now().strftime('studyforge_model_%Y%m%d_%H%M%S')
        directory=artifacts_dir()/Path(filename).stem
        directory.mkdir(parents=True,exist_ok=True)
        path=directory/(base+'.m')
    exporter(content, title, path)
    if not path.exists() or not path.stat().st_size:
        raise RuntimeError("导出文件为空")
    return path
