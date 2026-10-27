"""Opt-in one-question followup, immutable teaching scope and private answers."""
import asyncio
from concurrent.futures import ThreadPoolExecutor
import copy
import json

import httpx
import pytest
from docx import Document
from fastapi.testclient import TestClient

from studyforge import ai, conversations, exercises, lesson_checks, lesson_feedback, server
from studyforge.db import create_project, execute, j, many, now, one, uid, unj
from tests.test_lesson_focus import lesson, selected
from test_mobile_access import mobile_client, pair


def completed_lesson(tmp_path,**kwargs):
    project,chat,q,text,message=lesson(tmp_path,**kwargs)
    focus=selected(q)['id'] if '为什么取距离的一半' in text else None
    lesson_feedback.submit_questionnaire(q['id'],['motivation'],'我看懂计算，但不明白这个半径为什么可以取等号。',focus_section_id=focus)
    reply=uid()
    execute('INSERT INTO messages VALUES (?,?,?,?,?)',(reply,chat['id'],'assistant','COMPLETED_FOLLOWUP：开区间排除端点；任意正半径可以选一半，两个开邻域不会相交。',now()))
    execute("UPDATE lesson_questionnaires SET reply_status='completed',reply_message_id=? WHERE id=?",(reply,q['id']))
    return project,chat,q,text,message,reply


def item(spec):
    return {'question':'NEW_CHECK：数轴上两点为2和8。分别以它们为中心取半径3的开区间，它们是否有共同点？请说明理由，并比较半径4时会怎样。',
        'topic':spec['topic'],'solution':'PRIVATE_LESSON_ANSWER：前者无交集，中点5不属于开区间；半径4时有交集。',
        'rubric':[{'criterion':'边界判断','points':50},{'criterion':'说明开区间条件','points':50}],'source_url':'',
        'lesson_check_focus':'PRIVATE_LESSON_FOCUS：通过改变半径比较开区间端点与共同点，检查为何允许取一半。'}


def test_completed_followup_and_original_focus_are_frozen_without_scoring_the_student(tmp_path):
    p,c,q,_,_,reply=completed_lesson(tmp_path)
    lesson(tmp_path,project=p,chat=c,document='# LATER_LECTURE\n\n不是这次卡点。')
    spec=lesson_checks.build(q['id'],p['id'],c['id'])
    assert 'COMPLETED_FOLLOWUP' in spec['followup']['text'] and spec['followup_message_id']==reply
    assert '极限定义允许任意正数' in spec['lesson_focus']['text']
    assert 'LATER_LECTURE' not in j(spec)
    assert spec['basis'].endswith('not_mastery')
    assert spec==lesson_checks.build(q['id'],p['id'],c['id'])
    assert '严格只生成1题' in lesson_checks.requirements(spec)
    assert not many('SELECT * FROM mastery') and not many('SELECT * FROM exercise_sets')


@pytest.mark.parametrize('change',['not_submitted','running','failed','missing_reply','foreign_chat','foreign_project'])
def test_incomplete_or_foreign_followup_is_not_replaced_with_other_context(tmp_path,change):
    p,c,q,*_=completed_lesson(tmp_path)
    if change=='not_submitted':execute("UPDATE lesson_questionnaires SET status='pending' WHERE id=?",(q['id'],))
    if change in ('running','failed'):execute('UPDATE lesson_questionnaires SET reply_status=? WHERE id=?',(change,q['id']))
    if change=='missing_reply':execute('UPDATE lesson_questionnaires SET reply_message_id=NULL WHERE id=?',(q['id'],))
    if change=='foreign_chat':c=conversations.create(p['id'],'deepseek')
    if change=='foreign_project':p=create_project('其他项目');c=conversations.create(p['id'],'deepseek')
    with pytest.raises(ValueError):lesson_checks.build(q['id'],p['id'],c['id'])
    assert not many('SELECT * FROM lesson_understanding_checks')


def test_retry_does_not_reuse_changed_followup_and_large_teaching_text_is_explicitly_excerpted(tmp_path):
    p,c,q,_,_,reply=completed_lesson(tmp_path)
    before=lesson_checks.build(q['id'],p['id'],c['id'])
    text='长补讲。'*3000
    execute('UPDATE messages SET content=? WHERE id=?',(text,reply))
    after=lesson_checks.build(q['id'],p['id'],c['id'])
    assert after['signature']!=before['signature'] and after['followup']['excerpted']
    assert len(after['followup']['text'])==10000 and after['followup']['characters']==len(text)


@pytest.mark.parametrize('rubric',['含糊地给分',[],[{'criterion':'判断','points':True}],
    [{'criterion':'判断','points':80}],[{'criterion':'','points':100}],
    [{'criterion':'判断','points':110},{'criterion':'理由','points':-10}]])
def test_check_scoring_requires_explicit_positive_parts_adding_to_100(tmp_path,rubric):
    p,c,q,*_=completed_lesson(tmp_path);spec=lesson_checks.build(q['id'],p['id'],c['id'])
    question=item(spec);question['rubric']=rubric
    with pytest.raises(ValueError,match='评分'):lesson_checks.validate([question],spec)
    # Existing version-1 task snapshots remain resumable after this upgrade.
    if isinstance(rubric,str):lesson_checks.validate([question],spec|{'version':1})


def test_atomic_requests_reuse_running_job_and_preserve_retry_and_phone_identity(tmp_path):
    from studyforge.mobile_context import remote_device
    p,c,q,*_=completed_lesson(tmp_path);spec=lesson_checks.build(q['id'],p['id'],c['id'])
    payload={'conversation_id':c['id'],'lesson_check':spec}
    # Real simultaneous SQLite writers, not two sequential calls to a mock.
    with ThreadPoolExecutor(max_workers=4) as workers:
        values=list(workers.map(lambda _:lesson_checks.allocate_job(p['id'],payload,spec),range(4)))
    assert len({value[0] for value in values})==1 and sum(not value[1] for value in values)==1
    ident=values[0][0]
    assert len(many('SELECT * FROM jobs'))==2, 'Original lecture plus one check job'
    execute("UPDATE jobs SET status='paused' WHERE id=?",(ident,))
    assert lesson_checks.allocate_job(p['id'],payload,spec)==(ident,True)
    execute("UPDATE jobs SET status='failed' WHERE id=?",(ident,))
    token=remote_device.set('test-phone')
    try:new,reused=lesson_checks.allocate_job(p['id'],payload,spec)
    finally:remote_device.reset(token)
    assert new!=ident and not reused
    assert unj(one('SELECT payload_json FROM jobs WHERE id=?',(new,))['payload_json'])['_mobile_device_id']=='test-phone'


def test_existing_completed_check_is_reused_but_deleted_question_is_not(tmp_path):
    p,c,q,*_=completed_lesson(tmp_path);spec=lesson_checks.build(q['id'],p['id'],c['id'])
    job,_=lesson_checks.allocate_job(p['id'],{'conversation_id':c['id']},spec)
    ident=uid();question=item(spec)
    execute('INSERT INTO exercise_sets VALUES (?,?,?,?,?,?,?,?,?)',(ident,p['id'],'[]','小检查',j([{'index':0,'question':question['question'],'topic':spec['topic']}]),
        j([{'index':0,'solution':question['solution'],'rubric':question['rubric'],'lesson_check_focus':question['lesson_check_focus']}]),2,'assigned',now()))
    lesson_checks.register(job,ident,spec);execute("UPDATE jobs SET status='completed' WHERE id=?",(job,))
    assert lesson_checks.allocate_job(p['id'],{},spec)==(job,True)
    assert exercises.public_set(ident)['lesson_check']['questionnaire_id']==q['id']
    assert not exercises.public_set(ident,reveal=True).get('solutions')
    execute('DELETE FROM practice_options WHERE exercise_set_id=?',(ident,))
    execute('DELETE FROM exercise_sets WHERE id=?',(ident,))
    assert lesson_checks.latest(q['id'])['exercise_set_id'] is None
    assert not lesson_checks.allocate_job(p['id'],{},spec)[1]


def test_pending_reference_solution_and_design_focus_stay_private_through_word_and_public_api(tmp_path,monkeypatch):
    p,c,q,*_=completed_lesson(tmp_path);spec=lesson_checks.build(q['id'],p['id'],c['id'])
    request_log=[]
    class Client:
        async def complete(self,system,user,**kwargs):
            request_log.append((system,user))
            return j({'questions':[item(spec)]})
    async def search(*args):return []
    async def review(client,digest,questions,context_policy='',validate_revision=None):
        assert '具体疑问' in context_policy and 'COMPLETED_FOLLOWUP' in context_policy
        unchanged=copy.deepcopy(questions);unchanged[0].pop('lesson_check_focus')
        result={'questions':unchanged};validate_revision(result)
        return result['questions'],[{'stage':'受控审查','passed':True}]
    monkeypatch.setattr(exercises,'search_exercises',search);monkeypatch.setattr(exercises,'review_exercises',review)
    ident,_=asyncio.run(exercises.generate_set(Client(),p['id'],[],'SOURCE_EVIDENCE','小检查',lesson_spec=spec))
    job,_=lesson_checks.allocate_job(p['id'],{},spec);lesson_checks.register(job,ident,spec)
    public=exercises.public_set(ident,reveal=True)
    assert len(public['questions'])==1 and not public.get('solutions')
    assert 'PRIVATE_' not in j(public) and 'PRIVATE_' not in exercises.student_document(ident,include_solutions=True)
    assert 'COMPLETED_FOLLOWUP' in request_log[0][1] and 'SOURCE_EVIDENCE' in request_log[0][1]
    from studyforge.exporters import export_document
    word=export_document(exercises.student_document(ident),'docx','小检查','check.docx')
    assert 'PRIVATE_' not in '\n'.join(paragraph.text for paragraph in Document(word).paragraphs)


def test_changed_question_cannot_silently_keep_missing_design_metadata(tmp_path,monkeypatch):
    p,c,q,*_=completed_lesson(tmp_path);spec=lesson_checks.build(q['id'],p['id'],c['id'])
    class Client:
        async def complete(self,*args,**kwargs):return j({'questions':[item(spec)]})
    async def review(client,digest,questions,context_policy='',validate_revision=None):
        changed=copy.deepcopy(questions);changed[0]['question']='完全不同的新题';changed[0].pop('lesson_check_focus')
        validate_revision({'questions':changed})
        return changed,[]
    async def search(*args):return []
    monkeypatch.setattr(exercises,'search_exercises',search);monkeypatch.setattr(exercises,'review_exercises',review)
    with pytest.raises(ValueError,match='lesson_check_focus'):
        asyncio.run(exercises.generate_set(Client(),p['id'],[],'资料','小检查',lesson_spec=spec))
    assert not many('SELECT * FROM exercise_sets')


def test_native_generation_uses_current_teacher_and_original_scope_without_changing_current_selection(tmp_path,monkeypatch):
    p,c,q,*_=completed_lesson(tmp_path);spec=lesson_checks.build(q['id'],p['id'],c['id'])
    source=uid();path=tmp_path/'其他资料.txt';path.write_text('当前另选的资料，不属于原补讲。',encoding='utf-8')
    execute('INSERT INTO sources(id,project_id,name,kind,original_path,status,created_at) VALUES (?,?,?,?,?,?,?)',
        (source,p['id'],path.name,'.txt',str(path),'ready',now()))
    conversations.update(c['id'],p['id'],{'source_ids':[source]})
    captured=[]
    monkeypatch.setattr(server,'start_job',lambda ident,operation:(captured.append((ident,operation)),operation.close()))
    api=TestClient(server.app);url='/api/lesson-feedback/'+q['id']+'/check'
    payload={'project_id':p['id'],'conversation_id':c['id']}
    first=api.post(url,json=payload);assert first.status_code==200,first.text
    repeated=api.post(url,json=payload)
    assert repeated.json()['job_id']==first.json()['job_id'] and repeated.json()['reused']
    assert len(captured)==1 and conversations.detail(c['id'],p['id'])['source_ids']==[source]
    saved=unj(one('SELECT payload_json FROM jobs WHERE id=?',(first.json()['job_id'],))['payload_json'])
    assert saved['model']=='fixture-model' and saved['source_ids']==[] and saved['only_exercises']
    assert saved['library_snapshot']=={} and saved['lesson_check']['signature']==spec['signature']
    assert '当前另选的资料' not in saved['resolved_requirements']
    assert 'COMPLETED_FOLLOWUP' in saved['resolved_requirements']
    assert not many('SELECT * FROM mastery')


def test_native_generation_and_audit_save_one_private_question_with_pause_snapshot(tmp_path,monkeypatch):
    p,c,q,*_=completed_lesson(tmp_path);spec=lesson_checks.build(q['id'],p['id'],c['id'])
    requests=[];native=httpx.AsyncClient
    def respond(request):
        body=json.loads(request.content);requests.append(body)
        system=body['messages'][0]['content']
        if '习题设计者' in system:answer={'questions':[item(spec)]}
        elif '独立解题员' in system:answer={'answers':[{'index':0,'solution':'无交集；端点被排除。','issues':[]}]}
        else:answer={'passed':True,'issues':[],'independent_issue_decisions':[]}
        return httpx.Response(200,text='data: '+j({'choices':[{'delta':{'content':j(answer)},'finish_reason':'stop'}]})+'\n\ndata: [DONE]\n\n')
    monkeypatch.setattr(ai,'get_secret',lambda key:'fixture-key')
    monkeypatch.setattr(httpx,'AsyncClient',lambda *a,**kw:native(transport=httpx.MockTransport(respond)))
    async def search(*args):return []
    monkeypatch.setattr(exercises,'search_exercises',search)
    value=server.GenerateIn(project_id=p['id'],conversation_id=c['id'],source_ids=[],only_exercises=True,
        provider='deepseek',model='fixture-model',lesson_check=spec,resolved_requirements=lesson_checks.requirements(spec))
    job,_=lesson_checks.allocate_job(p['id'],value.model_dump(),spec)
    asyncio.run(server._generate_job(job,value))
    result=one('SELECT * FROM jobs WHERE id=?',(job,));assert result['status']=='completed',result['error']
    link=lesson_checks.latest(q['id']);data=exercises.public_set(link['exercise_set_id'],reveal=True)
    assert len(data['questions'])==1 and not data.get('solutions') and len(requests)==4
    assert 'COMPLETED_FOLLOWUP' in requests[0]['messages'][-1]['content']
    artifact=one("SELECT * FROM artifacts WHERE kind='exercise'")
    assert artifact['project_id']==p['id'] and unj(artifact['metadata_json'])['conversation_id']==c['id']
    assert 'PRIVATE_' not in '\n'.join(paragraph.text for paragraph in Document(artifact['path']).paragraphs)
    assert one('SELECT recovery_mode FROM practice_options')['recovery_mode']=='adaptive'
    assert not many('SELECT * FROM mastery')
    # Replaying the saved exercise stage must not call the model again.
    asyncio.run(server._generate_job(job,value));assert len(requests)==4


def test_phone_check_entry_keeps_device_identity_and_rejects_other_conversation(tmp_path,monkeypatch):
    p,c,q,*_=completed_lesson(tmp_path)
    monkeypatch.setattr(server,'start_job',lambda ident,operation:operation.close())
    manager,client=mobile_client()
    with client:
        pair(manager,client)
        url='/api/lesson-feedback/'+q['id']+'/check'
        another=conversations.create(p['id'],'deepseek')
        assert client.post(url,json={'project_id':p['id'],'conversation_id':another['id']}).status_code==400
        response=client.post(url,json={'project_id':p['id'],'conversation_id':c['id']})
        assert response.status_code==200,response.text
        payload=unj(one('SELECT payload_json FROM jobs WHERE id=?',(response.json()['job_id'],))['payload_json'])
        assert payload.get('_mobile_device_id')


def test_complete_submission_releases_this_check_reference_before_the_grader_runs(tmp_path,monkeypatch):
    from studyforge import homework, job_control
    p,c,q,*_=completed_lesson(tmp_path);spec=lesson_checks.build(q['id'],p['id'],c['id'])
    question=item(spec);ident=uid()
    execute('INSERT INTO exercise_sets VALUES (?,?,?,?,?,?,?,?,?)',(ident,p['id'],'[]','本次补讲检查',
        j([{'index':0,'question':question['question'],'topic':question['topic']}]),
        j([{'index':0,'solution':question['solution'],'rubric':question['rubric'],'lesson_check_focus':question['lesson_check_focus']}]),2,'assigned',now()))
    job,_=lesson_checks.allocate_job(p['id'],{},spec);lesson_checks.register(job,ident,spec)
    async def choose(*args,**kwargs):pass
    monkeypatch.setattr(server,'choose_generation_model',choose)
    # Save and export the actual submission, but leave asynchronous grading queued.
    monkeypatch.setattr(job_control,'start_job',lambda ident,operation:operation.close())
    api=TestClient(server.app)
    url=f'/api/exercises/{ident}/submit'
    body={'provider':'deepseek','model':'fixture-model','conversation_id':c['id'],
        'background':True,'detailed_feedback':True,'answers':['']}
    assert api.post(url,json=body).status_code==400
    assert not many('SELECT * FROM homework_attempts') and not many('SELECT * FROM submissions')
    assert 'PRIVATE_' not in exercises.student_document(ident,include_solutions=True)
    body['answers']=['两个半径3的开区间在5处都不含端点，因此没有共同点；半径4时区间重叠。']
    response=api.post(url,json=body);assert response.status_code==200,response.text
    attempt=homework.get_attempt(response.json()['id'])
    assert one('SELECT status FROM jobs WHERE id=?',(attempt['job_id'],))['status']=='queued'
    reference=one('SELECT path,metadata_json FROM artifacts WHERE id=?',(response.json()['solution_artifact']['id'],))
    text='\n'.join(p.text for p in Document(reference['path']).paragraphs)
    assert 'PRIVATE_LESSON_ANSWER' in text and 'PRIVATE_LESSON_FOCUS' in text
    assert '本题与补讲的关系' in text and 'solutions' in exercises.public_set(ident,reveal=True)
    assert unj(reference['metadata_json'])['conversation_id']==c['id']
    assert not many('SELECT * FROM submissions') and not many('SELECT * FROM mastery')


def test_generate_rebuilds_forged_check_and_rejects_a_different_source_scope(tmp_path,monkeypatch):
    p,c,q,*_=completed_lesson(tmp_path)
    monkeypatch.setattr(server,'start_job',lambda ident,operation:operation.close())
    api=TestClient(server.app)
    body={'project_id':p['id'],'conversation_id':c['id'],'provider':'deepseek','model':'fixture-model',
        'source_ids':[],'only_exercises':True,'instructions':'只围绕本次已完成补讲检查一次理解。',
        'lesson_check':{'questionnaire_id':q['id'],'topic':'FORGED_TOPIC','followup':{'text':'FORGED_REPLY'},'difficulty':5}}
    response=api.post('/api/generate',json=body);assert response.status_code==200,response.text
    saved=unj(one('SELECT payload_json FROM jobs WHERE id=?',(response.json()['job_id'],))['payload_json'])
    assert saved['lesson_check']==lesson_checks.build(q['id'],p['id'],c['id'])
    assert 'FORGED_' not in saved['resolved_requirements'] and saved['lesson_check']['difficulty']==2
    count=len(many('SELECT * FROM jobs'))
    body['source_ids']=['different-source-id']
    response=api.post('/api/generate',json=body)
    assert response.status_code==400 and '原资料范围' in response.text
    assert len(many('SELECT * FROM jobs'))==count


def test_check_keeps_original_book_inputs_after_current_chat_references_change(tmp_path,monkeypatch):
    from studyforge import library
    from test_library_generation_evidence import prepared,add_reading
    book,node,reading,p,conversation,snapshot=prepared(tmp_path)
    chat=one('SELECT * FROM conversations WHERE id=?',(conversation,))
    _,c,q,*_=completed_lesson(tmp_path,project=p,chat=chat)
    reference=library.reference_content(snapshot,'ORIGINAL_SOURCE')
    payload=unj(one('SELECT payload_json FROM jobs WHERE id=?',(q['job_id'],))['payload_json'])
    original={'library_snapshot':snapshot,'library_context':reference['text'],'library_evidence':library.reference_receipt(reference)}
    execute('UPDATE jobs SET payload_json=? WHERE id=?',(j(payload|original),q['job_id']))
    library.set_preferences(c['id'],p['id'],'off',[])
    add_reading(book,node,'newer-reader','NEWER_SOURCE')
    monkeypatch.setattr(server,'start_job',lambda ident,operation:operation.close())
    response=TestClient(server.app).post(f'/api/lesson-feedback/{q["id"]}/check',json={'project_id':p['id'],'conversation_id':c['id']})
    assert response.status_code==200,response.text
    saved=unj(one('SELECT payload_json FROM jobs WHERE id=?',(response.json()['job_id'],))['payload_json'])
    assert {key:saved[key] for key in original}==original
    records=library.generation_records(saved['library_snapshot'])
    assert records[0]['read_manifest']['library_reading']['id']==reading
    assert 'ORIGINAL_SOURCE' in records[0]['segments'][0]['text'] and 'NEWER_SOURCE' not in records[0]['segments'][0]['text']


def test_removing_an_unfinished_check_cancels_it_and_allows_an_explicit_new_request(tmp_path,monkeypatch):
    p,c,q,*_=completed_lesson(tmp_path)
    monkeypatch.setattr(server,'start_job',lambda ident,operation:operation.close())
    api=TestClient(server.app);url=f'/api/lesson-feedback/{q["id"]}/check'
    body={'project_id':p['id'],'conversation_id':c['id']}
    first=api.post(url,json=body).json()['job_id']
    removed=api.delete(f'/api/jobs/{first}');assert removed.status_code==200,removed.text
    assert lesson_checks.latest(q['id'])['status']=='cancelled'
    assert first not in {row['id'] for row in server.jobs(p['id'])}
    second=api.post(url,json=body);assert second.status_code==200,second.text
    assert second.json()['job_id']!=first and not second.json().get('reused')
    assert not many('SELECT * FROM exercise_sets') and not many('SELECT * FROM mastery')
