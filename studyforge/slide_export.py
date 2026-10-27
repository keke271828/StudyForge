"""Readable slide layout with rasterized math and preserved source text in notes."""
from __future__ import annotations

import hashlib
import unicodedata
from pathlib import Path

from .config import temp_dir
from .document_render import (blocks, inline_parts, math_image, owned_image, image_size,
                              EXPLANATION_COLOR, SOURCE_COLOR, EXPLANATION_FONT, SOURCE_FONT)


def wrap_text(text, width=46):
    lines, current, size = [], '', 0
    for char in text:
        weight = 1 if unicodedata.east_asian_width(char) in ('W','F') else .55
        if char == '\n' or size+weight > width:
            lines.append(current);current, size = '', 0
            if char == '\n': continue
        current += char; size += weight
    if current: lines.append(current)
    return lines or ['']


def inline_images(text, quote=False):
    from PIL import Image, ImageDraw, ImageFont
    color = SOURCE_COLOR if quote else EXPLANATION_COLOR
    key = hashlib.sha256(('ppt-v1:'+str(quote)+text).encode()).hexdigest()
    font_path = Path('C:/Windows/Fonts/simsun.ttc' if quote else 'C:/Windows/Fonts/msyh.ttc')
    font = ImageFont.truetype(str(font_path), 50) if font_path.exists() else ImageFont.truetype('DejaVuSans.ttf',50)
    probe = ImageDraw.Draw(Image.new('RGB',(1,1)))
    rows, current, width, height = [], [], 0, 70
    def add(item, item_width, item_height):
        nonlocal current, width, height
        if current and width+item_width > 1760:
            rows.append((current,height));current,width,height=[],0,70
        current.append((item,item_width,item_height));width+=item_width;height=max(height,item_height+16)
    for kind, part in inline_parts(text):
        if kind == 'math':
            image = Image.open(math_image(part, color=color)).convert('RGBA')
            if image.width > 1760:
                factor=1760/image.width
                if factor < .65: raise ValueError('PPT行内公式过宽，请拆分为分行公式后再导出')
                image=image.resize((1760,round(image.height*factor)))
            add(image,image.width,image.height)
        else:
            for char in part:
                if char == '\n':
                    rows.append((current,height));current,width,height=[],0,70;continue
                add(char,probe.textlength(char,font=font),60)
    if current: rows.append((current,height))
    images=[]
    for offset in range(0,len(rows),7):
        group=rows[offset:offset+7]
        width=min(1760,max(1,round(max(sum(w for _,w,_ in row) for row,_ in group))+2))
        result=Image.new('RGB',(width,max(1,sum(h for _,h in group))), '#F3F5F7' if quote else 'white')
        draw=ImageDraw.Draw(result);top=0
        for row, row_height in group:
            left=0
            for item,w,h in row:
                if isinstance(item,str): draw.text((left,top+(row_height-60)/2),item,font=font,fill='#'+color)
                else: result.paste(item,(round(left),round(top+(row_height-h)/2)),item)
                left+=w
            top+=row_height
        path=temp_dir()/f'slide-inline-{key}-{offset}.png';result.save(path);images.append(path)
    return images


def export_slides(content, title, path):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.oxml.xmlchemy import OxmlElement
    prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
    slide, y, section, count = None, 0, title, 0
    def text_box(text,x,top,width,height,size=18,quote=False,bold=False):
        shape=slide.shapes.add_textbox(Inches(x),Inches(top),Inches(width),Inches(height))
        frame=shape.text_frame;frame.word_wrap=True;frame.margin_left=frame.margin_right=0
        frame.margin_top=frame.margin_bottom=0
        for index,line in enumerate(text.splitlines()):
            paragraph=frame.paragraphs[0] if index==0 else frame.add_paragraph()
            paragraph.text=line;paragraph.font.size=Pt(size);paragraph.font.bold=bold
            paragraph.font.name='Times New Roman' if quote else EXPLANATION_FONT
            paragraph.font.color.rgb=RGBColor.from_string(SOURCE_COLOR if quote else EXPLANATION_COLOR)
            ea=OxmlElement('a:ea');ea.set('typeface',SOURCE_FONT if quote else EXPLANATION_FONT)
            paragraph.font._element.append(ea)
        return shape
    def new_slide(heading):
        nonlocal slide,y,count
        slide=prs.slides.add_slide(prs.slide_layouts[6]);count+=1;y=1.4
        displayed=heading if len(heading)<=75 else heading[:72]+'…'
        text_box(displayed,.55,.38,12.2,.85,26 if len(displayed)<40 else 22,bold=True)
        text_box(f'StudyForge · {count}',.55,7.08,12.2,.22,10)
        slide.notes_slide.notes_text_frame.text=heading+'\n'
    new_slide(title)
    text_box('学习讲解 · 按章节整理',.6,2.4,12,1,24)
    slide=None
    def ensure(height):
        nonlocal slide
        if height>5.3: raise ValueError('PPT单个内容块过高，请拆分公式、图表或段落后再导出')
        if slide is None or y+height>6.85: new_slide(section)
    def picture(image, max_height=5.1, caption='', natural_dpi=None):
        nonlocal y
        if natural_dpi:
            from PIL import Image
            with Image.open(image) as bitmap: width,height=bitmap.width/natural_dpi,bitmap.height/natural_dpi
            if width>12.1 or height>max_height: raise ValueError('PPT公式或图文过大，请拆分后再导出')
        else: width,height=image_size(image,12.1,max_height)
        ensure(height+.15)
        slide.shapes.add_picture(str(image),Inches(.6+(12.1-width)/2),Inches(y),width=Inches(width),height=Inches(height))
        y+=height+.15
        if caption: plain(caption)
    def plain(value,quote=False,code=False,bold=False):
        nonlocal y
        size=14 if code else 18
        lines=wrap_text(value,60 if code else 46)
        for offset in range(0,len(lines),10):
            text='\n'.join(lines[offset:offset+10]);height=len(lines[offset:offset+10])*(.28 if code else .36)+.12
            ensure(height);shape=text_box(text,.65,y,12.05,height,size,quote,bold)
            if code:
                for paragraph in shape.text_frame.paragraphs: paragraph.font.name='Consolas'
            if quote:
                shape.fill.solid();shape.fill.fore_color.rgb=RGBColor.from_string('F3F5F7')
            y+=height+.10
    for kind,value in blocks(content):
        if kind=='heading':
            if value[0]<=2: section=value[1];slide=None
            else: plain(value[1],bold=True)
            continue
        if slide is None: new_slide(section)
        if kind in ('math','quote_math'):
            image=math_image(value,color=SOURCE_COLOR if kind=='quote_math' else EXPLANATION_COLOR)
            picture(image,max_height=4,natural_dpi=200)
        elif kind=='paragraph' and (image:=owned_image(value)):
            picture(image[0],caption=image[1])
        elif kind=='table':
            for row in value:
                line=' | '.join(row)
                if any(k=='math' for k,_ in inline_parts(line)):
                    for image in inline_images(line): picture(image,max_height=4,natural_dpi=200)
                else: plain(line)
        elif kind=='code': plain(value,code=True)
        else:
            quote=kind=='quote';text=('• ' if kind=='bullet' else '')+value
            if any(k=='math' for k,_ in inline_parts(text)):
                for image in inline_images(text,quote): picture(image,max_height=4,natural_dpi=200)
            else: plain(''.join(p for _,p in inline_parts(text)),quote)
        # Plain text and LaTeX stay available for inspection/accessibility; the
        # visible slide uses typeset formulas, not raw command strings.
        slide.notes_slide.notes_text_frame.text += '\n'+(str(value))
    prs.save(path)
