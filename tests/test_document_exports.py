import asyncio
import json
import threading
import xml.etree.ElementTree as ET
from pathlib import Path

import fitz
import httpx
import pytest
from fastapi.testclient import TestClient

from studyforge import conversations, document_exports as exports, server
from studyforge.ai import save_settings
from studyforge.config import artifacts_dir
from studyforge.db import connect, create_job, create_project, execute, j, many, now, one, uid, unj, update_job
from studyforge.generation_control import active_generations, control_for
from studyforge.job_control import active_jobs, cancel_job, start_job

BODY = r'''# 为什么极限唯一

只在一个点非负不足以证明处处非负。取两个不相交的开邻域；严格小于号排除共同端点。

$$\mathbb{Q}(\sqrt2)=\{a+b\sqrt2\mid a,b\in\mathbb{Q}\}$$

> 原文：不同极限的距离为正。

[来源：极限.txt｜第 1 段]

## 系统性总结

半径可以取距离的一半，因为两个邻域是开区间。这里不另外布置习题。
'''


def snapshot(tmp_path, *, publish=True, status='completed'):
    project=create_project('数学学习');chat=conversations.create(project['id'],'deepseek')
    source=uid();path=tmp_path/'极限.txt';path.write_text('不同极限的距离为正。',encoding='utf-8')
    execute('INSERT INTO sources(id,project_id,name,kind,original_path,status,created_at) VALUES (?,?,?,?,?,?,?)',
        (source,project['id'],path.name,'.txt',str(path),'ready',now()))
    job=create_job(project['id'],'generate',{'conversation_id':chat['id'],'source_ids':[source]})
    meta={'generation_job_id':job,'conversation_id':chat['id'],'model':'原讲解模型','review_mode':'fast','source_reading':[{'name':path.name,'characters':10}]}
    exports.save(job,project['id'],'为什么极限唯一',BODY,[source],meta)
    if publish:
        output=server.export_document(BODY,'docx','为什么极限唯一',job+'.docx')
        artifact=server.add_artifact(project['id'],[source],'explanation','docx','为什么极限唯一',output,meta)
        exports.publish_lecture(one('SELECT * FROM document_snapshots WHERE job_id=?',(job,)))
    else:artifact=None
    update_job(job,status,100)
    return project,chat,job,source,artifact


def execute_export(project,ident,fmt,*,artifact=True):
    job,reused=exports.launch(project['id'],ident,fmt,uid(),artifact)
    assert not reused
    asyncio.run(exports.run_job(job))
    row=one('SELECT * FROM jobs WHERE id=?',(job,))
    assert row['status']=='completed',row['error']
    result=unj(row['result_json']);assert result['model_requests']==0
    return job,one('SELECT * FROM artifacts WHERE id=?',(result['artifacts'][0]['id'],))


@pytest.mark.parametrize('fmt',exports.FORMATS)
def test_local_formats_use_the_same_body_and_keep_project_chat_and_no_model(tmp_path,monkeypatch,fmt):
    project,chat,root,source,artifact=snapshot(tmp_path)
    monkeypatch.setattr(server,'client_for',lambda *a,**k:pytest.fail('另存不应创建模型客户端'))
    job,row=execute_export(project,artifact['id'],fmt)
    path=Path(row['path']);assert path.is_file() and path.stat().st_size>0
    meta=unj(row['metadata_json']);assert meta['generation_job_id']==root and meta['reexport_job_id']==job
    assert meta['conversation_id']==chat['id'] and meta['reexport']['model_requests']==0
    assert meta['reexport']['content_sha256']==one('SELECT content_sha256 FROM document_snapshots')['content_sha256']
    assert unj(row['source_ids_json'])==[source]
    public=TestClient(server.app).get(f"/api/projects/{project['id']}/artifacts").json()
    assert len(public)==2 and all(r['origin']['conversation_id']==chat['id'] for r in public)
    # Completed exports become selectable project resources without a manual import.
    sources=server.sources(project['id'])
    assert any(r.get('artifact_id')==row['id'] for r in sources)
    assert len(many('SELECT * FROM generation_dialogue_outputs'))==1
    assert len(many('SELECT * FROM lesson_questionnaires'))==1
    assert not many("SELECT * FROM execution_spans WHERE resource='model'")
    if fmt=='docx':
        from docx import Document
        assert '严格小于号' in '\n'.join(p.text for p in Document(path).paragraphs)
    elif fmt=='pdf':
        with fitz.open(path) as pdf:
            assert '严格小于号' in ''.join(page.get_text() for page in pdf)
            assert sum(len(page.get_images()) for page in pdf)>=1
    elif fmt=='pptx':
        from pptx import Presentation
        assert '严格小于号' in '\n'.join(s.text for slide in Presentation(path).slides for s in slide.shapes if s.has_text_frame)
    elif fmt=='mm': assert '严格小于号' in ET.tostring(ET.parse(path).getroot(),encoding='unicode')
    elif fmt in ('md','txt'):assert path.read_text(encoding='utf-8')==BODY


def test_complete_audited_body_survives_real_generation_layout_failure_and_recovers_without_ai(tmp_path,monkeypatch):
    save_settings({'chat_provider':'deepseek','deepseek_model':'lesson-model','background_learning_enabled':False})
    project=create_project('生成失败补救');chat=conversations.create(project['id'],'deepseek')
    source=uid();path=tmp_path/'极限.txt';path.write_text('不同极限的距离为正。',encoding='utf-8')
    execute('INSERT INTO sources(id,project_id,name,kind,original_path,status,created_at) VALUES (?,?,?,?,?,?,?)',
        (source,project['id'],path.name,'.txt',str(path),'ready',now()))
    requests=[];client=httpx.AsyncClient
    def respond(request):
        payload=json.loads(request.content);requests.append(payload)
        answer=j({'passed':True,'issues':[]}) if payload.get('response_format',{}).get('type')=='json_object' else BODY
        return httpx.Response(200,text='data: '+j({'choices':[{'delta':{'content':answer},'finish_reason':'stop'}]})+'\n\ndata: [DONE]\n\n')
    monkeypatch.setattr('studyforge.ai.get_secret',lambda key:'test-key')
    monkeypatch.setattr(httpx,'AsyncClient',lambda *args,**kwargs:client(transport=httpx.MockTransport(respond)))
    original=server.export_document
    def fail_pdf(*args,**kwargs):
        if args[1]=='pdf':raise ValueError('模拟仅 PDF 排版失败')
        return original(*args,**kwargs)
    monkeypatch.setattr(server,'export_document',fail_pdf)
    value=server.GenerateIn(project_id=project['id'],source_ids=[source],provider='deepseek',conversation_id=chat['id'],
        model='lesson-model',format='pdf',exercises=False,review_mode='fast')
    root=create_job(project['id'],'generate',value.model_dump());asyncio.run(server._generate_job(root,value))
    assert one('SELECT status,error FROM jobs WHERE id=?',(root,))['status']=='failed'
    assert one('SELECT content FROM document_snapshots WHERE job_id=?',(root,)),one('SELECT error FROM jobs WHERE id=?',(root,))
    assert len(requests)==2
    _,artifact=execute_export(project,root,'docx',artifact=False)
    assert artifact['format']=='docx' and len(requests)==2
    assert one('SELECT status FROM jobs WHERE id=?',(root,))['status']=='failed'
    assert len(many('SELECT * FROM generation_dialogue_outputs'))==1
    assert len(many('SELECT * FROM lesson_questionnaires'))==1
    assert not many('SELECT * FROM exercise_sets')


def test_idempotence_version_guard_and_cross_project_answer_protection(tmp_path):
    project,chat,root,source,artifact=snapshot(tmp_path);other=create_project('另一项目')
    request=uid();first,reused=exports.launch(project['id'],artifact['id'],'docx',request,True);assert not reused
    assert exports.launch(project['id'],artifact['id'],'docx',request,True)==(first,True)
    assert exports.launch(project['id'],artifact['id'],'docx',uid(),True)==(first,True)
    with pytest.raises(ValueError,match='同一次'):exports.launch(project['id'],artifact['id'],'pdf',request,True)
    with pytest.raises(LookupError):exports.launch(other['id'],artifact['id'],'docx',uid(),True)
    with pytest.raises(ValueError,match='另一版'):exports.save(root,project['id'],'改变正文',BODY+'新内容',[source],{})
    exercise=server.add_artifact(project['id'],[source],'exercise','docx','题单',Path(one('SELECT path FROM artifacts')['path']),{})
    with pytest.raises(ValueError,match='私有习题'):exports.launch(project['id'],exercise['id'],'txt',uid(),True)
    api=TestClient(server.app);assert api.post('/api/artifacts/'+artifact['id']+'/export',json={
        'project_id':other['id'],'format':'docx','request_id':uid()}).status_code==404
    assert not next(r for r in api.get(f"/api/projects/{project['id']}/artifacts").json() if r['id']==exercise['id'])['export_formats']
    assert BODY not in api.get(f"/api/projects/{project['id']}/jobs").text
    assert api.post('/api/jobs/'+root+'/export',json={'project_id':project['id'],'format':'zip','request_id':uid()}).status_code==422


def test_legacy_exact_message_can_reexport_after_history_deletion_or_cancel(tmp_path):
    project,chat,root,source,artifact=snapshot(tmp_path)
    execute('DELETE FROM document_snapshots WHERE job_id=?',(root,));update_job(root,'cancelled',100)
    execute('INSERT INTO deleted_job_records VALUES (?,?)',(root,now()))
    assert exports.available(many('SELECT * FROM artifacts'),True,project['id'])[artifact['id']]
    _,row=execute_export(project,artifact['id'],'txt')
    assert Path(row['path']).read_text(encoding='utf-8')==BODY
    with pytest.raises(ValueError):exports.launch(project['id'],root,'docx',uid())


def test_missing_legacy_record_is_not_reconstructed_and_deleted_chat_is_not_republished(tmp_path):
    project,chat,root,source,artifact=snapshot(tmp_path)
    execute('DELETE FROM document_snapshots WHERE job_id=?',(root,))
    execute('DELETE FROM generation_dialogue_outputs WHERE job_id=?',(root,))
    assert artifact['id'] not in exports.available(many('SELECT * FROM artifacts'),True,project['id'])
    with pytest.raises(ValueError,match='没有保存'):exports.resolve(project['id'],artifact['id'],True)
    exports.save(root,project['id'],'原讲解',BODY,[source],{'conversation_id':chat['id']})
    execute('INSERT INTO deleted_conversations VALUES (?,?)',(chat['id'],now()))
    _,row=execute_export(project,artifact['id'],'txt')
    assert not many('SELECT * FROM generation_dialogue_outputs')


def test_renamed_artifact_and_changed_source_keep_original_lecture(tmp_path):
    project,chat,root,source,artifact=snapshot(tmp_path)
    execute('UPDATE artifacts SET title=? WHERE id=?',('我重命名的讲解',artifact['id']))
    Path(one('SELECT original_path FROM sources WHERE id=?',(source,))['original_path']).write_text('已经改过的原文件',encoding='utf-8')
    _,row=execute_export(project,artifact['id'],'txt')
    assert row['title']=='我重命名的讲解' and Path(row['path']).read_text(encoding='utf-8')==BODY
    assert server.artifact_download_name(row).endswith('.txt') and '我重命名的讲解' in server.artifact_download_name(row)


def test_deleting_original_artifact_does_not_disable_feedback_after_another_export(tmp_path):
    project,chat,root,source,artifact=snapshot(tmp_path)
    questionnaire=one('SELECT * FROM lesson_questionnaires WHERE job_id=?',(root,))
    execute('DELETE FROM artifacts WHERE id=?',(artifact['id'],))
    _,row=execute_export(project,root,'docx',artifact=False)
    feedback=TestClient(server.app).get('/api/lesson-feedback/by-job/'+root).json()
    assert feedback['id']==questionnaire['id'] and feedback['ready'] and feedback['artifact_id']==row['id']


def test_parallel_exports_bound_real_layout_work_to_one(tmp_path,monkeypatch):
    project,chat,root,source,artifact=snapshot(tmp_path)
    jobs=[exports.launch(project['id'],root,fmt,uid())[0] for fmt in ('docx','md','txt')]
    original=exports.write_file;lock=threading.Lock();stats={'active':0,'peak':0}
    def measured(*args):
        import time
        with lock:stats['active']+=1;stats['peak']=max(stats['peak'],stats['active'])
        try:time.sleep(.04);return original(*args)
        finally:
            with lock:stats['active']-=1
    monkeypatch.setattr(exports,'write_file',measured)
    async def scenario():await asyncio.gather(*(start_job(job,exports.run_job(job)) for job in jobs))
    asyncio.run(scenario())
    assert stats['peak']==1 and all(one('SELECT status FROM jobs WHERE id=?',(job,))['status']=='completed' for job in jobs)
    assert len(many('SELECT * FROM lesson_questionnaires'))==1


@pytest.mark.parametrize('restart',[False,True])
def test_pause_and_restart_reuse_one_local_write(tmp_path,monkeypatch,restart):
    project,chat,root,source,artifact=snapshot(tmp_path);job,_=exports.launch(project['id'],root,'txt',uid())
    entered=threading.Event();release=threading.Event();calls=[];original=exports.write_file
    def blocked(*args):
        calls.append(args);entered.set();assert release.wait(10);return original(*args)
    monkeypatch.setattr(exports,'write_file',blocked)
    async def scenario():
        control=control_for(job);task=start_job(job,exports.run_job(job))
        while not entered.is_set():await asyncio.sleep(.01)
        await server.pause_generation(job);release.set()
        for _ in range(200):
            if one('SELECT status FROM jobs WHERE id=?',(job,))['status']=='paused' and control.in_flight==0:break
            await asyncio.sleep(.01)
        assert not many("SELECT * FROM artifacts WHERE json_extract(metadata_json,'$.reexport_job_id')=?",(job,))
        if restart:
            task.cancel();await asyncio.gather(task,return_exceptions=True)
            assert job not in active_generations
            from studyforge.desktop_lifecycle import recover_interrupted_jobs
            recover_interrupted_jobs();await server.resume_generation(job)
            task=active_jobs[job]
        else:await server.resume_generation(job)
        await task
    asyncio.run(scenario())
    assert one('SELECT status FROM jobs WHERE id=?',(job,))['status']=='completed'
    assert len(calls)==1 and len(many("SELECT * FROM artifacts WHERE json_extract(metadata_json,'$.reexport_job_id')=?",(job,)))==1


def test_cancel_waits_for_owned_file_work_and_does_not_publish_late(tmp_path,monkeypatch):
    project,chat,root,source,artifact=snapshot(tmp_path);job,_=exports.launch(project['id'],root,'txt',uid())
    entered=threading.Event();release=threading.Event();original=exports.write_file
    def blocked(*args):entered.set();assert release.wait(10);return original(*args)
    monkeypatch.setattr(exports,'write_file',blocked)
    async def scenario():
        start_job(job,exports.run_job(job))
        while not entered.is_set():await asyncio.sleep(.01)
        cancellation=asyncio.create_task(cancel_job(job));await asyncio.sleep(.03)
        assert not cancellation.done();release.set();await cancellation
    asyncio.run(scenario())
    assert not many("SELECT * FROM artifacts WHERE json_extract(metadata_json,'$.reexport_job_id')=?",(job,))
    assert not list(artifacts_dir().glob(job+'-*'))
    assert one('SELECT * FROM document_snapshots WHERE job_id=?',(root,))


def test_failed_partial_file_cleanup_and_fresh_export_can_retry(tmp_path,monkeypatch):
    project,chat,root,source,artifact=snapshot(tmp_path);job,_=exports.launch(project['id'],root,'txt',uid())
    import studyforge.exporters as module
    original=module.export_document
    def partial(content,fmt,title,filename):
        (artifacts_dir()/filename).write_text('半截文件',encoding='utf-8');raise RuntimeError('磁盘写入测试失败')
    monkeypatch.setattr(module,'export_document',partial);asyncio.run(exports.run_job(job))
    assert one('SELECT status FROM jobs WHERE id=?',(job,))['status']=='failed'
    assert not list(artifacts_dir().glob(job+'-*'))
    monkeypatch.setattr(module,'export_document',original);execute_export(project,root,'txt',artifact=False)


def test_async_endpoint_runs_real_local_job_and_restart_recovery_keeps_can_pause(tmp_path):
    project,chat,root,source,artifact=snapshot(tmp_path)
    async def scenario():
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=server.app),base_url='http://test') as api:
            response=await api.post('/api/artifacts/'+artifact['id']+'/export',json={
                'project_id':project['id'],'format':'txt','request_id':uid()})
            assert response.status_code==200,response.text
            job=response.json()['job_id'];await active_jobs[job]
            result=(await api.get('/api/jobs/'+job)).json();assert result['result']['model_requests']==0
            assert result['can_pause'] and result['status']=='completed'
            download=await api.get(result['result']['artifacts'][0]['download_url']);assert download.status_code==200
            assert download.text==BODY
    asyncio.run(scenario())
    job,_=exports.launch(project['id'],root,'md',uid())
    from studyforge.desktop_lifecycle import recover_interrupted_jobs
    assert recover_interrupted_jobs()['paused_generations']==1
    assert TestClient(server.app).get('/api/jobs/'+job).json()['can_pause']


def test_schema_migration_and_project_purge_keep_other_snapshots(tmp_path):
    from studyforge import db
    project,chat,root,source,artifact=snapshot(tmp_path)
    with connect() as conn:conn.execute('DROP TABLE document_snapshots');conn.execute('DROP INDEX document_export_request')
    db._initialized_databases.clear()
    exports.save(root,project['id'],'原讲解',BODY,[source],{'conversation_id':chat['id']})
    other=create_project('保留项目');otherjob=create_job(other['id'],'generate',{})
    exports.save(otherjob,other['id'],'另一讲解','原文。',[],{})
    _,row=execute_export(project,artifact['id'],'txt')
    api=TestClient(server.app);assert api.delete('/api/projects/'+project['id']).status_code==200
    response=api.request('DELETE','/api/recycle-bin/projects/'+project['id'],json={'confirmation_name':project['name']})
    assert response.status_code==200,response.text
    assert not one('SELECT * FROM document_snapshots WHERE job_id=?',(root,))
    assert one('SELECT * FROM document_snapshots WHERE job_id=?',(otherjob,))
    assert not Path(row['path']).exists()
