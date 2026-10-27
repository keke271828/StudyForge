"""Optional live AI check of read -> conversation requirements -> explanation."""
import asyncio
import json
import os
import tempfile
from pathlib import Path

from docx import Document

from studyforge import server
from studyforge.codex_bridge import bridge
from studyforge.db import create_job, create_project, execute, j, now, one, uid


async def main():
    root = Path(tempfile.mkdtemp(prefix='learning-live-', dir=Path(__file__).resolve().parents[1] / 'build'))
    os.environ['STUDYFORGE_DATA_DIR'] = str(root)
    project = create_project('因果与相关自检')
    source_id = uid()
    original = root / '因果与相关.txt'
    original.write_text('相关性是两个变量在观测数据中共同变化的关系，不能仅凭相关推出因果。'
                        '共同原因、反向因果和偶然都可能产生相关。社会科学的因果分析应说明理论机制、'
                        '证据、混杂因素及适用条件。不同研究设计的假设不同；即使采用实验也须考虑外部效度。', encoding='utf-8')
    execute('INSERT INTO sources(id,project_id,name,kind,original_path,segments_json,status,error,created_at) VALUES (?,?,?,?,?,?,?,?,?)',
            (source_id, project['id'], original.name, '.txt', str(original), j([dict(locator='预览', text='旧摘要：相关就代表因果。')]), 'ready', None, now()))
    try:
        read = server.ReadIn(project_id=project['id'], source_ids=[source_id])
        await server.choose_generation_model(read)
        print(j({'stage': 'reading', 'model': read.model, 'effort': read.reasoning_effort}), flush=True)
        read_job = create_job(project['id'], 'read', read.model_dump())
        await server._read_job(read_job, read)
        row = one('SELECT * FROM jobs WHERE id=?', (read_job,))
        assert row['status'] == 'completed', row['error']
        reading = json.loads(row['result_json'])
        chat = await server.chat_stream(server.ChatIn(project_id=project['id'], source_ids=[source_id],
             reading_id=reading['reading_id'], model=read.model, reasoning_effort=read.reasoning_effort,
             text='我的理解是：两者同时上涨，说明冰淇淋导致溺水。请先指出这个理解哪里不对。'
                  '这份资料的讲解请先用“冰淇淋销量与溺水人数”的日常例子解释共同原因，'
                  '然后用一张概念对照表区分相关、因果和混杂因素，并解释外部效度。不要布置习题。'))
        events = [json.loads(chunk.decode()) async for chunk in chat.body_iterator]
        assert not any(item['type'] == 'error' for item in events), events
        conversation_id = next(item['value'] for item in events if item['type'] == 'conversation')
        answer = one("SELECT content FROM messages WHERE conversation_id=? AND role='assistant' ORDER BY created_at DESC", (conversation_id,))['content']
        (root / '导师答复.md').write_text(answer, encoding='utf-8')
        print(j({'stage': 'generation', 'reading_id': reading['reading_id'], 'conversation_id': conversation_id}), flush=True)
        value = server.GenerateIn(project_id=project['id'], source_ids=[source_id], reading_id=reading['reading_id'],
                    conversation_id=conversation_id, model=read.model, reasoning_effort=read.reasoning_effort,
                    generation_quality='custom', instructions='讲解尽量控制在 1200 字以内，正文不要出题。', summary_level=2)
        generation_job = create_job(project['id'], 'generate', value.model_dump())
        await server._generate_job(generation_job, value)
        job = one('SELECT * FROM jobs WHERE id=?', (generation_job,))
        assert job['status'] == 'completed', job['error']
        artifact = one("SELECT * FROM artifacts WHERE kind='explanation'")
        handoff = json.loads(artifact['metadata_json'])['mentor_handoff']
        assert handoff['status'] == 'completed' and handoff['mentor']['model'] == read.model
        assert handoff['mentor_thread_id'] == one('SELECT remote_thread_id FROM conversations WHERE id=?', (conversation_id,))['remote_thread_id']
        misconceptions = [item for item in handoff['packet']['understanding'] if item['status'] == 'misconception']
        assert misconceptions and any('冰淇淋' in evidence['quote'] for item in misconceptions for evidence in item['evidence']), handoff
        (root / '导师交接自检.json').write_text(j(handoff), encoding='utf-8')
        document = Document(artifact['path'])
        text = '\n'.join(p.text for p in document.paragraphs) + '\n'.join(c.text for table in document.tables for row in table.rows for c in row.cells)
        assert '冰淇淋' in text and '外部效度' in text and document.tables and '系统性总结' in text
        print(j({'ok': True, 'model': read.model, 'tables': len(document.tables),
                 'characters': len(text), 'misconceptions': len(misconceptions), 'handoff_verified': True,
                 'artifact': artifact['path'], 'teacher_answer': str(root / '导师答复.md')}), flush=True)
    finally:
        await bridge.close()


if __name__ == '__main__':
    asyncio.run(main())
