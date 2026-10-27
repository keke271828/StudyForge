"""Bounded live teaching samples for multiple roles. No simulated quality score.

All sources, attempts and outputs live in --output/data; real learning data is
never imported. The selected provider uses its existing account credentials.
Run only with --live and an explicit provider/model. Read the output before
claiming that teaching quality passed: source matching is not semantic proof.
"""
from __future__ import annotations

import argparse
import asyncio
import json
import os
import time
from pathlib import Path

TUTORS = {
    'tutor_math': ('柯西不等式的证明动机', '我知道平方和非负，但不懂证明柯西不等式时为什么考虑sum((a_i-t*b_i)^2)。请解释目标、t的作用和怎么想到的，处理b全为0的情形。不要出题，控制在500汉字左右。',
                   ['解释构造与目标差值的关系','处理b全为0，不除以0','不把一个构造说成唯一方法']),
    'tutor_physics': ('高斯定律的推理边界', '一个闭曲面内总电荷为0，我认为这说明曲面上电场处处为0。我的推理哪里错？请用一个已经讲完的具体例子说明，不出题，不强行比喻，500字以内。',
                      ['区分净通量为0与局部电场为0','有效例子符合闭合面与外法向','不把没有内电荷等同没有电场']),
    'tutor_social': ('相关、因果与外部效度', '观察到夏天冰淇淋销量与溺水人数一起上升，所以吃冰淇淋导致溺水；我觉得只要加大样本就能排除共同原因，还能推广到所有人。请指出我的推理缺口，最多500字，不布置习题。',
                     ['大样本不自动消除混杂','共同原因示例是可能解释而非已证事实','外部效度与内部因果识别不同']),
    'tutor_biology': ('负反馈与稳定性', '输出偏高就抑制、偏低就促进，这不是必然稳定吗？解释负反馈为什么不保证稳定，说明时延或响应强度怎样影响；没有具体生物系统资料，不要编造机制。400字以内，不出题。',
                      ['反馈方向不等于稳定性','正确说明时延或过强响应可能振荡','不凭结构相似认定机制相同']),
}
LECTURES = {
    'lecture_math': ('配方证明与条件', '设a、b为实数。由(a-b)^2>=0，展开再移项，得到a^2+b^2>=2ab；等号当且仅当a=b。选取平方是因为展开式里有目标中的交叉项，这提供一种构造线索，并不证明该方法唯一。',
                     ['实数条件与等号条件','构造动机，不声称唯一','引用真实，末尾总结，不追加习题']),
    'lecture_philosophy': ('描述与规范的论证', '某种做法很普遍，这是描述性判断。仅凭普遍性，不能推出它应当被接受；规范结论还需要评价原则。材料不主张任何普遍做法都不可接受，也没有说明这个区分的首创者。',
                           ['描述性与规范性准确区分','缺少前提不等于结论已经被证伪','不补造首创者或强套物理']),
    'lecture_chemistry': ('给定速率模型', '本资料只假定一个反应A+B→C满足v=k[A][B]。浓度单位为mol/L，速率单位为mol/(L*s)，所以k的单位为L/(mol*s)。这是给定模型，不能从反应方程式推出所有反应都服从这个速率式。',
                          ['模型中各量与单位','反应方程式不自动决定速率级数','解释单位推导，不越界推广']),
}
OTHER_CASES = ['exercise_transfer','goal_checkpoint_social','grading_equivalent','feynman_misconception','remediation_evidence','coding_feedback']


class Metered:
    def __init__(self, client, calls):self.client,self.calls=client,calls
    def __getattr__(self, name):return getattr(self.client,name)
    async def complete(self, system, user, **kwargs):
        started=time.perf_counter()
        result=await self.client.complete(system,user,**kwargs)
        record={'seconds':round(time.perf_counter()-started,2),'json':bool(kwargs.get('json_mode'))}
        if kwargs.get('json_mode'):record['fixture_response']=result  # Synthetic benchmark data only, never user history.
        self.calls.append(record)
        return result
    def fork(self):return Metered(self.client.fork(),self.calls)
    async def close(self):await self.client.close()


async def evaluate(args):
    output=Path(args.output).resolve();output.mkdir(parents=True,exist_ok=True)
    os.environ['STUDYFORGE_DATA_DIR']=str(output/'data')
    from studyforge import ai,coding,exercises,homework
    from studyforge.db import create_project,execute,j,now,uid
    from studyforge.teaching import TEACHER_PERSONA,CHAT_RULES,TEACHING_TASK_REMINDER,finish_explanation
    from studyforge.teaching_styles import style_policy,language_policy
    from studyforge.learning_modes import mode_policy
    from studyforge.exporters import export_document
    from studyforge.lesson_quality import citation_checks
    from studyforge.parsers import parse_file
    from docx import Document
    ai.save_settings({'chat_provider':args.provider,'agent_tools_enabled':False,'agent_network_enabled':False,
                      'agent_web_video_enabled':False,'background_learning_enabled':False,'output_language':'zh'})
    project=create_project('多角色真实教学抽样')
    rows=[];calls=[]
    async def no_external_sources(*a,**kw):return []
    # This benchmark tests source-bounded exercise design and private reviews,
    # not web search quality or Docker execution. Restore patched boundaries.
    old_search,old_tests=exercises.search_exercises,coding._run_assignment_tests
    exercises.search_exercises=no_external_sources
    coding._run_assignment_tests=lambda *a,**kw:{'ran':False,'reason':'本次抽样未启动Docker，不能声称测试已执行'}
    def new_client():
        if args.provider=='codex':
            from studyforge.codex_bridge import CodexBridge
            return Metered(ai.CodexClient(private_bridge=CodexBridge(),model=args.model,effort=args.effort),calls)
        return Metered(ai.client_for(args.provider,model=args.model,effort=args.effort),calls)
    def set_fixture(title, question, answer, rubric, status='assigned', language=None):
        ident=uid();q={'index':0,'question':question,'topic':title}
        if language:q['language']=language
        execute('INSERT INTO exercise_sets VALUES (?,?,?,?,?,?,?,?,?)',
                (ident,project['id'],'[]',title,j([q]),j([{'index':0,'solution':answer,'rubric':rubric,'test_code':''}]),1,status,now()))
        return ident
    def docx(title,text,name):
        path=export_document(text,'docx',title,name+'.docx')
        read=Document(path)
        return str(path),{'paragraphs':len(read.paragraphs),'tables':len(read.tables),'readable':True}
    try:
        for name in args.case or list(TUTORS)+list(LECTURES)+OTHER_CASES:
            client=new_client();started=time.perf_counter();before=len(calls)
            async def sample():
                if name in TUTORS:
                    title,question,criteria=TUTORS[name]
                    text=await client.complete(TEACHER_PERSONA+CHAT_RULES+style_policy('clear')+language_policy('zh')+mode_policy('deep'),question+TEACHING_TASK_REMINDER)
                    path,inspection=docx(title,text,name)
                    return {'title':title,'text':text,'document':path,'inspection':inspection,'human_review_criteria':criteria}
                if name in LECTURES:
                    title,original,criteria=LECTURES[name]
                    path=output/(name+'.txt');path.write_text(original,encoding='utf-8')
                    segments=parse_file(path)
                    digest='\n'.join(f'[来源 {path.name}｜{x["locator"]}]\n{x["text"]}' for x in segments)
                    text=await ai.generate_explanation(client,digest,3,[],requirements='只使用本次原文；讲清概念、方法动机与适用条件，400—700汉字，不出题，不写发现史。')
                    text=await finish_explanation(client,text)
                    file,inspection=docx(title,text,name)
                    return {'title':title,'text':text,'document':file,'inspection':inspection,
                            'source_checks':citation_checks(text,[{'name':path.name,'segments':segments}]),'human_review_criteria':criteria}
                if name=='exercise_transfer':
                    ident,reports=await exercises.generate_set(client,project['id'],[],
                        '材料：相关并不足以证明因果。可能有共同原因，同时影响两变量；样本量增大不自动排除这种偏差。外部效度涉及其他人群与条件。',
                        '因果辨析与迁移',2,requirements='只用本材料，原创2题，分别检测共同原因与适用范围，题目不泄漏答案。')
                    text=exercises.student_document(ident)
                    path,inspection=docx('迁移题学生版',text,name)
                    return {'exercise_set_id':ident,'document':path,'text':text,'inspection':inspection,'private_reviews':reports,
                            'human_review_criteria':['改变应用结构而非只改数字','题意确定且私有解经真实求解复核','学生版无答案']}
                if name=='grading_equivalent':
                    ident=set_fixture('等价证明','证明对实数a、b，有a^2+b^2>=2ab，并给出等号条件。',
                        '由(a-b)^2>=0展开，等号当且仅当a=b。','非负依据40分，等价变形40分，等号条件20分。允许等价正确证明。')
                    result=await exercises.grade_submission(client,ident,['令u=a+b、v=a-b。由v^2>=0，a^2+b^2-2ab=v^2>=0，所以结论成立。等号当且仅当v=0，即a=b。'])
                    return {'grade':result,'human_review_criteria':['等价证明应得100分','不能因为引入u但未用到而扣正确性分','未另行安排未请求习题']}
                if name=='goal_checkpoint_social':
                    from studyforge import learning_goals
                    from studyforge.db import one,unj
                    original='描述性判断说明事实是什么，规范性判断讨论应该怎样。仅凭一种做法很普遍，不能推出它应当被接受，规范结论需要评价原则。这不意味着普遍做法都不值得接受。本材料没有说明该区分的首创者。'
                    path=output/'描述与规范.txt';path.write_text(original,encoding='utf-8');source=uid()
                    execute('INSERT INTO sources(id,project_id,name,kind,original_path,segments_json,status,created_at) VALUES (?,?,?,?,?,?,?,?)',
                        (source,project['id'],path.name,'.txt',str(path),'[]','ready',now()))
                    goal=learning_goals.save(project['id'],{'name':'描述与规范','domain':'哲学','intent':'deep','source_ids':[source],
                        'criterion':'说明事实与评价的区别，解释规范结论缺少的前提，在陌生但资料足以支持的情境中指出论证边界。'})
                    spec=learning_goals.check_spec(learning_goals.get(goal['id'],project['id']))
                    digest='\n'.join(f'[来源 {path.name}｜{s["locator"]}]\n{s["text"]}' for s in parse_file(path))
                    ident,reports=await exercises.generate_set(client,project['id'],[source],digest,'描述与规范理解检查',
                        requirements='只用给定材料，题面精简；私有解说明理由，评分标准明确。',check_spec=spec)
                    learning_goals.register_check(ident,project['id'],spec)
                    public=exercises.public_set(ident);text=exercises.student_document(ident);file,inspection=docx('三题理解检查',text,name)
                    return {'goal':goal,'exercise_set_id':ident,'student_questions':public['questions'],'document':file,'text':text,
                        'private_solutions':unj(one('SELECT solutions_json FROM exercise_sets WHERE id=?',(ident,))['solutions_json']),
                        'private_reviews':reports,'inspection':inspection,'human_review_criteria':[
                            '三题分别检测概念、论证理由和有意义的迁移边界','题面不提前给结论或完整解析',
                            '不编造理论首创者，不强套物理','私有解析符合给定事实且明确说明评价前提，不把不足以推出说成结论已被证伪']}
                if name=='feynman_misconception':
                    ident=set_fixture('高斯定律','说明闭曲面净电通量为0是否意味着面上电场处处为0。',
                        '不意味着。净通量为零可以来自正负通量抵消，均匀非零电场通过闭曲面的净通量为零。','区分积分约束与逐点电场',status='feynman_required')
                    explanation='我认为净电通量是电场的总和。面内电荷为零就意味着积分为零，所以曲面上每一点电场都为零。无论外部有没有电荷都这样，因为零的积分只能由处处为零组成。'
                    result=await exercises.grade_feynman(client,ident,explanation)
                    return {'grade':result,'human_review_criteria':['不能通过80分线','指出积分为零不意味着逐点为零','不把流畅复述当正确理解']}
                if name=='remediation_evidence':
                    q=[{'index':0,'question':'说明封闭面净电通量为零的含义。','topic':'高斯定律'},
                       {'index':1,'question':'计算2+3，并写出步骤。','topic':'加法'}]
                    row={'answers_json':j(['0','2+3=4']),'options_json':j({'teaching':{'explanation_style':'clear','output_language':'zh','learning_mode':'deep'},
                        'profile':{'observations':[{'topic':'高斯定律','observation':'学生两个月前询问过高斯定律是什么，尚无当前掌握证据。'}]}})}
                    worksheet={'questions_json':j(q),'solutions_json':j([{'index':0,'solution':'净通量零不是电场处处零。'},{'index':1,'solution':'2+3=5'}])}
                    analysis=await homework.analyze(client,row,worksheet,{'scores':[0,30],'feedback':['缺少说明','加法算错']},[])
                    text='\n\n'.join(f'## 第{x["index"]+1}题\n状态：{x["confidence"]}\n{x["cause"]}\n{x["explanation"]}\n核对：{x["check"]}\n待询问：{x.get("clarification","")}' for x in analysis['items'])
                    path,inspection=docx('错因与补讲抽样',text,name)
                    return {'analysis':analysis,'text':text,'document':path,'inspection':inspection,
                            'human_review_criteria':['仅最终答案0不能诊断不了解概念，先提问','旧画像只是线索','2+3=4可定位计算错误，示范更正与核对']}
                if name=='coding_feedback':
                    ident=set_fixture('返回两数之和','实现Python函数add(a,b)，合法输入为整数，返回a+b。',
                        'def add(a,b):\n    return a+b','只按合法整数输入返回两数之和，不要求类型检查',language='python')
                    result=await coding.grade_code(client,ident,'def add(a,b):\n    return a-b')
                    return {'grade':result,'human_review_criteria':['指出减法与求和目标不符','说明修正与可检查的例子','测试未运行时不假称通过、不定最终分']}
                raise ValueError('未知样例：'+name)
            try:
                value=await asyncio.wait_for(sample(),args.timeout*(4 if name in ('exercise_transfer','goal_checkpoint_social') else 1))
                row={'case':name,'status':'completed',**value,'semantic_review':None}
            except Exception as exc:row={'case':name,'status':'failed','error':str(exc) or type(exc).__name__,'validation_cause':str(exc.__cause__) if exc.__cause__ else None,'semantic_review':None}
            finally:await client.close()
            row.update(seconds=round(time.perf_counter()-started,2),model_requests=len(calls)-before)
            rows.append(row)
            (output/'report.json').write_text(json.dumps({'provider':args.provider,'model':args.model,'effort':args.effort or 'service_default',
                 'simulated':False,'real_learning_data_used':False,'cases':rows,'model_requests':len(calls),'calls':calls,
                 'notice':'教学准确性、方法解释和适配性需逐份人工审阅；流程完成和来源匹配不是教学质量认证。'},ensure_ascii=False,indent=2),encoding='utf-8')
            print(json.dumps({'case':name,'status':row['status'],'seconds':row['seconds'],'model_requests':row['model_requests'],'error':row.get('error')},ensure_ascii=False),flush=True)
    finally:
        exercises.search_exercises,coding._run_assignment_tests=old_search,old_tests
    return output/'report.json'


def main():
    parser=argparse.ArgumentParser()
    parser.add_argument('--live',action='store_true')
    parser.add_argument('--provider',required=True)
    parser.add_argument('--model',required=True)
    parser.add_argument('--effort')
    parser.add_argument('--output',required=True)
    parser.add_argument('--case',action='append',choices=list(TUTORS)+list(LECTURES)+OTHER_CASES)
    parser.add_argument('--timeout',type=int,default=180)
    args=parser.parse_args()
    if not args.live:parser.error('明确指定--live才会消耗实际模型额度；不存在模拟教学质量分数')
    print(asyncio.run(evaluate(args)),flush=True)


if __name__=='__main__':main()
