"""Reopen real browser downloads without PowerShell's legacy stdin encoding."""
import os
from pathlib import Path

from docx import Document


def main():
    root = Path(os.environ['STUDYFORGE_MOBILE_CASE']).parent
    text = '\n'.join(p.text for p in Document(root / 'downloaded.docx').paragraphs)
    assert '系统性总结' in text and '原文' in text
    solutions = '\n'.join(p.text for p in Document(root / 'homework-solutions.docx').paragraphs)
    feedback = '\n'.join(p.text for p in Document(root / 'homework-feedback.docx').paragraphs)
    code = '\n'.join(p.text for p in Document(root / 'coding-feedback.docx').paragraphs)
    assert 'PHONE_REFERENCE_VARIABLE' in solutions and 'PHONE_REFERENCE_ADD' in solutions
    assert '错因分析' in feedback and '针对性讲解' in feedback and '变量符号' in feedback
    assert '中文代码说明' in code and '没有最终分数' in code
    print('Downloaded explanation, solutions, teaching feedback and code feedback reopened successfully')


if __name__ == '__main__':
    main()
