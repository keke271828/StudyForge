from __future__ import annotations

import html
import base64
import json
import copy
import re
import sqlite3
import tempfile
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path, PurePosixPath
from typing import Callable

from .media import transcribe, video_frames, video_frame_timestamps, video_metadata_path
from .config import derived_dir

Segment = dict[str, object]
Parser = Callable[[Path], list[Segment]]
PARSERS: dict[str, Parser] = {}


def _safe_zip_read(archive: zipfile.ZipFile, name: str, max_bytes: int = 100 * 1024 * 1024) -> bytes:
    info = archive.getinfo(name)
    if info.file_size > max_bytes:
        raise ValueError("压缩包内部文件过大，已停止解析")
    return archive.read(name)


def register(*extensions: str):
    def decorator(func: Parser) -> Parser:
        for ext in extensions:
            PARSERS[ext.lower()] = func
        return func
    return decorator


def segment(locator: str, text: str, kind: str = "text", **meta: object) -> Segment:
    return {"locator": locator, "text": text.strip(), "kind": kind, "meta": meta}


def _clean_html(value: str) -> str:
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(value, "html.parser")
    for tag in soup.select('script, style, nav, footer, noscript'):
        tag.decompose()
    article = soup.select_one('.Post-RichTextContainer .RichText, article, .RichText.ztext, main') or soup
    for image in article.find_all('img'):
        label = image.get('data-tex') or image.get('data-formula') or image.get('alt')
        image.replace_with(f" 公式或图像：{label} " if label else ' [图像：需结合原网页核对] ')
    return html.unescape(article.get_text("\n", strip=True))


def decode_text(data: bytes) -> tuple[str, str]:
    """Decode without replacing unreadable bytes with silent U+FFFD glyphs."""
    if data.startswith((b'\xff\xfe', b'\xfe\xff')):
        return data.decode('utf-16'), 'utf-16'
    if data.startswith(b'\xef\xbb\xbf'):
        return data.decode('utf-8-sig'), 'utf-8-sig'
    # UTF-16 documents without a BOM are common Windows exports.
    if b'\x00' in data[:400]:
        sample = data[:400]
        encodings = ('utf-16-be', 'utf-16-le') if sample[::2].count(0) > sample[1::2].count(0) else ('utf-16-le', 'utf-16-be')
        for encoding in encodings:
            try:
                text = data.decode(encoding)
                if '\x00' not in text:
                    return text, encoding
            except UnicodeError:
                pass
    for encoding in ('utf-8', 'gb18030', 'big5'):
        try:
            return data.decode(encoding), encoding
        except UnicodeError:
            pass
    raise ValueError('无法可靠识别文件编码，请另存为 UTF-8 后导入；已保留原文件')


def _ocr(path: Path) -> str:
    try:
        from rapidocr import RapidOCR
        from .local_compute import cached_model, cpu_threads, inference_slot
        with inference_slot():
            engine = cached_model('ocr', 'rapidocr', lambda: RapidOCR(params={
                'EngineConfig.onnxruntime.intra_op_num_threads': cpu_threads(),
                'EngineConfig.onnxruntime.inter_op_num_threads': 1,
            }))
            result = engine(str(path))
        if hasattr(result, "txts"):
            return "\n".join(result.txts or [])
        if isinstance(result, tuple):
            return "\n".join(row[1] for row in (result[0] or []))
    except Exception as exc:
        return f"[OCR 未完成：{type(exc).__name__}: {exc}]"
    return ""


@register(".txt", ".md", ".markdown", ".csv", ".tsv", ".html", ".htm", ".json", ".tex")
def parse_text(path: Path) -> list[Segment]:
    value, encoding = decode_text(path.read_bytes())
    if path.suffix.lower() in (".html", ".htm"):
        value = _clean_html(value)
    lines = value.splitlines()
    blocks: list[Segment] = []
    for start in range(0, len(lines), 100):
        text = "\n".join(lines[start:start + 100]).strip()
        if text:
            blocks.append(segment(f"行 {start + 1}-{min(start + 100, len(lines))}", text, encoding=encoding))
    return blocks


CODE_EXTENSIONS = (".py", ".js", ".ts", ".tsx", ".jsx", ".cpp", ".c", ".h", ".hpp",
                   ".java", ".rs", ".go", ".cs", ".r", ".jl", ".sql", ".sh", ".ps1",
                   ".m", ".wl", ".wls", ".nb", ".mdl")


@register(*CODE_EXTENSIONS)
def parse_code(path: Path) -> list[Segment]:
    return [dict(item, kind="code", meta={"language": path.suffix.lstrip(".")})
            for item in parse_text(path)]


@register(".ipynb")
def parse_notebook(path: Path) -> list[Segment]:
    notebook = json.loads(path.read_text(encoding="utf-8-sig"))
    parts = []
    for i, cell in enumerate(notebook.get('cells', [])):
        parts.append(segment(f'单元 {i + 1}', ''.join(cell.get('source', [])), cell.get('cell_type', 'code'),
                             language=notebook.get('metadata', {}).get('kernelspec', {}).get('language', 'python')))
        for j, output in enumerate(cell.get('outputs', [])):
            data = output.get('data', {})
            value = output.get('text') or data.get('text/latex') or data.get('text/plain') or output.get('traceback', [])
            if value:
                parts.append(segment(f'单元 {i + 1} 输出 {j + 1}', ''.join(value), 'output'))
            if data.get('image/png'):
                folder = derived_dir(path, 'notebook')
                image = folder / f'{i}-{j}.png'
                image.write_bytes(base64.b64decode(''.join(data['image/png'])))
                parts.append(segment(f'单元 {i + 1} 图 {j + 1}', '', 'image', image_path=str(image)))
    return parts


@register(".pdf")
def parse_pdf(path: Path) -> list[Segment]:
    import fitz
    parts = []
    with fitz.open(path) as pdf:
        for index, page in enumerate(pdf):
            value = page.get_text("text").strip()
            image_path = ""
            has_visuals = bool(page.get_images(full=True) or page.get_drawings())
            if has_visuals or len(value) < 30:
                folder = derived_dir(path, 'pdf-pages')
                image = folder / f"page-{index + 1}.png"
                page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5)).save(image)
                image_path = str(image)
            if len(value) < 30:
                value = _ocr(Path(image_path))
            parts.append(segment(f"第 {index + 1} 页", value, "page", image_path=image_path))
    return parts


@register(".docx")
def parse_docx(path: Path) -> list[Segment]:
    from docx import Document
    from docx.oxml.ns import qn
    doc = Document(path)
    parts = []
    def text(node):
        kind = node.tag.rsplit('}', 1)[-1]
        if kind == 'oMath':
            return '$' + omml_latex(node) + '$'
        if kind == 't':
            return node.text or ''
        if kind in ('tab', 'br', 'cr'):
            return '\n' if kind != 'tab' else '\t'
        return ''.join(text(child) for child in node)

    def read_body(root, prefix, container_part=None):
        for index, node in enumerate(root):
            if node.tag not in (qn('w:p'), qn('w:tbl')):
                continue
            locator = f'{prefix} {index + 1}'
            if node.tag == qn('w:tbl'):
                value = '\n'.join(' | '.join(text(cell) for cell in row.findall(qn('w:tc')))
                                  for row in node.findall(qn('w:tr')))
            else:
                value = text(node)
            if value.strip():
                parts.append(segment(locator, value, 'table' if node.tag == qn('w:tbl') else 'paragraph'))
            for image_index, image_node in enumerate(node.iter(qn('a:blip'))):
                rid = image_node.get(qn('r:embed'))
                part = (container_part or doc.part).related_parts.get(rid)
                if part:
                    folder = derived_dir(path, 'docx-images')
                    image = folder / f'{prefix}-{index}-{image_index}{Path(str(part.partname)).suffix}'
                    image.write_bytes(part.blob)
                    parts.append(segment(f'{locator} 图 {image_index + 1}', '', 'image', image_path=str(image)))
    read_body(doc.element.body, '正文位置')
    for i, section in enumerate(doc.sections):
        read_body(section.header._element, f'节 {i + 1} 页眉', section.header.part)
        read_body(section.footer._element, f'节 {i + 1} 页脚', section.footer.part)
    from lxml import etree
    with zipfile.ZipFile(path) as archive:
        for filename, label in [('word/footnotes.xml', '脚注'), ('word/endnotes.xml', '尾注')]:
            if filename in archive.namelist():
                root = etree.fromstring(_safe_zip_read(archive, filename))
                for note in root:
                    if int(note.get(qn('w:id'), '0')) >= 0:
                        read_body(note, f"{label} {note.get(qn('w:id'))}")
    return parts


def omml_latex(node) -> str:
    """Preserve original Word formulas in the text sent to the model."""
    kind = node.tag.rsplit('}', 1)[-1]
    if kind.endswith('Pr'):
        return ''
    if kind == 't':
        return node.text or ''
    children = {child.tag.rsplit('}', 1)[-1]: child for child in node}
    def part(name):
        return omml_latex(children[name]) if name in children else ''
    if kind == 'f':
        return '\\frac{' + part('num') + '}{' + part('den') + '}'
    if kind in ('sSup', 'sSub', 'sSubSup'):
        return '{' + part('e') + '}' + ('_{' + part('sub') + '}' if 'sub' in children else '') + ('^{' + part('sup') + '}' if 'sup' in children else '')
    if kind == 'rad':
        return '\\sqrt' + ('[' + part('deg') + ']' if part('deg') else '') + '{' + part('e') + '}'
    if kind == 'acc':
        accent = next((child.get('{http://schemas.openxmlformats.org/officeDocument/2006/math}val')
                       for child in node.iter() if child.tag.rsplit('}', 1)[-1] == 'chr'), '^')
        command = {'→': 'vec', '^': 'hat', 'ˆ': 'hat', '˙': 'dot', '¨': 'ddot', '~': 'tilde', '¯': 'bar', '‾': 'bar'}.get(accent, 'hat')
        return '\\' + command + '{' + part('e') + '}'
    if kind == 'm':
        rows = [' & '.join(omml_latex(cell) for cell in row if cell.tag.rsplit('}', 1)[-1] == 'e')
                for row in node if row.tag.rsplit('}', 1)[-1] == 'mr']
        return '\\begin{matrix}' + ' \\\\ '.join(rows) + '\\end{matrix}'
    return ''.join(omml_latex(child) for child in node)


@register(".pptx")
def parse_pptx(path: Path) -> list[Segment]:
    from pptx import Presentation
    deck = Presentation(path)
    parts = []
    for i, slide in enumerate(deck.slides):
        texts = []
        def read_shapes(shapes):
            for j, shape in enumerate(shapes):
                if shape.has_text_frame and shape.text.strip():
                    texts.append(shape.text)
                if shape.has_table:
                    texts.extend(' | '.join(cell.text for cell in row.cells) for row in shape.table.rows)
                if shape.has_chart:
                    texts.append('图表数据：' + '\n'.join(f'{series.name}: {list(series.values)}' for series in shape.chart.series))
                if hasattr(shape, 'shapes'):
                    read_shapes(shape.shapes)
                if hasattr(shape, 'image'):
                    folder = derived_dir(path, 'pptx-images')
                    image = folder / f'{i}-{shape.shape_id}.{shape.image.ext}'
                    image.write_bytes(shape.image.blob)
                    parts.append(segment(f'幻灯片 {i + 1} 图 {shape.shape_id}', '', 'image', image_path=str(image)))
        read_shapes(slide.shapes)
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            texts.append('演讲者备注：' + slide.notes_slide.notes_text_frame.text)
        parts.append(segment(f"幻灯片 {i + 1}", "\n".join(texts), "slide"))
    return parts


@register(".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff")
def parse_image(path: Path) -> list[Segment]:
    return [segment("图像", _ocr(path), "image", image_path=str(path))]


def _xml_label(node: ET.Element) -> str:
    return node.attrib.get("TEXT") or node.attrib.get("text") or node.attrib.get("title") or ""


@register(".mm")
def parse_freemind(path: Path) -> list[Segment]:
    root = ET.parse(path).getroot()
    parts: list[Segment] = []

    def visit(node: ET.Element, trail: list[str]) -> None:
        title = _xml_label(node) or next((_clean_html(ET.tostring(child, encoding='unicode'))
                                        for child in node.findall('richcontent') if child.get('TYPE') == 'NODE'), '')
        branch = trail + ([title] if title else [])
        details = [title]
        details.extend('备注：' + _clean_html(ET.tostring(child, encoding='unicode'))
                       for child in node.findall('richcontent') if child.get('TYPE') == 'NOTE')
        if node.get('LINK'): details.append('链接：' + node.get('LINK'))
        if any(details):
            parts.append(segment(" / ".join(branch) or '未命名节点', '\n'.join(details), "mindmap_node", depth=len(branch)))
        for child in node.findall("node"):
            visit(child, branch)

    for node in root.findall("node"):
        visit(node, [])
    return parts


@register(".opml")
def parse_opml(path: Path) -> list[Segment]:
    root = ET.parse(path).getroot()
    parts: list[Segment] = []

    def visit(node: ET.Element, trail: list[str]) -> None:
        title = _xml_label(node)
        branch = trail + ([title] if title else [])
        details = [title]
        for key in ('_note', 'description', 'url', 'htmlUrl', 'xmlUrl'):
            if node.get(key): details.append(f'{key}：{_clean_html(node.get(key))}')
        if any(details):
            parts.append(segment(" / ".join(branch) or '未命名节点', '\n'.join(details), "mindmap_node", depth=len(branch)))
        for child in node.findall("outline"):
            visit(child, branch)

    body = root.find("body")
    if body is not None:
        for node in body.findall("outline"):
            visit(node, [])
    return parts


@register(".xmind")
def parse_xmind(path: Path) -> list[Segment]:
    parts: list[Segment] = []
    with zipfile.ZipFile(path) as archive:
        if "content.json" in archive.namelist():
            sheets = json.loads(_safe_zip_read(archive, "content.json"))

            def visit(topic: dict, trail: list[str]) -> None:
                title = topic.get("title", "")
                branch = trail + ([title] if title else [])
                details = [title]
                for kind, note in (topic.get('notes') or {}).items():
                    content = note.get('content') if isinstance(note, dict) else note
                    if isinstance(content, str) and content.strip(): details.append('备注：' + _clean_html(content))
                if topic.get('labels'): details.append('标签：' + '、'.join(map(str, topic['labels'])))
                if topic.get('href'): details.append('链接：' + str(topic['href']))
                if any(details):
                    parts.append(segment(" / ".join(branch) or '未命名节点', '\n'.join(details), "mindmap_node", depth=len(branch)))
                children = topic.get("children", {})
                for collection in children.values():
                    for child in collection:
                        visit(child, branch)

            for sheet in sheets:
                visit(sheet.get("rootTopic", {}), [sheet.get("title", "画布")])
        elif "content.xml" in archive.namelist():
            root = ET.fromstring(_safe_zip_read(archive, "content.xml"))
            local = lambda node: node.tag.rsplit("}", 1)[-1]

            def visit_xml(topic: ET.Element, trail: list[str]) -> None:
                title = next((child.text or "" for child in topic if local(child) == "title"), "")
                branch = trail + ([title] if title else [])
                details = [title]
                for child in topic:
                    if local(child) == 'notes':
                        details.append('备注：' + _clean_html(ET.tostring(child, encoding='unicode')))
                    elif local(child) == 'labels':
                        details.append('标签：' + '、'.join(''.join(item.itertext()) for item in child))
                link = next((value for key, value in topic.attrib.items() if key.rsplit('}', 1)[-1] == 'href'), None)
                if link: details.append('链接：' + link)
                if any(details):
                    parts.append(segment(" / ".join(branch) or '未命名节点', '\n'.join(details), "mindmap_node", depth=len(branch)))
                for container in topic:
                    if local(container) == "children":
                        for collection in container:
                            if local(collection) == "topic":
                                visit_xml(collection, branch)
                            else:
                                for child in collection:
                                    if local(child) == "topic":
                                        visit_xml(child, branch)

            for sheet in root.iter():
                if local(sheet) != "sheet":
                    continue
                sheet_title = next((child.text or "画布" for child in sheet if local(child) == "title"), "画布")
                for topic in sheet:
                    if local(topic) == "topic":
                        visit_xml(topic, [sheet_title])
        else:
            raise ValueError("XMind 文件没有可识别的内容")
    return parts


@register(".apkg", ".colpkg")
def parse_anki(path: Path) -> list[Segment]:
    from .anki_import import decompress, media_map
    from bs4 import BeautifulSoup
    from PIL import Image
    import io
    from urllib.parse import unquote
    parts = []
    with zipfile.ZipFile(path) as archive:
        entry = next((name for name in ('collection.anki21b', "collection.anki21", "collection.anki2") if name in archive.namelist()), None)
        if not entry:
            raise ValueError("Anki 卡组缺少 collection 数据库")
        with tempfile.TemporaryDirectory() as folder:
            db_path = Path(folder) / "collection.sqlite3"
            content = _safe_zip_read(archive, entry)
            db_path.write_bytes(decompress(content) if entry == 'collection.anki21b' else content)
            db = sqlite3.connect(db_path)
            try:
                rows = db.execute("SELECT id, flds, tags FROM notes ORDER BY id").fetchall()
            finally:
                db.close()
        mapping = media_map(_safe_zip_read(archive, 'media'), entry == 'collection.anki21b') if 'media' in archive.namelist() else {}
        by_name = {name: number for number, name in mapping.items()}
        for index, (note_id, fields, tags) in enumerate(rows):
            raw_fields = fields.split('\x1f')
            clean = [_clean_html(item) for item in raw_fields]
            locator = f'卡片 {index + 1}'
            parts.append(segment(locator, '\n'.join(clean), 'anki_note',
                                 note_id=note_id, fields=clean, raw_fields=raw_fields, tags=tags.strip().split(),
                                 field_order='原笔记字段顺序；自定义模板与填空卡不等同于固定正反面'))
            for field_index, field in enumerate(raw_fields):
                for image_index, image in enumerate(BeautifulSoup(field, 'html.parser').find_all('img')):
                    name = unquote(str(image.get('src', '')))
                    label = f'{locator} / 字段 {field_index + 1} / 图片 {image_index + 1}'
                    item = by_name.get(name)
                    if item is None or item not in archive.namelist():
                        parts.append(segment(label, '[图片缺失或外部图像无法确认]', 'notice'))
                        continue
                    payload = _safe_zip_read(archive, item)
                    if entry == 'collection.anki21b': payload = decompress(payload)
                    target = derived_dir(path, 'anki-images') / f'anki-{index}-{field_index}-{image_index}.png'
                    try:
                        with Image.open(io.BytesIO(payload)) as source:
                            source.convert('RGB').save(target)
                    except (OSError, ValueError):
                        parts.append(segment(label, '[此图像格式无法解析，保留原卡组核对]', 'notice'))
                    else:
                        parts.append(segment(label, image.get('alt') or '[卡片图片，需视觉核对]',
                                             'image', note_id=note_id, image_path=str(target)))
    return parts


@register(".slx")
def parse_simulink(path: Path) -> list[Segment]:
    parts = []
    with zipfile.ZipFile(path) as archive:
        for name in archive.namelist():
            if name.startswith("simulink/") and name.endswith(".xml"):
                root = ET.fromstring(_safe_zip_read(archive, name))
                for item in root.iter():
                    if item.tag.endswith("Block"):
                        label = item.attrib.get("Name", "Block")
                        props = [f"{x.attrib.get('Name', '')}: {x.text or ''}" for x in item if x.tag.endswith("P")]
                        parts.append(segment(f"{name} / {label}", "\n".join(props), "model_block"))
    return parts


@register(".mp3", ".wav", ".m4a", ".flac", ".ogg")
def parse_audio(path: Path) -> list[Segment]:
    return [segment(f"{p['start']:.1f}-{p['end']:.1f} 秒", p["text"], "transcript",
                    start=p["start"], end=p["end"]) for p in transcribe(path)]


@register(".mp4", ".mov", ".mkv", ".webm", ".avi")
def parse_video(path: Path, *, frame_strategy: str = 'dense', prepared_transcript=None) -> list[Segment]:
    import av
    from PIL import Image, ImageChops, ImageStat
    with av.open(str(path)) as container:
        has_audio = bool(container.streams.audio)
    captions = []
    metadata = video_metadata_path(path)
    if metadata.is_file():
        try:
            data = json.loads(metadata.read_text(encoding='utf-8'))
            if data.get('media_bytes') == path.stat().st_size:
                captions = data.get('subtitles', [])
        except (OSError, ValueError):
            pass
    if frame_strategy not in ('dense', 'smart'): raise ValueError('未知视频画面方案')
    parts = copy.deepcopy(prepared_transcript) if prepared_transcript is not None else [segment(f"{p['start']:.1f}-{p['end']:.1f} 秒", p['text'], 'transcript',
                     start=p['start'], end=p['end'], method='video_subtitle') for p in captions] if captions else (
        parse_audio(path) if has_audio else [segment("音轨", "[视频没有可转录的音轨]", "notice")])
    previous_image = None
    previous_ocr = ""
    from .task_timing import timed_call, phase
    frames = (timed_call('frame_extract', '提取视频采样画面', video_frames, path, scene_threshold=.08, max_frames=4320)
              if frame_strategy == 'smart' else timed_call('frame_extract', '提取视频采样画面', video_frames, path))
    timestamps = video_frame_timestamps(frames)
    from .video_sampling import KeyframeSelector
    selector = KeyframeSelector() if frame_strategy == 'smart' else None
    from .frame_ocr import ExactFrameOCR
    frame_ocr = ExactFrameOCR(_ocr)
    ocr_frames = 0
    with phase('video_ocr', '采样画面 OCR（含共享推理等待）'):
        for index, frame in enumerate(frames):
            if selector and not selector.retain(frame, timestamps.get(frame.name), final=index == len(frames) - 1):
                continue
            ocr = frame_ocr.read(frame)
            ocr_frames += 1
            with Image.open(frame) as source:
                reduced = source.convert("L").resize((64, 36))
            difference = ImageStat.Stat(ImageChops.difference(reduced, previous_image)).mean[0] if previous_image else 100
            if selector or previous_image is None or difference > 12 or ocr.strip() != previous_ocr.strip() or index % 6 == 0:
                stamp = timestamps.get(frame.name)
                label = f'约 {stamp:g} 秒画面' if stamp is not None else f'第 {index + 1} 帧（时间无法确认）'
                parts.append(segment(label, ocr or "[画面含非文字内容，需视觉分析]",
                                     "video_frame", time=stamp, image_path=str(frame)))
                previous_ocr = ocr
            previous_image = reduced
    cap = 4320 if selector else 2160
    if len(frames) >= cap:
        parts.append(segment("画面采样上限", f"[已达到 {cap} 帧上限；更长视频的后续画面需要分段导入]", "notice"))
    from .video_records import recognized_parts
    parts = recognized_parts(sorted(parts, key=lambda item: item["meta"].get("start", item["meta"].get("time", 0))))
    if parts:
        parts[0].setdefault('meta', {})['video_sampling'] = {
            'strategy': frame_strategy, 'candidate_frames': len(frames), 'ocr_frames': ocr_frames,
            'ocr_requests': frame_ocr.requests, 'ocr_reused_frames': frame_ocr.reused,
            'retained_frames': sum(p['kind'] == 'video_frame' for p in parts),
            'limitation': '画面采用采样识别，未逐帧观看；智能合并尚待真实课程对照验收'}
    return parts


@register('.sfvideo')
def parse_video_record(path: Path) -> list[Segment]:
    from .video_records import read_record
    return read_record(path)


@register('.zip')
def parse_code_archive(path: Path) -> list[Segment]:
    """Read exported code packages as source data; never extract or run code."""
    parts = []
    with zipfile.ZipFile(path) as archive:
        entries = [item for item in archive.infolist() if not item.is_dir()]
        if len(entries) > 300 or sum(item.file_size for item in entries) > 30 * 1024 * 1024:
            raise ValueError('代码包过大，请拆分后引用')
        seen = set()
        for entry in entries:
            name = PurePosixPath(entry.filename.replace('\\', '/'))
            if name.is_absolute() or '..' in name.parts or ':' in str(name) or str(name) in seen:
                raise ValueError('代码包含无效或重复的文件路径')
            seen.add(str(name))
            content = _safe_zip_read(archive, entry.filename, 5 * 1024 * 1024)
            try:
                text, encoding = decode_text(content)
                if '\x00' in text:
                    raise ValueError('二进制内容')
            except (UnicodeError, ValueError):
                parts.append(segment(str(name), '[此文件不能按文本解析，需单独打开核对]', 'notice'))
                continue
            if name.suffix.lower() == '.ipynb':
                notebook = json.loads(text)
                for index, cell in enumerate(notebook.get('cells', []), 1):
                    value = cell.get('source', '')
                    parts.append(segment(f'{name} · 单元 {index}', ''.join(value) if isinstance(value, list) else value,
                                         cell.get('cell_type', 'text'), archive_entry=str(name)))
            else:
                lines = text.splitlines()
                for start in range(0, len(lines), 100):
                    parts.append(segment(f'{name} · 行 {start + 1}—{min(start + 100, len(lines))}',
                                         '\n'.join(lines[start:start + 100]), 'code', encoding=encoding, archive_entry=str(name)))
    return parts


def parse_file(path: Path) -> list[Segment]:
    parser = PARSERS.get(path.suffix.lower())
    if not parser:
        raise ValueError(f"暂不支持解析 {path.suffix or '无扩展名'}；原文件已保留")
    segments = parser(path)
    if not segments:
        raise ValueError("没有提取到可用内容；原文件已保留")
    return segments
