"""Check the real frozen mobile gateway with isolated files and no AI calls."""
from __future__ import annotations

import json
import io
import os
from pathlib import Path
import secrets
import socket
import sqlite3
import ssl
import subprocess
import time
from urllib.parse import unquote
import uuid

import httpx
from cryptography import x509
from cryptography.hazmat.primitives import serialization
from docx import Document


def free_port() -> int:
    with socket.socket() as listener:
        listener.bind(('127.0.0.1', 0))
        return listener.getsockname()[1]


def checked(response: httpx.Response):
    response.raise_for_status()
    return response.json()


def main() -> None:
    repo = Path(__file__).resolve().parent.parent
    root = repo / 'build' / ('packaged-mobile-' + uuid.uuid4().hex)
    data = root / 'data'
    (data / 'mobile').mkdir(parents=True)
    (data / 'settings.json').write_text(json.dumps({'chat_provider': 'deepseek',
        'background_learning_enabled': False}), encoding='utf-8')
    desktop_port, phone_port = free_port(), free_port()
    while desktop_port == phone_port:
        phone_port = free_port()
    (data / 'mobile' / 'connection.json').write_text(json.dumps({'auto_start': True,
        'port': phone_port, 'enable_tailscale': True, 'remember_days': 365}), encoding='utf-8')
    environment = {**os.environ, 'STUDYFORGE_DATA_DIR': str(data),
        'STUDYFORGE_PORT': str(desktop_port), 'STUDYFORGE_PARENT_PID': str(os.getpid()),
        'STUDYFORGE_DESKTOP_TOKEN': secrets.token_hex(32), 'PYTHONUTF8': '1',
        'PATH': str(Path(os.environ.get('SystemRoot', 'C:/Windows')) / 'System32')}
    environment.pop('PYTHONPATH', None)
    environment.pop('PYTHONHOME', None)
    executable = repo / 'dist' / 'win-unpacked' / 'resources' / 'backend' / 'StudyForgeBackend.exe'
    assert executable.is_file(), executable
    report = {'ok': False, 'checks': [], 'root': str(root)}
    process = None
    phone = None
    with (root / 'backend.log').open('wb') as log, httpx.Client(
            base_url=f'http://127.0.0.1:{desktop_port}', trust_env=False, timeout=15,
            headers={'X-StudyForge-Desktop-Token': environment['STUDYFORGE_DESKTOP_TOKEN']}) as desktop:
        def start():
            nonlocal process
            process = subprocess.Popen([str(executable)], cwd=root, env=environment,
                stdout=log, stderr=log, creationflags=subprocess.CREATE_NO_WINDOW)
            for _ in range(150):
                if process.poll() is not None:
                    raise RuntimeError('Frozen backend exited: ' + str(process.returncode))
                try:
                    if checked(desktop.get('/api/health'))['ok']:
                        return checked(desktop.get('/api/mobile/status'))
                except httpx.HTTPError:
                    time.sleep(.2)
            raise RuntimeError('Frozen mobile startup timed out')

        def stop():
            nonlocal process
            if process and process.poll() is None:
                try:
                    checked(desktop.post('/api/desktop/shutdown'))
                finally:
                    process.terminate()
                    process.wait(timeout=15)
            process = None

        try:
            status = start()
            assert status['running'] and status['auto_start'], status.get('error')
            assert status['gateway_version'] == 2
            report['checks'].append('frozen backend restores saved HTTPS mobile service without Python or Node in PATH')
            ca_file = data / 'mobile' / 'tls' / 'studyforge-mobile-ca.cer'
            ca = ca_file.read_bytes()
            context = ssl.create_default_context(cadata=x509.load_der_x509_certificate(ca).public_bytes(
                serialization.Encoding.PEM).decode('ascii'))
            phone = httpx.Client(base_url=f'https://127.0.0.1:{phone_port}', verify=context,
                                 trust_env=False, timeout=15)
            for route in ('/', '/app.js', '/style.css', '/homework.js', '/lesson-feedback.js', '/source_scope.js',
                          '/message_history.js', '/vendor/katex.js', '/sw.js'):
                response = phone.get(route)
                response.raise_for_status()
                assert response.content
            qr = checked(desktop.get('/api/mobile/qr', params={'address': status['addresses'][0]}))
            assert '<svg' in qr['svg']
            for route in ('/user_guide.js', '/user_guide.css', '/help/guide.html', '/help/guide.js', '/help/guide.css'):
                response = phone.get(route)
                response.raise_for_status()
                assert len(response.content) > 500, route
            response = phone.get('/help/guide.html')
            assert '按此范围解析' in response.text and '完整提交' in response.text
            assert response.headers['X-Frame-Options'] == 'SAMEORIGIN'
            assert "frame-ancestors 'self'" in response.headers['Content-Security-Policy']
            report['checks'].append('trusted TLS, phone assets, bundled user guide and real SVG QR generation work from the frozen package')
            pairing = checked(phone.post('/mobile/pair', json={'code': status['pairing_code'], 'name': 'Frozen acceptance phone'}))
            checked(desktop.post(f"/api/mobile/devices/{pairing['device_id']}/approve", json={'approved': True}))
            ready = checked(phone.get(f"/mobile/pair/{pairing['device_id']}", headers={'x-mobile-pair-ticket': pairing['ticket']}))
            assert ready['status'] == 'active'
            checked(phone.get('/mobile/session'))
            assert checked(desktop.put('/api/settings', json={'api_connection_mode': 'direct'}))['api_connection_mode'] == 'direct'
            assert checked(desktop.get('/api/settings'))['api_connection_mode'] == 'direct'
            assert phone.get('/api/settings').status_code == 403
            checked(desktop.put('/api/settings', json={'api_connection_mode': 'system'}))
            report['checks'].append('desktop API route persists while sensitive global settings remain forbidden to paired phone')
            project = checked(phone.post('/api/projects', json={'name': '安装版手机验收'}))
            content = '安装版手机上传测试：原文件先入库，选择解析后才读取。'.encode('utf-8')
            uploaded = checked(phone.post('/api/import/file', data={'project_id': project['id']},
                files={'file': ('手机资料.txt', content, 'text/plain')}))
            assert uploaded['status'] == 'stored' and uploaded['job_id'] is None
            parsed = checked(phone.post(f"/api/sources/{uploaded['source_id']}/parse", json={'mode': 'full'}))
            for _ in range(100):
                job = checked(phone.get(f"/api/jobs/{parsed['job_id']}"))
                if job['status'] in ('completed', 'failed'):
                    break
                time.sleep(.1)
            assert job['status'] == 'completed', job.get('error')
            report['checks'].append('phone pairing, project creation, file staging and explicit parsing work through HTTPS')
            # A test-owned published file checks the real authenticated download
            # route. It is a local fixture, not an AI-generated teaching sample.
            artifact_path = data / 'artifacts' / 'phone-download.docx'
            artifact_path.parent.mkdir(parents=True, exist_ok=True)
            document = Document()
            document.add_paragraph('手机成果下载验收：文件内容与已发布的原文件一致。')
            document.save(artifact_path)
            artifact_id = uuid.uuid4().hex
            with sqlite3.connect(data / 'studyforge.sqlite3') as database:
                database.execute('INSERT INTO artifacts '
                    '(id,project_id,source_ids_json,kind,format,path,title,metadata_json,created_at) '
                    'VALUES (?,?,?,?,?,?,?,?,?)', (artifact_id, project['id'], '[]', 'lesson', 'docx',
                    str(artifact_path), '手机下载验收', json.dumps({'download_filename': '手机下载验收.docx'},
                    ensure_ascii=False), '2026-10-14T00:00:00+00:00'))
            listing = checked(phone.get(f"/api/projects/{project['id']}/artifacts"))
            assert any(item['id'] == artifact_id for item in listing)
            downloaded = phone.get(f'/api/artifacts/{artifact_id}/download')
            downloaded.raise_for_status()
            assert downloaded.content == artifact_path.read_bytes()
            assert '手机下载验收.docx' in unquote(downloaded.headers['Content-Disposition'])
            assert '文件内容与已发布的原文件一致' in Document(io.BytesIO(downloaded.content)).paragraphs[0].text
            report['checks'].append('paired phone lists and downloads the exact published DOCX with its readable Chinese filename')
            stop()
            restarted = start()
            assert restarted['running']
            assert ca_file.read_bytes() == ca
            phone.close()
            # Recreate the client to avoid reusing a connection to the old server.
            remembered = httpx.Client(base_url=f'https://127.0.0.1:{phone_port}', verify=context,
                                      trust_env=False, timeout=15, cookies=phone.cookies)
            phone = remembered
            assert checked(phone.get('/mobile/session'))['device']['id'] == pairing['device_id']
            assert any(item['id'] == project['id'] for item in checked(phone.get('/api/projects')))
            report['checks'].append('restarting the frozen backend preserves CA, trusted device session and project data')
            report['ok'] = True
        except BaseException as error:
            report['error'] = repr(error)
            raise
        finally:
            if phone:
                phone.close()
            stop()
            report['own_backend_stopped'] = process is None
            (root / 'result.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
