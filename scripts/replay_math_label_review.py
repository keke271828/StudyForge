"""Replay saved cloud review JSON through the final guard without new requests."""
import asyncio
import hashlib
import json
import os
from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[1]


async def main(name):
    out = ROOT / 'build' / name
    if out.exists() and any(out.iterdir()):
        raise SystemExit('Use a fresh folder.')
    out.mkdir(parents=True)
    os.environ['STUDYFORGE_DATA_DIR'] = str(out / 'data')
    os.environ['STUDYFORGE_SEMANTIC_MODEL_DIR'] = str(out / 'isolated-model')
    from studyforge.review import review_explanation
    from studyforge.exporters import export_docx
    from docx import Document
    from lxml import etree

    source_dir = ROOT / 'build/supplement-conditions-real-20261017-1'
    case = next(c for c in json.loads((source_dir / 'inputs.json').read_text(encoding='utf-8'))['cases']
                if c['id'] == 'historical_review')
    previous = json.loads((source_dir / 'results.json').read_text(encoding='utf-8'))
    rows = []
    for arm in ('baseline', 'candidate'):
        records = json.loads((source_dir / ('historical_review-' + arm + '-records.json')).read_text(encoding='utf-8'))
        old = next(r for r in previous if r['case'] == case['id'] and r['arm'] == arm)
        class Replay:
            calls = 0
            async def complete(self, *args, **kwargs):
                self.calls += 1
                assert self.calls == 1
                return records[0]['response']
        client = Replay()
        text, reports = await review_explanation(client, case['source'], case['draft'], mode='fast',
                                                explanation_style='balanced', requirements=case['question'], summary_level=3)
        assert client.calls == 1 and reports[0]['locally_corrected_issues'] == [0, 1]
        assert reports[0]['passed'] is False and '尚未再次独立复核' in text
        assert '待核实事项' not in text
        assert '才自然得到' not in text
        result = out / (arm + '-replayed-final.md')
        result.write_text(text, encoding='utf-8')
        old_path = source_dir / old['answer_file']
        old_text = old_path.read_bytes().decode('utf-8').replace('\r\n', '\n')
        assert hashlib.sha256(old_text.encode()).hexdigest() == old['answer_sha256']
        if arm == 'candidate':
            assert old_text == text  # The already matching record remains unchanged.
        words = []
        if arm == 'baseline':
            for content, filename in ((old_text, '格式修复前_旧核查交付.docx'), (text, '格式修复后_核查修订交付.docx')):
                path = out / filename
                export_docx(content, filename.removesuffix('.docx'), path)
                document = Document(path)
                word_text = ''.join(etree.fromstring(document.part.blob).xpath('//*[local-name()="t"]/text()'))
                assert '核查说明' in word_text and document.paragraphs
                if content is text:
                    assert '才自然得到' not in word_text and '尚未再次独立复核' in word_text
                words.append({'file': filename, 'paragraphs': len(document.paragraphs),
                              'tables': len(document.tables), 'structurally_reopened': True})
        rows.append({'arm': arm, 'before_file': str(old_path), 'after_file': result.name,
                     'before_resolved': old['reports'][0].get('locally_corrected_issues', []),
                     'after_resolved': reports[0]['locally_corrected_issues'],
                     'same_saved_review_response': True, 'replayed_response_count': client.calls,
                     'new_external_model_requests': 0, 'reports': reports, 'word_files': words,
                     'original_draft_sha256': hashlib.sha256(case['draft'].encode()).hexdigest(),
                     'saved_review_response_sha256': hashlib.sha256(records[0]['response'].encode()).hexdigest(),
                     'after_sha256': hashlib.sha256(result.read_bytes()).hexdigest(),
                     'meaning': 'patch_application_not_factual_or_independent_review_certification'})
    (out / 'results.json').write_text(json.dumps({'passed': True, 'source': case['source'],
        'requirements': case['question'], 'rows': rows, 'new_external_model_requests': 0,
        'total_replayed_responses': 2}, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({'passed': True, 'replayed_responses': 2, 'new_external_model_requests': 0,
                      'newly_applied_issue': 'baseline issue 1; previously matching candidate unchanged'}, ensure_ascii=True))


if __name__ == '__main__':
    asyncio.run(main(sys.argv[1]))
