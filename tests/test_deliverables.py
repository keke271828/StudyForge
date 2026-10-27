import asyncio
import json
import sqlite3
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

import fitz
import pytest
from docx import Document
from pptx import Presentation

from studyforge import server, deliverables
from studyforge.db import create_project,create_job,execute,j,now,one,many,uid,unj
from studyforge.exporters import export_document
from studyforge.document_render import blocks


SOURCE='给定恒定参数k>0，模型x(t)=cos(kt)，x(0)=1。这里是给定模型，不宣称适用所有系统。'
LESSON=('# 模型与条件\n\n> 给定恒定参数k>0。 [原文.txt｜行 1-1]\n\n'
        '这里仅解释给定模型，参数$k>0$；导数为$-k\\sin(kt)$。\n\n'
        '## 推导\n$$\nx(t)=\\cos(kt),\\quad \\frac{dx}{dt}=-k\\sin(kt)\n$$\n\n'
        '## 系统性总结\n保留参数条件和模型边界。')


def read_cards(path):
    with zipfile.ZipFile(path) as archive:
        with sqlite3.connect(':memory:') as connection:
            connection.deserialize(archive.read('collection.anki2'))
            return connection.execute('SELECT guid,flds FROM notes ORDER BY id').fetchall()


def test_atomic_cards_have_separate_questions_sources_safe_html_and_stable_update_guids():
    content={'cards':[{'question':'给定模型的参数k需要满足什么条件？','answer':'> k>0。','source':'[原文.txt｜行 1-1]'},
                      {'question':'给定模型的初始值是什么？','answer':'$x(0)=1$。 <script>bad()</script>','source':'原文'}]}
    first=export_document(j(content),'apkg','模型','cards.apkg');rows=read_cards(first)
    assert len(rows)==2 and all('来源：' in row[1] for row in rows)
    assert '<blockquote' in rows[0][1] and '<script>' not in rows[1][1] and r'\(x(0)=1\)' in rows[1][1]
    content['cards'][0]['answer']='> k为正数。'
    second=read_cards(export_document(j(content),'apkg','模型','updated.apkg'))
    assert rows[0][0]==second[0][0]
    with pytest.raises(ValueError,match='整章'): export_document(LESSON,'apkg','模型','chapter.apkg')


@pytest.mark.parametrize('payload',[
    {'cards':[]},{'cards':[{'question':'','answer':'A'}]},
    {'cards':[{'question':'Q','answer':'A'*2001}]},
    {'cards':[{'question':'Q','answer':'A'},{'question':'Q','answer':'B'}]},
])
def test_bad_cards_are_rejected_instead_of_silently_cutting_or_duplicating(payload):
    with pytest.raises(ValueError): export_document(j(payload),'apkg','卡片','bad.apkg')


def test_ppt_math_is_rendered_and_long_chapters_paginate_with_originals_in_notes():
    content=LESSON+'\n\n## 长段落\n'+('概念与边界需要逐项解释，不能在页尾被截断。'*80)+'最后条件仍须保留。'
    path=export_document(content,'pptx','排版','deck.pptx');prs=Presentation(path)
    assert len(prs.slides)>=6
    pictures=[shape for slide in prs.slides for shape in slide.shapes if shape.shape_type==13]
    assert len(pictures)>=2 and all(shape.height<prs.slide_height for shape in pictures)
    text='\n'.join(shape.text for slide in prs.slides for shape in slide.shapes if shape.has_text_frame)
    assert '\\frac' not in text and '最后条件仍须保留' in text
    assert '\\frac' in '\n'.join(s.notes_slide.notes_text_frame.text for s in prs.slides)
    for slide in prs.slides:
        for shape in slide.shapes:
            assert shape.left>=0 and shape.top>=0 and shape.top+shape.height<=prs.slide_height


def test_quote_math_and_code_are_not_reinterpreted_as_explanation_or_headings():
    content='> $$\n> F=ma\n> $$\n\n~~~python\n# not a heading\ntext=r"\\[not math\\]"\n~~~\n\n\\[x^2+1\\]'
    parsed=list(blocks(content));assert [kind for kind,_ in parsed]==['quote_math','code','math']
    assert '\\[not math\\]' in parsed[1][1]
    doc=Document(export_document(content,'docx','来源公式','quote.docx'))
    quoted=next(p for p in doc.paragraphs if p.style.name=='Source Quote')
    assert quoted._p.xpath('.//m:oMath')
    with fitz.open(export_document(content,'pdf','来源公式','quote.pdf')) as pdf:
        assert len(pdf)>=1 and sum(len(p.get_images()) for p in pdf)>=2


def test_notebook_keeps_cells_separate_and_does_not_execute_original_error_examples():
    content='# 原文错误示例\n```python\nthis is invalid Python\n```\n\n## 可运行模型或示例\n参数为正。\n```python\nk=1\n```\n解释输出。\n```python\nx=k+1\n```\n\n## 系统性总结\n边界。'
    path=export_document(content,'ipynb','模型','model.ipynb');value=json.loads(path.read_text(encoding='utf-8'))
    cells=value['cells'];code=[c for c in cells if c['cell_type']=='code']
    assert len(code)==2 and ''.join(code[0]['source'])=='k=1'
    assert all(c['outputs']==[] and c['execution_count'] is None for c in code)
    assert all(c.get('id') for c in cells) and len({c['id'] for c in cells})==len(cells)
    assert 'invalid Python' in ''.join(cells[0]['source']) and '系统性总结' in ''.join(cells[-1]['source'])
    assert all(c['syntax_checked'] for c in deliverables.inspection(path)['checks'])


@pytest.mark.parametrize('heading', ['Runnable model or example', '可运行模型或示例 / Runnable model or example'])
@pytest.mark.parametrize('fmt', ['py','ipynb'])
def test_translated_delivery_heading_excludes_original_error_code(heading, fmt):
    content = '# Original sample\n```python\nBAD ORIGINAL CODE\n```\n\n## '+heading+'\n```python\nx = 1\n```\n\n## Systematic summary\nPreserve conditions.'
    path = export_document(content, fmt, 'Example', 'translated.'+fmt)
    if fmt == 'py':
        text = path.read_text(encoding='utf-8')
        assert 'BAD ORIGINAL CODE' not in text and 'x = 1' in text and 'Execution status:' in text
    else:
        cells = json.loads(path.read_text(encoding='utf-8'))['cells']
        assert [''.join(c['source']) for c in cells if c['cell_type']=='code'] == ['x = 1']
        assert 'BAD ORIGINAL CODE' in ''.join(cells[0]['source']) and 'Execution status:' in ''.join(cells[0]['source'])


@pytest.mark.parametrize('fmt,label,code',[('py','python','k=1\nprint(k)'),('js','javascript','const k=1;'),
    ('ts','typescript','const k: number = 1;'),('cpp','c++','int main(){return 0;}'),
    ('m','matlab','k=1; x=cos(k);'),('wl','wolfram','k=1; x=Cos[k];')])
def test_code_formats_choose_final_delivery_and_report_actual_verification(fmt,label,code):
    content=f'# 原文\n```{label}\nBAD ORIGINAL EXAMPLE\n```\n\n## 可运行模型或示例\n```{label}\n{code}\n```'
    path=export_document(content,fmt,'模型',f'model.{fmt}');text=path.read_text(encoding='utf-8')
    assert 'BAD ORIGINAL' not in text and code in text and '尚未执行验证' in text
    status=deliverables.inspection(path)
    assert status['executed'] is False and status['checks'][0]['syntax_checked']==(fmt in ('py','js'))


@pytest.mark.parametrize('name',['../outside.py','/absolute.py','C:/secret.py','sub\\x.py','CON.py','sub/aux.txt','name?.py','sub/../x.py'])
def test_bundle_rejects_windows_invalid_and_escaping_paths(name):
    with pytest.raises(ValueError): export_document(j({'files':[{'path':name,'content':'x=1'}]}),'zip','程序','bad.zip')


def test_bundle_keeps_file_names_and_dependencies_and_does_not_run_code(tmp_path):
    marker=tmp_path/'executed.txt'
    code=f"from pathlib import Path\nPath({str(marker)!r}).write_text('bad')"
    files={'files':[{'path':'src/main.py','content':code},{'path':'README.md','content':'参数k>0；尚未运行；python src/main.py'},
                    {'path':'requirements.txt','content':'# 标准库，无额外依赖'}]}
    path=export_document(j(files),'zip','程序','bundle.zip')
    with zipfile.ZipFile(path) as archive:
        assert 'src/main.py' in archive.namelist() and archive.read('src/main.py').decode()==code
        status=json.loads(archive.read('StudyForge-validation.json'))
        assert status['executed'] is False and status['checks'][0]['syntax_checked']
    assert not marker.exists()


def test_mindmap_preserves_nested_concepts_and_code_headings_are_not_nodes():
    content='# 章\n## 概念\n### 条件\n条件细节。\n```python\n# not heading\nprint(1)\n```'
    tree=ET.parse(export_document(content,'mm','结构','nested.mm'))
    assert tree.find("./node/node[@TEXT='章']/node[@TEXT='概念']/node[@TEXT='条件']") is not None
    assert not any(n.get('TEXT')=='not heading' for n in tree.iter('node'))


def test_word_math_keeps_functions_and_units_upright_and_double_struck_symbols():
    from studyforge.document_render import word_math
    from lxml import etree
    result=etree.tostring(word_math(r'\cos(x)+\mathrm{kg}+\mathbb{R}'),encoding='unicode')
    assert '<m:t>cos</m:t>' in result and result.count('m:val="p"')>=3 and 'ℝ' in result


def test_table_formula_absolute_value_and_escaped_pipes_stay_in_the_same_cell():
    content='| 对象 | 条件 |\n| --- | --- |\n| x | $|x|\\le1$ |\n| 记号 | a\\|b |'
    kind,rows=list(blocks(content))[0]
    assert kind=='table' and rows[1]==['x',r'$|x|\le1$'] and rows[2]==['记号','a|b']
    for fmt in ('docx','pdf','pptx'): assert export_document(content,fmt,'绝对值',f'absolute.{fmt}').is_file()


def test_javascript_hashbang_stays_first_and_syntax_is_checked_without_running(tmp_path):
    marker=tmp_path/'executed.txt'
    content='```javascript\n#!/usr/bin/env node\nthrow new Error("must not run");\n```'
    path=export_document(content,'js','脚本','hashbang.js')
    assert path.read_text(encoding='utf-8').startswith('#!/usr/bin/env node\n// 运行状态')
    assert deliverables.inspection(path)['checks'][0]['syntax_checked'] and not marker.exists()


def test_matlab_script_uses_legal_name_and_local_functions_do_not_override_it():
    from fastapi.testclient import TestClient
    project=create_project('中文模型')
    content='```matlab\nx=1;\nfunction y=local_helper(x)\ny=x;\nend\n```'
    path=export_document(content,'m','中文模型','opaque-id.m')
    assert deliverables.matlab_identifier(path.stem) and path.stem.startswith('studyforge_model_')
    artifact=server.add_artifact(project['id'],[],'explanation','m','中文模型',path)
    assert artifact['download_filename']==path.name
    renamed=TestClient(server.app).put('/api/artifacts/'+artifact['id'],json={'name':'我的新模型'})
    assert renamed.status_code==200 and renamed.json()['title']=='我的新模型'
    assert deliverables.matlab_identifier(Path(renamed.json()['download_filename']).stem)


@pytest.mark.parametrize('code,entry',[
    ('% 注释\nfunction y=oscillator_model(t)\ny=cos(t);\nend','oscillator_model'),
    ('%{\n说明\n%}\nfunction show_model % 说明\ndisp(1);\nend','show_model'),
    ('classdef (Sealed) SampleModel\nend','SampleModel')])
def test_matlab_function_and_class_names_match_files_and_versions_do_not_overwrite(code,entry):
    from fastapi.testclient import TestClient
    project=create_project('模型');content='```matlab\n'+code+'\n```'
    first=export_document(content,'m','中文标题','version-one.m')
    second=export_document(content,'m','另一标题','version-two.m')
    assert first.name==second.name==entry+'.m' and first!=second and first.is_file()
    artifact=server.add_artifact(project['id'],[],'explanation','m','中文标题',first)
    renamed=TestClient(server.app).put('/api/artifacts/'+artifact['id'],json={'name':'中文成果改名'})
    assert renamed.json()['title']=='中文成果改名' and renamed.json()['download_filename']==entry+'.m'


@pytest.mark.parametrize('name',['中文.m','mismatch.m'])
def test_matlab_bundle_names_must_be_identifiers_and_match_primary_function(name):
    with pytest.raises(ValueError,match='MATLAB文件名'):
        deliverables.bundle_files({'files':[{'path':name,'content':'function y=real_model(x)\ny=x;\nend'}]})


def test_bundle_file_cannot_also_be_used_as_directory():
    with pytest.raises(ValueError,match='文件和目录'):
        deliverables.bundle_files({'files':[{'path':'model.py','content':'x=1'}, {'path':'model.py/sub.py','content':'x=2'}]})


def test_legacy_invalid_matlab_code_does_not_break_artifact_listing():
    project=create_project('旧成果');from studyforge.config import artifacts_dir
    path=artifacts_dir()/'旧模型.m';path.write_text('function 123bad\nend',encoding='utf-8')
    artifact=server.add_artifact(project['id'],[],'explanation','m','旧模型',path)
    assert deliverables.matlab_identifier(Path(artifact['download_filename']).stem)


@pytest.mark.parametrize('exercise_format',['docx','apkg'])
def test_lecture_cards_do_not_change_independent_exercise_format_or_leak_private_answers(tmp_path,monkeypatch,exercise_format):
    project=create_project('独立格式');sid=uid();path=tmp_path/'原文.txt';path.write_text(SOURCE,encoding='utf-8')
    execute('INSERT INTO sources(id,project_id,name,kind,original_path,status,created_at) VALUES (?,?,?,?,?,?,?)',
            (sid,project['id'],'原文.txt','.txt',str(path),'ready',now()))
    long_question='请解释并论证：'+'这个模型的条件需要说明。'*40
    class Client:
        model='same'
        async def complete(self,system,user,**kwargs):
            if '【Anki格式整理】' in system: return j({'cards':[{'question':'k的条件是什么？','answer':'k>0','source':'[原文.txt｜行 1-1]'}]})
            return LESSON
    monkeypatch.setattr(server,'client_for',lambda *args,**kwargs:Client())
    async def create_set(*args,**kwargs):
        set_id=uid()
        execute('INSERT INTO exercise_sets VALUES (?,?,?,?,?,?,?,?,?)',
                (set_id,project['id'],j([sid]),'题组',j([{'question':long_question,'topic':'模型','index':0}]),
                 j([{'solution':'SECRET_SAMPLE_SOLUTION','rubric':'私有标准'}]),1,'assigned',now()))
        return set_id,[]
    monkeypatch.setattr(server,'generate_set',create_set)
    value=server.GenerateIn(project_id=project['id'],source_ids=[sid],model='same',format='apkg',exercises=True,
                            exercise_format=exercise_format,review_mode='express')
    job=create_job(project['id'],'generate',value.model_dump());asyncio.run(server._generate_job(job,value))
    row=one('SELECT * FROM jobs WHERE id=?',(job,));assert row['status']=='completed',row['error']
    outputs=many('SELECT * FROM artifacts');assert len(outputs)==2
    lecture=next(x for x in outputs if x['kind']=='explanation');exercise=next(x for x in outputs if x['kind']=='exercise')
    assert lecture['format']=='apkg' and exercise['format']==exercise_format
    contents='\n'.join(fields for _,fields in read_cards(Path(exercise['path']))) if exercise_format=='apkg' else '\n'.join(p.text for p in Document(exercise['path']).paragraphs)
    assert long_question in contents and 'SECRET_SAMPLE_SOLUTION' not in contents and '私有标准' not in contents


@pytest.mark.parametrize('fmt',['apkg','zip','py','pptx'])
def test_generation_shape_uses_same_client_one_artifact_and_preserves_no_exercises(tmp_path,monkeypatch,fmt):
    project=create_project('格式');sid=uid();path=tmp_path/'原文.txt';path.write_text(SOURCE,encoding='utf-8')
    execute('INSERT INTO sources(id,project_id,name,kind,original_path,status,created_at) VALUES (?,?,?,?,?,?,?)',
            (sid,project['id'],'原文.txt','.txt',str(path),'ready',now()))
    calls=[]
    class Client:
        model='same-mentor';effort='high'
        async def complete(self,system,user,**kwargs):
            calls.append((id(self),system,user))
            if '【Anki格式整理】' in system:
                return j({'cards':[{'question':'该模型的k满足什么条件？','answer':'k>0。','source':'[原文.txt｜行 1-1]'},
                                   {'question':'模型是否普适？','answer':'并非普适。','source':'[原文.txt｜第99页]'}]})
            if '【代码包格式整理】' in system: return j({'files':[{'path':'model.py','content':'k=1\nprint(k)'},{'path':'README.md','content':'参数及条件；未运行'}]})
            if '实现为可运行的 Python' in system: return '```python\nk=1\nprint(k)\n```'
            return LESSON
    instance=Client();monkeypatch.setattr(server,'client_for',lambda *args,**kwargs:instance)
    value=server.GenerateIn(project_id=project['id'],source_ids=[sid],model='same-mentor',format=fmt,exercises=False,review_mode='express')
    job=create_job(project['id'],'generate',value.model_dump());asyncio.run(server._generate_job(job,value))
    row=one('SELECT * FROM jobs WHERE id=?',(job,));assert row['status']=='completed',row['error']
    outputs=many('SELECT * FROM artifacts');assert len(outputs)==1 and outputs[0]['format']==fmt
    assert len({row[0] for row in calls})==1 and len(calls)==(1 if fmt=='pptx' else 2)
    assert many('SELECT * FROM exercise_sets')==[]
    meta=unj(outputs[0]['metadata_json']);assert meta['delivery']['executed'] is False
    if fmt=='apkg':
        assert meta['delivery']['card_count']==2 and meta['delivery']['source_warning_count']==1
        assert any('来源定位需核对' in fields for _,fields in read_cards(Path(outputs[0]['path'])))
    if fmt in ('apkg','zip'): assert meta['delivery']['format_model_requests']==1
