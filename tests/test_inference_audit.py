"""Check the real review path, unchanged review budgets, and saved evidence."""
import asyncio
import hashlib
import json

import pytest
from docx import Document

from studyforge import review
from studyforge.exporters import export_document


def test_fast_review_requests_claim_and_example_checks_once():
    class Client:
        calls = []
        async def complete(self, system, user, **kwargs):
            self.calls.append((system, user))
            return json.dumps({'passed': True, 'issues': []})
    client = Client()
    result, reports = asyncio.run(review.review_explanation(client, '原文', '现有讲解', mode='fast'))
    assert len(client.calls) == 1 and result == '现有讲解'
    assert '通常、往往、倾向' in client.calls[0][0]
    assert '解释同一个对象' in client.calls[0][0]
    assert '假设的示意按其假设核对' in client.calls[0][0]
    assert '其他有依据的结论应保留' in client.calls[0][0]
    receipt = reports[0]['inference_scope']
    assert receipt['criteria_sha256'] == hashlib.sha256(review.INFERENCE_SCOPE_RULES.encode()).hexdigest()
    assert receipt['meaning'] == 'requested_checks_not_accuracy_certification'


@pytest.mark.parametrize('mode,bad,calls', [
    ('express', False, 0), ('express', True, 0), ('fast', False, 1), ('fast', True, 1),
    ('balanced', False, 1), ('balanced', True, 2), ('strict', False, 3), ('strict', True, 7),
])
def test_inference_checks_do_not_add_review_rounds_or_expand_scope(mode, bad, calls):
    class Client:
        requests = []
        async def complete(self, system, user, json_mode=False):
            self.requests.append((system, user, json_mode))
            if not json_mode: return '# 修订\n只说明范围内能成立的结论。'
            failed = bad and 'INITIAL_DRAFT' in user
            return json.dumps({'passed': not failed, 'issues': [{'problem': '本句增加了没有依据的方向判断'}] if failed else []})
    client = Client()
    _, reports = asyncio.run(review.review_explanation(client, '第一节。其他章节。', 'INITIAL_DRAFT', mode=mode,
        requirements='只讲第一节，不新增习题，不扩写其他章。', summary_level=1, section='本次只检查第一节'))
    assert len(client.requests) == calls
    audits = [row for row in client.requests if row[2]]
    for system, user, _ in audits:
        assert system.count(review.INFERENCE_SCOPE_RULES) == 1
        assert '只讲第一节，不新增习题' in user and '本次只检查第一节' in user
    receipts = [row for row in reports if 'inference_scope' in row]
    assert len(receipts) == len(audits)
    if mode == 'express': assert reports[0]['passed'] is None and not receipts


def test_local_scope_correction_and_unresolved_doubt_survive_word_export():
    class Client:
        calls = 0
        async def complete(self, *args, **kwargs):
            self.calls += 1
            return json.dumps({'passed': False, 'issues': [
                {'location': '补充结论', 'problem': '没有排除其他路径，不能确定效应缩小。'},
                {'location': '例子', 'problem': '例子尚未解释原来的问题，需要回到同一个对象。'}],
                'corrections': [{'issue_index': 0, 'before': '通常会把效应压小。',
                                'after': '缺少对其他作用路径的限制，不能确定效应大小或方向；原随机分配的比较仍有价值。'}]})
    client = Client()
    draft = '# 范围与条件\n\n> 分配提醒。[材料｜段1]\n\n通常会把效应压小。\n\n例子需要核对。\n\n## 系统性总结\n区分两种对象。'
    result, reports = asyncio.run(review.review_explanation(client, '没有排除其他路径。', draft, mode='fast'))
    assert client.calls == 1 and reports[0]['locally_corrected_issues'] == [0]
    assert '> 分配提醒。[材料｜段1]' in result and '通常会把效应压小。' not in result
    assert '待核实事项' in result and '例子尚未解释原来的问题' in result
    assert '尚未再次独立复核' in result and len(reports[0]['issues']) == 2
    path = export_document(result, 'docx', '讲解范围核验', 'inference-check.docx')
    text = '\n'.join(paragraph.text for paragraph in Document(path).paragraphs)
    assert '不能确定效应大小或方向' in text and '例子尚未解释原来的问题' in text
    assert '原随机分配的比较仍有价值' in text


def test_native_generation_delivers_quote_anchored_fix_and_the_actual_review_receipt(tmp_path, monkeypatch):
    import httpx
    from studyforge import conversations, server
    from studyforge.ai import save_settings
    from studyforge.db import create_project, create_job, execute, j, many, now, one, uid, unj
    from tests.test_review_quote_anchors import PROSE, FIXED, patch
    save_settings({'agent_tools_enabled': False, 'background_learning_enabled': False})
    project = create_project('讲解依据核验'); chat = conversations.create(project['id'], 'deepseek', 'fixture-model')
    source, path = uid(), tmp_path/'原文.txt'; path.write_text('当前条件未限制方向，不能确定效应缩小。', encoding='utf-8')
    execute('INSERT INTO sources(id,project_id,name,kind,original_path,status,created_at) VALUES (?,?,?,?,?,?,?)',
        (source, project['id'], path.name, '.txt', str(path), 'ready', now()))
    draft = '# 讲解范围\n\n'+PROSE+'\n\n## 系统性总结\n保留成立部分，不把未明条件当事实。'
    calls = []; native = httpx.AsyncClient
    def respond(request):
        body = json.loads(request.content); calls.append(body)
        if body.get('response_format', {}).get('type') == 'json_object':
            assert body['messages'][0]['content'].count(review.INFERENCE_SCOPE_RULES) == 1
            answer = j({'passed': False, 'issues': [{'problem': '新增方向判断无依据'}], 'corrections': [patch()]})
        else: answer = draft
        return httpx.Response(200, text='data: '+j({'choices': [{'delta': {'content': answer}, 'finish_reason': 'stop'}]})+'\n\ndata: [DONE]\n\n')
    monkeypatch.setattr('studyforge.ai.get_secret', lambda key: 'fixture-only')
    monkeypatch.setattr(httpx, 'AsyncClient', lambda *a, **kw: native(transport=httpx.MockTransport(respond)))
    value = server.GenerateIn(project_id=project['id'], conversation_id=chat['id'], source_ids=[source],
        provider='deepseek', model='fixture-model', format='docx', exercises=False, review_mode='fast')
    job = create_job(project['id'], 'generate', value.model_dump())
    asyncio.run(server._generate_job(job, value))
    assert one('SELECT status,error FROM jobs WHERE id=?', (job,)) == {'status': 'completed', 'error': None}
    assert len(calls) == 2, 'One writing call plus the existing one review call; no extra review or rewrite.'
    saved = one('SELECT * FROM document_snapshots WHERE job_id=?', (job,))
    report = unj(saved['metadata_json'])['reviews'][0]
    assert report['inference_scope']['version'] == review.INFERENCE_SCOPE_VERSION
    assert report['locally_corrected_issues'] == [0] and FIXED in saved['content']
    artifact = one("SELECT * FROM artifacts WHERE kind='explanation'")
    word = '\n'.join(paragraph.text for paragraph in Document(artifact['path']).paragraphs)
    assert FIXED in word and PROSE not in word and '尚未再次独立复核' in word
    assert not many('SELECT * FROM exercise_sets')
